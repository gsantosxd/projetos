import os, re, json, time, sqlite3, threading, unicodedata, urllib.parse, webbrowser, sys, logging, socket, glob, shutil, subprocess, hashlib
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, date, timedelta, timezone
from logging.handlers import RotatingFileHandler
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from email.utils import parsedate_to_datetime
from xml.etree import ElementTree

try:
    import requests
    from bs4 import BeautifulSoup
except Exception:
    requests = None
    BeautifulSoup = None

APP_TITLE="Tô no Corre"
APP_VERSION="0.9.4-beta-estável"
AUTHOR_LINKEDIN="https://www.linkedin.com/in/gabriel-santos-8667bb1b7/"
AUTHOR_GITHUB="https://github.com/gsantosxd"
AUTHOR_EMAIL="gsantosxd3@gmail.com"
DONATION_PIX="27998886868"
PRIVACY_NOTICE_VERSION="3"
TERMS_OF_USE_VERSION="2"
PCD_CONSENT_VERSION="1"
PRIVACY_CONTACT=f"{AUTHOR_EMAIL} | {AUTHOR_LINKEDIN}"
APP_DIR=os.path.dirname(os.path.abspath(__file__))

def resource_path(name):
    """Localiza recursos tanto no código-fonte quanto no pacote PyInstaller."""
    root=getattr(sys,"_MEIPASS",APP_DIR)
    return os.path.join(root,name)
APP_DATA_VERSION="BetaEstavel3"
LEGACY_APP_DATA_VERSIONS=()
if getattr(sys,"frozen",False):
    APP_DATA_ROOT=os.path.join(os.environ.get("LOCALAPPDATA",os.path.dirname(sys.executable)),"ToNoCorre")
    BASE_DIR=os.path.join(APP_DATA_ROOT,APP_DATA_VERSION)
    os.makedirs(BASE_DIR,exist_ok=True)
else:
    APP_DATA_ROOT=APP_DIR
    BASE_DIR=APP_DIR
DB_PATH=os.path.join(BASE_DIR,"vagas.db")
PROFILE_PATH=os.path.join(BASE_DIR,"perfil.json")
CV_PATH=os.path.join(BASE_DIR,"curriculo.txt")
CV_FILE_PATH=os.path.join(BASE_DIR,"curriculo_original")
CACHE_PATH=os.path.join(BASE_DIR,"cache_detalhes.json")
SOURCE_RESULTS_PATH=os.path.join(BASE_DIR,"cache_resultados_fontes.json")
LOG_PATH=os.path.join(BASE_DIR,"to_no_corre.log")
BACKUP_DIR=os.path.join(BASE_DIR,"backups")

def apply_windows_data_permissions():
    """Garante acesso do usuário antes de abrir qualquer arquivo e restringe a pasta."""
    if os.name!="nt" or not getattr(sys,"frozen",False):return False
    domain=os.environ.get("USERDOMAIN","").strip();username=os.environ.get("USERNAME","").strip()
    account=(domain+"\\"+username).strip("\\")
    if not account:return False
    try:
        flags=getattr(subprocess,"CREATE_NO_WINDOW",0)
        # Arquivos existentes recebem uma ACE efetiva (sem OI/CI). Aplicar
        # somente OI/CI com /T os deixaria com uma regra apenas herdável, sem
        # acesso ao próprio arquivo na execução seguinte.
        repair=subprocess.run([
            "icacls",BASE_DIR,"/inheritance:r",
            "/grant:r",f"{account}:F","*S-1-5-18:F","*S-1-5-32-544:F","/T","/C"
        ],capture_output=True,text=True,timeout=20,creationflags=flags)
        if repair.returncode!=0:return False
        # A pasta recebe, adicionalmente, regras herdáveis para novos arquivos.
        protect=subprocess.run([
            "icacls",BASE_DIR,
            "/grant",f"{account}:(OI)(CI)F",
            "/grant","*S-1-5-18:(OI)(CI)F",
            "/grant","*S-1-5-32-544:(OI)(CI)F","/C"
        ],capture_output=True,text=True,timeout=20,creationflags=flags)
        return protect.returncode==0
    except Exception:
        return False

# A permissão precisa ser aplicada antes do RotatingFileHandler. Fazer isso
# depois pode deixar o log aberto sem a nova ACL efetiva para a próxima execução.
if getattr(sys,"frozen",False):apply_windows_data_permissions()

def configure_logging():
    root=logging.getLogger()
    for handler in list(root.handlers):
        try:handler.close()
        finally:root.removeHandler(handler)
    handler=RotatingFileHandler(LOG_PATH,maxBytes=1_000_000,backupCount=3,encoding="utf-8")
    handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
    root.setLevel(logging.INFO);root.addHandler(handler)

configure_logging()
LOGGER=logging.getLogger("to_no_corre")

def migrate_legacy_data_directory():
    """Importa uma única vez os dados da última edição beta para a pasta permanente."""
    if not getattr(sys,"frozen",False):return False
    migrated=False
    for version in LEGACY_APP_DATA_VERSIONS:
        legacy=os.path.join(APP_DATA_ROOT,version)
        marker=os.path.join(BASE_DIR,f".migrated_{version}")
        if os.path.isfile(marker) or not os.path.isdir(legacy):continue
        try:
            legacy_db=os.path.join(legacy,"vagas.db")
            if os.path.isfile(legacy_db) and not os.path.isfile(DB_PATH):
                source=sqlite3.connect(f"file:{urllib.parse.quote(os.path.abspath(legacy_db))}?mode=ro",uri=True,timeout=8)
                destination=sqlite3.connect(DB_PATH)
                try:source.backup(destination)
                finally:destination.close();source.close()
                migrated=True
            for name in ("perfil.json","curriculo.txt","cache_detalhes.json"):
                source_path=os.path.join(legacy,name);target_path=os.path.join(BASE_DIR,name)
                if os.path.isfile(source_path) and not os.path.exists(target_path):
                    shutil.copy2(source_path,target_path);migrated=True
            for source_path in glob.glob(os.path.join(legacy,"curriculo_original.*")):
                target_path=os.path.join(BASE_DIR,os.path.basename(source_path))
                if not os.path.exists(target_path):shutil.copy2(source_path,target_path);migrated=True
            legacy_browser=os.path.join(legacy,"browser_profiles")
            current_browser=os.path.join(BASE_DIR,"browser_profiles")
            if os.path.isdir(legacy_browser) and not os.path.exists(current_browser):
                shutil.copytree(legacy_browser,current_browser);migrated=True
            with open(marker,"w",encoding="utf-8") as stream:stream.write(datetime.now(timezone.utc).isoformat())
        except Exception:
            LOGGER.exception("Não foi possível migrar os dados da edição %s",version)
    return migrated

migrate_legacy_data_directory()

PRIVACY_NOTICE=f"""Tô no Corre {APP_VERSION} — Aviso de privacidade

Finalidade
O aplicativo lê o currículo para localizar, organizar e comparar vagas. O resultado é apenas uma recomendação ao próprio usuário e não decide contratações.

Dados mantidos neste computador
Currículo original e texto extraído; áreas, cursos e preferências identificadas; cidades; vagas, fila, candidaturas e descartes; cache, diagnóstico e perfil separado do navegador usado nas candidaturas.

Comunicações externas
As fontes de vagas recebem termos profissionais e localidades usados nas pesquisas. O currículo completo não é enviado às fontes durante a busca. Ao iniciar uma candidatura, o navegador acessa o site escolhido e pode preencher dados ou anexar o currículo para revisão do usuário. O aplicativo não envia a candidatura automaticamente.

PCD
A opção de vagas para PCD é facultativa e pode permitir inferência relacionada à saúde. Ela exige uma confirmação separada e pode ser desativada a qualquer momento.

Armazenamento e exclusão
Não há servidor próprio, telemetria nem venda de dados. Os arquivos permanecem localmente até o usuário usar “Limpar tudo / trocar currículo”. A limpeza remove currículo, banco, backups, cache, diagnóstico, cookies e sessões do navegador desta edição. Administradores do computador ou pessoas com acesso à conta do Windows ainda podem acessar dados locais.

Serviços independentes
Sites de vagas possuem suas próprias políticas de privacidade e passam a tratar os dados fornecidos durante a candidatura.

Canal de privacidade: {PRIVACY_CONTACT}."""

TERMS_OF_USE=f"""Tô no Corre {APP_VERSION} — Termos de Uso (versão {TERMS_OF_USE_VERSION})

1. Finalidade e fase beta
O Tô no Corre é uma ferramenta gratuita em fase beta para auxiliar o próprio usuário a localizar, comparar e organizar vagas. Ele não é agência de emprego, recrutador ou representante dos sites consultados e não garante entrevista, contratação, disponibilidade ou exatidão dos anúncios.

2. Conferência e candidatura
Compatibilidade, modalidade, localização, salário, requisitos e datas são estimativas baseadas nas informações disponíveis. O usuário deve conferir o anúncio original antes de se candidatar. O aplicativo pode abrir páginas, preparar campos e anexos, mas a revisão e a decisão final de enviar uma candidatura são sempre do usuário.

3. Fontes independentes
Gupy, LinkedIn, Google, Indeed e as demais fontes são serviços independentes, sujeitos aos próprios termos, políticas, disponibilidade e limitações. A presença de uma marca ou link não indica parceria, patrocínio ou aprovação. Uma fonte pode alterar ou impedir o acesso sem aviso.

4. Uso responsável
O aplicativo destina-se a uso pessoal e lícito. O usuário não deve empregá-lo para sobrecarregar serviços, contornar controles de acesso, enviar candidaturas enganosas, tratar dados de outra pessoa sem autorização ou praticar atos que violem direitos e regras aplicáveis. Recursos automatizados devem permanecer sob supervisão humana.

5. Dados e troca de usuário
Os dados são mantidos localmente conforme o Aviso de Privacidade. Quem usar o mesmo computador para outra pessoa deve executar “Limpar tudo / trocar currículo” antes de carregar um novo currículo.

6. Disponibilidade, alterações e encerramento
Por se tratar de beta gratuito, funções podem apresentar falhas, mudar ou ficar temporariamente indisponíveis. Alterações relevantes destes termos ou do aviso de privacidade exigirão nova aceitação por versão. O usuário pode deixar de usar o aplicativo e apagar os dados locais a qualquer momento.

7. Responsabilidade e direitos
O desenvolvedor busca manter informações claras e medidas razoáveis de segurança, mas não controla anúncios, processos seletivos ou serviços externos. Nada nestes termos limita direitos ou responsabilidades que não possam ser afastados pela legislação aplicável.

8. Contato
E-mail: {AUTHOR_EMAIL}
LinkedIn: {AUTHOR_LINKEDIN}
GitHub: {AUTHOR_GITHUB}"""

def protect_local_data_directory():
    """Restringe a pasta da edição ao usuário do Windows no aplicativo empacotado."""
    return apply_windows_data_permissions()

if getattr(sys,"frozen",False):protect_local_data_directory()

class ThreadSafeConnection(sqlite3.Connection):
    """Serializa transações feitas pela interface e pelos workers em segundo plano."""
    def __init__(self,*args,**kwargs):
        super().__init__(*args,**kwargs);self._write_lock=threading.RLock()
    def execute(self,*args,**kwargs):
        with self._write_lock:return super().execute(*args,**kwargs)
    def executemany(self,*args,**kwargs):
        with self._write_lock:return super().executemany(*args,**kwargs)
    def commit(self):
        with self._write_lock:return super().commit()
    def rollback(self):
        with self._write_lock:return super().rollback()
    def __enter__(self):
        self._write_lock.acquire()
        try:return super().__enter__()
        except Exception:self._write_lock.release();raise
    def __exit__(self,*args):
        try:return super().__exit__(*args)
        finally:self._write_lock.release()

def connect_database(path):
    conn=sqlite3.connect(path,check_same_thread=False,factory=ThreadSafeConnection,timeout=8)
    conn.execute("PRAGMA busy_timeout=8000")
    if path not in (":memory:",""):
        conn.execute("PRAGMA journal_mode=WAL")
        conn.execute("PRAGMA synchronous=NORMAL")
    return conn

def load_json_file(path,default):
    try:
        with open(path,encoding="utf-8") as f:return json.load(f)
    except Exception:return default

def save_json_file(path,data):
    temp_path=""
    try:
        os.makedirs(os.path.dirname(os.path.abspath(path)),exist_ok=True)
        temp_path=path+f".{os.getpid()}.{threading.get_ident()}.tmp"
        with open(temp_path,"w",encoding="utf-8") as f:
            json.dump(data,f,ensure_ascii=False,indent=2);f.flush();os.fsync(f.fileno())
        os.replace(temp_path,path)
    except Exception:
        LOGGER.exception("Não foi possível salvar o arquivo JSON: %s",path)
        try:
            if temp_path and os.path.isfile(temp_path):os.remove(temp_path)
        except OSError:pass

DETAIL_CACHE=load_json_file(CACHE_PATH,{})
CACHE_LOCK=threading.Lock()
DETAIL_CACHE_DIRTY=0
DETAIL_CACHE_LAST_SAVE=time.time()
SOURCE_API_CACHE={}
SOURCE_API_CACHE_LOCK=threading.Lock()
SOURCE_RESULTS_CACHE=load_json_file(SOURCE_RESULTS_PATH,{"version":1,"entries":{}})
if not isinstance(SOURCE_RESULTS_CACHE,dict) or SOURCE_RESULTS_CACHE.get("version")!=1:
    SOURCE_RESULTS_CACHE={"version":1,"entries":{}}
SOURCE_RESULTS_LOCK=threading.Lock()
SOURCE_RESULT_METRICS={}

def source_search_signature(source,p):
    relevant={
        "source":source,
        # A revisão invalida apenas caches do LinkedIn quando a estratégia de
        # descoberta muda, sem obrigar as demais fontes a repetir consultas.
        "collector_revision":2 if "linkedin" in norm(source) else 1,
        "gupy":p.get("consultas_gupy",[]),"linkedin":p.get("consultas_linkedin",[]),
        "google":p.get("consultas_google",[]),"english":p.get("consultas_ingles",[]),
        "cities":p.get("cidades_presencial",[]),"state":p.get("estado_local",""),
        "remote":bool(p.get("aceitar_remoto",True)),
        "internships":internship_search_mode(p),
        "internship_areas":internship_selected_areas(p),
        "apprentice":bool(p.get("buscar_jovem_aprendiz",False)),
        "entry":bool(p.get("buscar_vagas_inicio_carreira",False)),
        "international":international_search_enabled(p),
        # As fontes já devolvem resultados recortados pela idade. Portanto,
        # caches criados para 7 dias não podem ser reutilizados numa busca de
        # 60 dias, pois não contêm as vagas mais antigas que voltaram a valer.
        "publication_age_days":int(p.get(
            "idade_maxima_dias",p.get("idade_maxima_vaga_dias",60))),
    }
    raw=json.dumps(relevant,ensure_ascii=False,sort_keys=True,separators=(",",":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()

def persistent_source_fetch(source,p,fetcher,cooldown_seconds=3600,max_cached=600):
    """Acumula resultados válidos e protege fontes públicas contra consultas repetidas."""
    if p.get("_disable_persistent_source_cache",False):return fetcher()
    key=source_search_signature(source,p)
    now=time.time()
    with SOURCE_RESULTS_LOCK:
        entry=dict(SOURCE_RESULTS_CACHE.get("entries",{}).get(key,{}) or {})
        cached=list(entry.get("jobs",[]) or [])
        updated=float(entry.get("updated_at",0) or 0)
    cached=[job for job in cached if isinstance(job,dict) and early_date_allowed(job,p)]
    if cached and now-updated<max(60,int(cooldown_seconds)):
        with SOURCE_RESULTS_LOCK:SOURCE_RESULT_METRICS[source]={"fresh":0,"cached":len(cached),"cooldown":True}
        LOGGER.info("%s: %s vaga(s) reutilizadas do cache durante intervalo seguro",source,len(cached))
        return cached
    try:fresh=fetcher() or []
    except Exception:
        if cached:
            with SOURCE_RESULTS_LOCK:SOURCE_RESULT_METRICS[source]={"fresh":0,"cached":len(cached),"fallback":True}
            LOGGER.exception("%s falhou; usando %s vaga(s) recentes do cache",source,len(cached))
            return cached
        raise
    if not fresh:
        with SOURCE_RESULTS_LOCK:SOURCE_RESULT_METRICS[source]={"fresh":0,"cached":len(cached),"fallback":bool(cached)}
        return cached
    combined=dedupe_multisource(list(fresh)+cached,p)
    combined=[job for job in combined if early_date_allowed(job,p)][:max(1,int(max_cached))]
    with SOURCE_RESULTS_LOCK:
        SOURCE_RESULTS_CACHE.setdefault("entries",{})[key]={"updated_at":now,"source":source,"jobs":combined}
        SOURCE_RESULT_METRICS[source]={"fresh":len(fresh),"cached":max(0,len(combined)-len(fresh)),"total":len(combined)}
        snapshot={"version":1,"entries":dict(SOURCE_RESULTS_CACHE["entries"])}
    save_json_file(SOURCE_RESULTS_PATH,snapshot)
    return combined

def flush_detail_cache(force=False):
    """Persiste o cache em lotes para evitar reescrever todo o JSON a cada vaga."""
    global DETAIL_CACHE_DIRTY,DETAIL_CACHE_LAST_SAVE
    with CACHE_LOCK:
        elapsed=time.time()-DETAIL_CACHE_LAST_SAVE
        if not force and DETAIL_CACHE_DIRTY<20 and elapsed<3:return False
        if not DETAIL_CACHE_DIRTY and not force:return False
        snapshot=dict(DETAIL_CACHE)
        DETAIL_CACHE_DIRTY=0
        DETAIL_CACHE_LAST_SAVE=time.time()
    save_json_file(CACHE_PATH,snapshot)
    return True

def prune_backups(limit=7):
    backups=sorted(glob.glob(os.path.join(BACKUP_DIR,"vagas_*.db")),reverse=True)
    for old in backups[max(1,int(limit)):]:
        try:os.remove(old)
        except OSError:LOGGER.warning("Não foi possível remover o backup antigo: %s",old)

def backup_database():
    """Cria um backup SQLite consistente por dia e conserva somente os sete mais recentes."""
    if not os.path.isfile(DB_PATH) or os.path.getsize(DB_PATH)==0:return ""
    os.makedirs(BACKUP_DIR,exist_ok=True)
    today=datetime.now().strftime("%Y%m%d")
    existing=glob.glob(os.path.join(BACKUP_DIR,f"vagas_{today}_*.db"))
    if existing:return existing[-1]
    target=os.path.join(BACKUP_DIR,f"vagas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db")
    source=sqlite3.connect(f"file:{urllib.parse.quote(os.path.abspath(DB_PATH))}?mode=ro",uri=True,timeout=8)
    destination=sqlite3.connect(target)
    try:source.backup(destination)
    finally:destination.close();source.close()
    prune_backups();return target


UA="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/131 Safari/537.36"

BRAZIL_STATE_NAMES={
    "AC":"Acre","AL":"Alagoas","AP":"Amapá","AM":"Amazonas","BA":"Bahia","CE":"Ceará",
    "DF":"Distrito Federal","ES":"Espírito Santo","GO":"Goiás","MA":"Maranhão","MT":"Mato Grosso",
    "MS":"Mato Grosso do Sul","MG":"Minas Gerais","PA":"Pará","PB":"Paraíba","PR":"Paraná",
    "PE":"Pernambuco","PI":"Piauí","RJ":"Rio de Janeiro","RN":"Rio Grande do Norte",
    "RS":"Rio Grande do Sul","RO":"Rondônia","RR":"Roraima","SC":"Santa Catarina",
    "SP":"São Paulo","SE":"Sergipe","TO":"Tocantins"
}

def norm(t):
    t=unicodedata.normalize("NFKD", str(t or ""))
    t="".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"\s+"," ",t.lower()).strip()

ENGLISH_PROFESSIONAL_TERMS=[
    (r"\b(?:native|fluent) english\b","ingles fluente"),
    (r"\badvanced english\b","ingles avancado"),
    (r"\bbachelor(?:'s)? degree in law\b","bacharelado em direito"),
    (r"\b(?:bachelor(?:'s)? degree in )?business administration\b","graduacao em administracao"),
    (r"\bcomputer science\b","ciencia da computacao"),
    (r"\binformation systems\b","sistemas de informacao"),
    (r"\bsoftware engineering\b","engenharia de software"),
    (r"\bhuman resources\b","recursos humanos"),
    (r"\baccounting\b","contabilidade"),
    (r"\badministrative assistants?\b","assistente administrativo"),
    (r"\badministrative analysts?\b","analista administrativo"),
    (r"\boffice assistants?\b","assistente administrativo"),
    (r"\bcustomer (?:service|support)\b","atendimento cliente"),
    (r"\bcustomer success\b","sucesso cliente atendimento"),
    (r"\btechnical support\b","suporte tecnico"),
    (r"\bit support\b","suporte tecnologia da informacao"),
    (r"\binformation technology\b","tecnologia da informacao"),
    (r"\bsoftware development\b","desenvolvimento software"),
    (r"\bweb development\b","desenvolvimento web"),
    (r"\bdata analysis\b","analise dados"),
    (r"\bproject management\b","gestao projetos"),
    (r"\bcontract management\b","gestao contratos"),
    (r"\blegal assistants?\b","assistente juridico"),
    (r"\blaw students?\b","estudante direito"),
    (r"\blaw degree\b","graduacao direito"),
    (r"\blaw firm\b","escritorio juridico"),
    (r"\bcase management\b","gestao processos juridicos"),
    (r"\blegal documents?\b","documentos juridicos"),
    (r"\bbachelor(?:'s)? degree\b","bacharelado"),
    (r"\bundergraduate\b","graduacao cursando"),
    (r"\bcurrently studying\b","cursando"),
    (r"\bhigh school\b","ensino medio"),
    (r"\binternships?\b","estagio"),
    (r"\bintern\b","estagiario"),
    (r"\bentry[- ]level\b","junior"),
    (r"\byears? of experience\b","anos de experiencia"),
    (r"\bmicrosoft office\b","pacote office"),
    (r"\bspreadsheets?\b","planilhas"),
    (r"\bfiling\b","arquivo documentacao"),
    (r"\bdeadline management\b","controle prazos"),
    (r"\benglish\b","ingles"),
    (r"\bspanish\b","espanhol"),
]

def semantic_norm(text):
    """Normaliza conceitos profissionais PT/EN sem modificar o texto armazenado."""
    value=norm(text)
    for pattern,replacement in ENGLISH_PROFESSIONAL_TERMS:
        value=re.sub(pattern,replacement,value)
    return re.sub(r"\s+"," ",value).strip()

def detect_cv_language(text):
    value=norm(text)
    english=len(re.findall(r"\b(?:the|and|with|experience|skills|education|work|degree|professional|responsibilities)\b",value))
    portuguese=len(re.findall(r"\b(?:de|com|experiencia|habilidades|formacao|trabalho|graduacao|profissional|responsabilidades)\b",value))
    if english>=5 and english>portuguese*1.4:return "Inglês"
    if portuguese>=3 and portuguese>english*1.4:return "Português"
    return "Misto/indefinido"

ENGLISH_LEVELS=("Não informado","Básico","Intermediário","Fluente")

def normalize_english_level(value):
    level=norm(value)
    if any(token in level for token in ("fluente","fluent","avancado","advanced","nativo","native")):return "Fluente"
    if any(token in level for token in ("intermediario","intermediate")):return "Intermediário"
    if any(token in level for token in ("basico","basic","iniciante","beginner")):return "Básico"
    return "Não informado"

def detect_english_level(text,document_language=""):
    value=norm(text)
    patterns=(
        ("Fluente",r"\b(?:ingles|english)\s*[-:–—]?\s*(?:fluente|fluent|avancado|advanced|nativo|native)\b"),
        ("Intermediário",r"\b(?:ingles|english)\s*[-:–—]?\s*(?:intermediario|intermediate)\b"),
        ("Básico",r"\b(?:ingles|english)\s*[-:–—]?\s*(?:basico|basic|iniciante|beginner)\b"),
    )
    for level,pattern in patterns:
        if re.search(pattern,value):return level
    if document_language=="Inglês":return "Fluente"
    return "Não informado"

def english_level(profile):
    return normalize_english_level(profile.get("nivel_ingles",""))

def required_english_level(text):
    """Lê apenas exigências explícitas, sem inferir pelo idioma do anúncio."""
    value=semantic_norm(text)
    if re.search(r"\b(?:ingles (?:fluente|avancado|nativo)|fluent english|advanced english|native english|english fluency)\b",value):
        return "Fluente"
    if re.search(r"\b(?:ingles intermediario|intermediate (?:english|ingles))\b",value):
        return "Intermediário"
    if re.search(r"\b(?:ingles basico|basic (?:english|ingles))\b",value):
        return "Básico"
    return "Não informado"

def international_search_enabled(profile):
    """Centraliza a decisão para não misturar fontes internacionais por acidente."""
    return bool(profile.get("buscar_vagas_internacionais",False) and english_level(profile)=="Fluente")

def clean(t):
    if not t:return ""
    if BeautifulSoup:
        return BeautifulSoup(str(t),"html.parser").get_text(" ",strip=True)
    return re.sub(r"<[^>]+>"," ",str(t))

def compact_json(value):
    if value in (None,"",[],{}):return ""
    if isinstance(value,str):return value
    try:return json.dumps(value,ensure_ascii=False,separators=(",",":"))
    except Exception:return clean(value)

def structured_persistence(job,location_info):
    mode=location_info.get("mode","")
    eligible=1 if mode.startswith("Remoto Brasil") and "confirmado" in mode else 0 if mode.startswith("Remoto sem") else None
    return (
        str(job.get("workplace_type_raw") or job.get("workplace_type") or ""),
        str(job.get("workplace_source") or ""),
        compact_json(job.get("structured_location_json") or job.get("structured_location")),
        compact_json(job.get("applicant_location_requirements")),
        eligible,
        datetime.now(timezone.utc).isoformat(timespec="seconds")
    )

def default_profile():
    """Configuração neutra para a primeira execução, sem dados de desenvolvedor."""
    return {
        "perfil_curriculo_versao":0,
        "areas_curriculo_detectadas":[],"competencias_curriculo_detectadas":[],
        "termos_curriculo_detectados":[],"cursos_curriculo_detectados":[],
        "competencias_perfil":[],"termos_perfil":[],
        "formacao_curriculo_detectada":[],"perfil_inicio_carreira_detectado":False,
        "buscar_vagas_inicio_carreira":False,"buscar_jovem_aprendiz":False,"perfil_inicio_carreira":False,
        "consultas_br":[],"consultas_gupy":[],"consultas_linkedin":[],"consultas_google":[],"consultas_ingles":[],
        "cidades_presencial":[],"cidades_presencial_hibrido":[],"estado_local":"",
        "localidades_excluidas":[],
        "aceitar_remoto":True,"buscar_estagios":False,"modo_estagios":"nao_buscar",
        "areas_estagio":[],"areas_estagio_manual":False,"buscar_vagas_pcd":False,
        "consentimento_pcd_versao":"","consentimento_pcd_em":"",
        "nivel_ingles":"Não informado","nivel_ingles_manual":False,
        "buscar_vagas_internacionais":False,"preferencia_internacional_manual":False,
        "mostrar_compativeis_fora_regiao":False,
        "idade_maxima_vaga_dias":60,"idade_maxima_dias":60,
        "descartar_vagas_encerradas":True,"descartar_vagas_exclusivas_pcd":True,
        "descartar_superior_completo_obrigatorio":True,"descartar_experiencia_especifica_anos":5,
        "navegador_automacao":"automatico","enriquecimento_inicial_linkedin":60,
        "cache_detalhes_horas":24,"usar_tres_niveis_decisao":True
    }

def load_profile():
    profile=load_json_file(PROFILE_PATH,None)
    if not isinstance(profile,dict):
        profile=default_profile();save_json_file(PROFILE_PATH,profile)
    else:
        changed=False
        if "modo_estagios" not in profile:
            profile["modo_estagios"]="incluir" if profile.get("buscar_estagios",False) else "nao_buscar";changed=True
        if "areas_estagio" not in profile:
            profile["areas_estagio"]=list(profile.get("cursos_curriculo_detectados",[]));changed=True
        if "areas_estagio_manual" not in profile:
            profile["areas_estagio_manual"]=False;changed=True
        if "buscar_vagas_pcd" not in profile:
            # Compatibilidade com perfis anteriores, que armazenavam a opção invertida.
            profile["buscar_vagas_pcd"]=not bool(profile.get("descartar_vagas_exclusivas_pcd",True));changed=True
        if "buscar_vagas_internacionais" not in profile:
            # Preserva o comportamento das pessoas que já usavam a edição em inglês.
            profile["buscar_vagas_internacionais"]=profile.get("idioma_curriculo_detectado")=="Inglês";changed=True
        if "preferencia_internacional_manual" not in profile:
            profile["preferencia_internacional_manual"]=False;changed=True
        if "nivel_ingles" not in profile:
            inferred="Fluente" if profile.get("idioma_curriculo_detectado")=="Inglês" else "Não informado"
            profile["nivel_ingles"]=inferred;changed=True
        if "nivel_ingles_manual" not in profile:
            profile["nivel_ingles_manual"]=False;changed=True
        if changed:save_json_file(PROFILE_PATH,profile)
    return profile

INTERNSHIP_MODE_LABELS={
    "nao_buscar":"Não buscar estágios",
    "incluir":"Estágios junto às demais vagas",
    "somente":"Somente estágios",
}

def internship_search_mode(p):
    value=norm((p or {}).get("modo_estagios","" )).replace(" ","_")
    aliases={"nao_buscar":"nao_buscar","não_buscar":"nao_buscar","incluir":"incluir",
             "junto":"incluir","somente":"somente","somente_estagios":"somente"}
    # Aceita alterações feitas pelo campo booleano legado durante a transição.
    if value=="nao_buscar" and (p or {}).get("buscar_estagios",False):return "incluir"
    if value in aliases:return aliases[value]
    return "incluir" if (p or {}).get("buscar_estagios",False) else "nao_buscar"

def internship_selected_areas(p):
    selected=(p or {}).get("areas_estagio",[])
    if not isinstance(selected,list):selected=[]
    selected=[clean(value).strip() for value in selected if clean(value).strip()]
    return list(dict.fromkeys(selected or (p or {}).get("cursos_curriculo_detectados",[])))

def is_internship_query(value):
    return bool(re.search(r"\b(?:estagio|estagiari[oa]|intern|internship)\b",semantic_norm(value)))

def internship_area_queries(areas):
    portuguese=[];english=[]
    families=[
        ({"direito"},["direito","jurídico","contencioso","contratos","direito trabalhista"],
         ["legal intern","law internship"]),
        ({"ads","analise e desenvolvimento de sistemas","tecnologia da informacao","sistemas de informacao",
          "ciencia da computacao","engenharia de software"},
         ["TI","ADS","desenvolvimento de sistemas","desenvolvimento de software","programação",
          "suporte técnico","infraestrutura TI","dados"],
         ["IT intern","software development intern","technical support internship","data intern"]),
        ({"administracao","administracao de empresas"},
         ["administração","administrativo","financeiro","comercial","operações"],
         ["administrative intern","finance intern","operations internship"]),
        ({"ciencias contabeis","contabilidade"},["contabilidade","contábil","fiscal","auditoria","financeiro"],
         ["accounting intern","audit internship","finance intern"]),
        ({"gestao de recursos humanos","gestao de rh","recursos humanos","psicologia"},
         ["recursos humanos","RH","recrutamento e seleção","departamento pessoal"],
         ["human resources intern","recruiting intern"]),
        ({"marketing","publicidade e propaganda","publicidade","comunicacao social"},
         ["marketing","marketing digital","comunicação","conteúdo","mídias sociais"],
         ["marketing intern","communications internship","social media intern"]),
        ({"engenharia civil"},["engenharia civil","obras","projetos","orçamento de obras","planejamento de obras"],
         ["civil engineering intern"]),
        ({"enfermagem"},["enfermagem","assistência em saúde","saúde ocupacional"],["nursing intern"]),
        ({"pedagogia"},["pedagogia","educação","apoio pedagógico"],["education intern"]),
        ({"design"},["design","design gráfico","UX","UI"],["design intern","UX intern","UI intern"]),
        ({"logistica"},["logística","suprimentos","transportes","estoque"],["logistics intern","supply chain intern"]),
    ]
    for area in areas:
        normalized=semantic_norm(area)
        family=next((item for item in families if normalized in item[0]),None)
        if family:
            portuguese += [f"estágio {term}" for term in [area]+family[1]]
            english += family[2]
        else:
            portuguese += [f"estágio {area}",f"estagiário {area}",f"programa de estágio {area}"]
            english += [f"{area} intern",f"{area} internship"]
    return list(dict.fromkeys(portuguese)),list(dict.fromkeys(english))

def read_cv():
    try:
        with open(CV_PATH,encoding="utf-8") as stream:
            return stream.read()
    except:return ""

def parse_cv_file(path):
    ext=os.path.splitext(path)[1].lower()
    if ext==".pdf":
        try:
            from pypdf import PdfReader
            return "\n".join((page.extract_text() or "") for page in PdfReader(path).pages).strip()
        except Exception as e: raise RuntimeError("Não consegui ler este PDF. Tente outro arquivo.") from e
    if ext==".docx":
        try:
            from docx import Document
            return "\n".join(p.text for p in Document(path).paragraphs).strip()
        except Exception as e: raise RuntimeError("Não consegui ler este DOCX.") from e
    if ext==".txt":
        with open(path,encoding="utf-8",errors="ignore") as stream:
            return stream.read().strip()
    raise RuntimeError("Use um currículo em PDF, DOCX ou TXT.")

def cv_profile_summary(text):
    n=semantic_norm(text); areas=[]; skills=[]; education=[]
    rules=[
        ("Jurídico",["direito","juridico","jurídico","peticao","petição","processual","tribunal","pje","eproc"]),
        ("Administrativo",["administrativo","documentacao","documentação","arquivo","planilha","office"]),
        ("Atendimento",["atendimento","cliente","publico","público","suporte"]),
        ("Tecnologia / TI",["analise e desenvolvimento de sistemas","tecnologia da informacao","tecnologia da informação","software","programacao","programação","help desk","service desk"]),
    ]
    for label,terms in rules:
        if any(norm(t) in n for t in terms):areas.append(label)
    skill_rules=[("Atendimento","atendimento"),("Excel","excel"),("Pacote Office","office"),("Suporte","suporte"),
                 ("Processos jurídicos","processos juridicos"),("Controle de prazos","prazo"),("Contratos","contrat"),
                 ("Sistemas","sistema"),("Inglês","ingles"),("Espanhol","espanhol")]
    for label,term in skill_rules:
        if term in n and label not in skills:skills.append(label)
    if "cursando" in n or "graduando" in n: education.append("Ensino superior em andamento")
    if "ensino medio" in n or "ensino médio" in text.lower(): education.append("Ensino médio")
    stop={"para","com","uma","das","dos","que","por","como","the","and","from","de","da","do",
          "em","no","na","ao","aos","nas","nos","seu","sua","anos","experiencia","profissional",
          "empresa","cargo","atividades","responsavel","formacao","curso","conhecimento"}
    words=re.findall(r"[a-z][a-z0-9+#.-]{2,}",n);counts={}
    for word in words:
        if word not in stop and not word.isdigit():counts[word]=counts.get(word,0)+1
    keywords=[word for word,_count in sorted(counts.items(),key=lambda item:(-item[1],item[0]))[:24]]
    courses=[]
    # Formação deve ser extraída linha a linha. `norm(text)` remove quebras e antes
    # fazia uma linha de certificações virar, incorretamente, nome de curso.
    for raw_line in str(text or "").splitlines():
        line=semantic_norm(raw_line)
        match=re.search(
            r"(?:graduacao|bacharelado|tecnologo|curso superior|curso tecnico|cursando)\s+"
            r"(?:(?:em|de)\s+)?(.{3,70})",line)
        if not match:
            match=re.search(r"formacao academica\s*[:\-]\s*(.{3,70})",line)
        if not match:
            match=re.search(r"^(.{3,60}?)\s*[-|]\s*(?:cursando|em andamento|\d+\s*(?:periodo|semestre))\b",line)
        if not match:continue
        course=re.split(r"[.,;|•]|\s+-\s+|\s+(?:na|no|pela|pelo|universidade|faculdade)\s+",match.group(1))[0].strip(" .-")
        course=re.sub(r"^\d+[ºo]?\s*(?:periodo|semestre)\s*(?:de|em)?\s*","",course).strip()
        course=re.sub(r"\s+\d+[ºo]?\s*(?:periodo|semestre).*$","",course).strip()
        if course and not any(x in course for x in ("ensino medio","certificacoes","cursos livres")) and course not in courses:
            courses.append(course)
    if not courses:
        if re.search(r"\b(?:graduacao|bacharelado|cursando).{0,35}\bdireito\b",n):courses.append("direito")
        if "analise e desenvolvimento de sistemas" in n:courses.append("análise e desenvolvimento de sistemas")
    return {"areas":areas or ["Geral"],"skills":skills[:10],"education":education,
            "keywords":keywords,"courses":courses[:6]}

def adapt_profile_to_cv(p,text):
    prof=cv_profile_summary(text); n=norm(text); queries=[];english_queries=[]
    if not p.get("areas_estagio_manual",False):
        p["areas_estagio"]=list(prof["courses"])
    selected_internship_areas=internship_selected_areas(p)
    language=detect_cv_language(text)
    detected_english=detect_english_level(text,language)
    if not p.get("nivel_ingles_manual",False):p["nivel_ingles"]=detected_english
    # O alcance considera a fluência declarada, não o idioma em que o documento
    # foi redigido. A escolha manual continua soberana ao trocar o currículo.
    if not p.get("preferencia_internacional_manual",False):
        p["buscar_vagas_internacionais"]=(english_level(p)=="Fluente")
    normalized_cv=semantic_norm(text)
    has_high_school="Ensino médio" in prof["education"]
    declares_no_experience=bool(re.search(
        r"\b(sem experiencia|nao possuo experiencia|primeiro emprego|nenhuma experiencia|no experience)\b",
        normalized_cv))
    has_professional_history=bool(re.search(
        r"\b(experiencia profissional|professional experience|historico profissional|"
        r"atuacao profissional|trabalhei como|work history|employment history)\b",normalized_cv)) and not declares_no_experience
    entry_profile_detected=bool(has_high_school and not prof["courses"] and not has_professional_history)
    broad_entry_search=bool(p.get("buscar_vagas_inicio_carreira",False))
    apprentice_search=bool(p.get("buscar_jovem_aprendiz",False))
    entry_profile=bool(broad_entry_search or apprentice_search)
    # Um curriculo redigido predominantemente em ingles demonstra dominio funcional
    # do idioma para a triagem. Textos mistos/indefinidos nao recebem essa inferencia.
    if language=="Inglês" and "Inglês fluente" not in prof["skills"]:
        prof["skills"].append("Inglês fluente")
    internship_mode=internship_search_mode(p)
    include_internships=internship_mode!="nao_buscar"
    only_internships=internship_mode=="somente"
    if "Jurídico" in prof["areas"]:
        # A busca pública da Gupy é muito literal para cargos compostos. Termos
        # amplos descobrem os anúncios; curso, localização e compatibilidade
        # continuam sendo validados pelos filtros internos.
        # Os cargos específicos ficam primeiro porque produziram as combinações
        # mais úteis na versão de referência. Termos amplos complementam o alcance.
        queries += ["assistente jurídico","auxiliar jurídico","analista jurídico júnior","controladoria jurídica",
                    "assistente de contratos","recuperação de crédito","paralegal","legal operations",
                    "jurídico","contratos","cobrança"]
        english_queries += ["legal assistant","junior legal analyst","paralegal","legal operations",
                            "contract assistant","compliance assistant"]
        if include_internships:queries += ["estágio","estágio direito","estágio jurídico"]
        if include_internships:english_queries += ["legal intern","law internship"]
    if "Tecnologia / TI" in prof["areas"]:
        queries += ["suporte n1","help desk","service desk","assistente de suporte",
                    "analista de suporte júnior","suporte de sistemas"]
        english_queries += ["IT support","technical support","help desk","service desk",
                            "junior support analyst","systems support"]
        if include_internships:queries += ["estágio TI","estágio ADS","estágio suporte técnico"]
        if include_internships:english_queries += ["IT intern","technical support internship"]
    if "Administrativo" in prof["areas"]:
        queries += ["assistente administrativo","auxiliar administrativo","analista administrativo júnior",
                    "assistente de operações","assistente comercial","assistente de cadastro",
                    "assistente financeiro","backoffice"]
        english_queries += ["administrative assistant","office assistant","operations assistant",
                            "back office assistant","junior administrative analyst","contract administrator"]
        if include_internships:queries += ["estágio administrativo","estágio administração","estágio operações"]
        if include_internships:english_queries += ["administrative intern","operations internship"]
    if "Atendimento" in prof["areas"]:
        queries += ["atendimento ao cliente","assistente de atendimento","analista de atendimento júnior"]
        english_queries += ["customer service","customer support","customer success",
                            "client services assistant","support specialist"]
        if include_internships:queries += ["estágio atendimento","estágio suporte ao cliente"]
        if include_internships:english_queries += ["customer service intern"]
    if broad_entry_search:
        queries += ["auxiliar administrativo","atendente de loja","operador de caixa",
                    "repositor de mercadorias","estoquista","auxiliar de logística",
                    "recepcionista","auxiliar de produção","auxiliar de serviços gerais",
                    "auxiliar de cozinha","ajudante de carga e descarga","operador de telemarketing",
                    "auxiliar de escritório","auxiliar de cadastro","atendente de mercado",
                    "atendente de balcão","atendente de drogaria","atendente de lanchonete",
                    "atendente de padaria","operador de supermercado","empacotador",
                    "auxiliar de estoque","auxiliar de expedição",
                    "auxiliar de depósito","auxiliar de almoxarifado","auxiliar de armazém",
                    "ajudante de entregas","separador de pedidos",
                    "conferente de mercadorias","auxiliar de produção","ajudante de produção",
                    "ajudante geral","auxiliar de embalagem","auxiliar de montagem",
                    "auxiliar de limpeza","auxiliar de alimentação","auxiliar de loja",
                    "auxiliar de atendimento","assistente operacional","promotor de vendas",
                    "assistente administrativo","assistente de atendimento"]
    if apprentice_search:
        apprentice_queries=["jovem aprendiz","aprendiz administrativo","aprendiz de atendimento",
                            "aprendiz de logística","aprendiz de produção","aprendiz de loja",
                            "aprendiz auxiliar administrativo","aprendiz ensino médio"]
        # Opt-in explícito: consulta aprendizagem antes dos cargos gerais para
        # que os limites de cada fonte não escondam esses resultados.
        queries=apprentice_queries+queries
    for course in prof["courses"][:2]:
        if include_internships:queries.append(f"estágio {course}")
        queries += [f"assistente {course}",f"analista júnior {course}"]
    if include_internships and selected_internship_areas:
        focused_pt,focused_en=internship_area_queries(selected_internship_areas)
        queries=[query for query in queries if not is_internship_query(query)]+focused_pt
        english_queries=[query for query in english_queries if not is_internship_query(query)]+focused_en
    if only_internships and not any(is_internship_query(query) for query in queries):
        useful=[word for word in prof["keywords"] if len(word)>=5][:2]
        queries=["estágio"]+[f"estágio {word}" for word in useful]
    if not queries:
        useful=[word for word in prof["keywords"] if len(word)>=5][:4]
        queries=[f"assistente {word}" for word in useful[:2]]+[f"analista júnior {word}" for word in useful[2:]]
    if not queries:queries=["assistente","auxiliar","atendimento"]
    def valid_query(value):
        value=clean(value).strip()
        nv=norm(value)
        return value if 2<len(value)<=70 and "•" not in value and " periodo" not in nv and " semestre" not in nv else ""
    motors=p.get("motores_busca",{}) if isinstance(p.get("motores_busca",{}),dict) else {}
    legacy=[]
    if "Jurídico" in prof["areas"]:legacy+=motors.get("juridico",[])[:14]+motors.get("estagio_direito",[])[:8]
    if "Tecnologia / TI" in prof["areas"]:legacy+=motors.get("estagio_ads",[])[:10]
    if "Administrativo" in prof["areas"] or "Atendimento" in prof["areas"]:legacy+=motors.get("geral",[])[:14]
    previous=list(p.get("consultas_br",[]))+list(p.get("consultas_gupy",[]))
    if selected_internship_areas and include_internships:
        # Uma nova seleção de áreas invalida termos de estágio herdados. Vagas
        # comuns continuam preservadas apenas no modo misto.
        previous=[query for query in previous if not is_internship_query(query)]
        legacy=[query for query in legacy if not is_internship_query(query)]
        if only_internships:previous=[];legacy=[]
    primary=[valid_query(x) for x in queries+previous+legacy]
    # Somente a escolha explícita Fluente reserva consultas para cargos em inglês.
    # Nos demais níveis, toda a capacidade das fontes fica disponível ao mercado nacional.
    english=[valid_query(x) for x in english_queries] if english_level(p)=="Fluente" else []
    # Títulos em inglês também são usados por empresas brasileiras. Por isso as
    # consultas bilíngues domésticas independem da opção de fontes internacionais.
    combined=[]
    if language=="Português":
        # Em currículos em português, preserve primeiro a cobertura nacional.
        # Os equivalentes ingleses continuam presentes, mas não substituem cargos
        # brasileiros quando a fonte limita a quantidade de consultas.
        combined=primary+english
    else:
        for index in range(max(len(primary),len(english))):
            if index<len(primary) and primary[index]:combined.append(primary[index])
            if index<len(english) and english[index]:combined.append(english[index])
    if only_internships:combined=[x for x in combined if is_internship_query(x)]
    elif not include_internships:combined=[x for x in combined if not is_internship_query(x)]
    combined=list(dict.fromkeys(x for x in combined if x))
    english_norms={norm(value) for value in english if value}
    saved_pt=[valid_query(x) for x in p.get("consultas_linkedin",[])
              if valid_query(x) and norm(valid_query(x)) not in english_norms]
    if selected_internship_areas and include_internships:
        saved_pt=[query for query in saved_pt if not is_internship_query(query)]
    pt_pool=[valid_query(x) for x in queries+legacy+saved_pt]
    linkedin=[]
    if language=="Português":
        linkedin=pt_pool+english
    else:
        for index in range(max(len(pt_pool),len(english))):
            if index<len(pt_pool) and pt_pool[index]:linkedin.append(pt_pool[index])
            if index<len(english) and english[index]:linkedin.append(english[index])
    if only_internships:linkedin=[x for x in linkedin if is_internship_query(x)]
    elif not include_internships:linkedin=[x for x in linkedin if not is_internship_query(x)]
    linkedin=list(dict.fromkeys(x for x in linkedin if x))
    p["consultas_br"]=combined[:60 if broad_entry_search else 45]
    p["consultas_gupy"]=combined[:55 if broad_entry_search else 40]
    p["consultas_linkedin"]=linkedin[:45 if broad_entry_search else 35]
    p["consultas_vagas_gupy"]=[query for query in combined if not is_internship_query(query)][:55 if broad_entry_search else 40]
    p["consultas_estagio_gupy"]=[query for query in combined if is_internship_query(query)][:40]
    p["consultas_vagas_linkedin"]=[query for query in linkedin if not is_internship_query(query)][:45 if broad_entry_search else 35]
    p["consultas_estagio_linkedin"]=[query for query in linkedin if is_internship_query(query)][:35]
    p["consultas_ingles"]=list(dict.fromkeys(english))[:20]
    english_set={norm(query) for query in english}
    search_places=[]
    if (broad_entry_search or only_internships):
        state_code=str(p.get("estado_local","") or "").strip().upper()
        search_places=[" ".join(x for x in (norm(city).title(),state_code) if x)
                       for city in p.get("cidades_presencial",[]) if norm(city)]
    google_queries=[]
    google_base_limit=24 if only_internships else (15 if broad_entry_search else 12)
    for index,query in enumerate(combined[:google_base_limit]):
        suffix="jobs Brazil" if norm(query) in english_set else "vagas Brasil"
        place=f' "{search_places[index%len(search_places)]}"' if search_places else ""
        google_queries.append(f'"{query}" {suffix}{place}')
    if only_internships and p.get("aceitar_remoto",True):
        for query in combined[:8]:
            google_queries.append(f'"{query}" vagas remoto Brasil')
    if broad_entry_search:
        generic_queries=[
            '"sem experiência" "ensino médio" vagas Brasil',
            '"não exige experiência" vagas Brasil',
            '"primeiro emprego" "ensino médio" vagas Brasil',
            '"ensino médio completo" auxiliar vagas Brasil',
            '"ensino médio completo" atendente vagas Brasil',
        ]
        for index,query in enumerate(generic_queries):
            place=f' "{search_places[index%len(search_places)]}"' if search_places else ""
            google_queries.append(query+place)
    google_cap=32 if only_internships else (20 if broad_entry_search else 12)
    p["consultas_google"]=list(dict.fromkeys(google_queries))[:google_cap]
    p["consultas_vagas_google"]=[query for query in p["consultas_google"]
                                 if not is_internship_query(query)][:20 if broad_entry_search else 12]
    stage_google=[]
    for index,query in enumerate([value for value in combined if is_internship_query(value)][:24]):
        place=f' "{search_places[index%len(search_places)]}"' if search_places else ""
        stage_google.append(f'"{query}" vagas Brasil{place}')
    if p.get("aceitar_remoto",True):
        stage_google += [f'"{query}" vagas remoto Brasil'
                         for query in [value for value in combined if is_internship_query(value)][:8]]
    p["consultas_estagio_google"]=list(dict.fromkeys(stage_google))[:32]
    p["areas_curriculo_detectadas"]=prof["areas"]
    p["competencias_curriculo_detectadas"]=prof["skills"]
    p["termos_curriculo_detectados"]=prof["keywords"]
    p["cursos_curriculo_detectados"]=prof["courses"]
    p["formacao_curriculo_detectada"]=prof["education"]
    p["perfil_inicio_carreira_detectado"]=entry_profile_detected
    p["perfil_inicio_carreira"]=entry_profile
    p["idioma_curriculo_detectado"]=language
    p["perfil_curriculo_versao"]=4
    p["competencias_perfil"]=list(dict.fromkeys(prof["skills"]+prof["keywords"]))[:40]
    p["termos_perfil"]=list(dict.fromkeys(prof["keywords"]+prof["areas"]))[:40]
    return prof

def profile_terms(p):
    return [norm(x) for x in p.get("termos_perfil",p.get("termos_relevancia",[])) if norm(x)]

def profile_skills(p):
    return [norm(x) for x in p.get("competencias_perfil",p.get("competencias",[])) if norm(x)]

def profile_courses(p):
    return [norm(x) for x in p.get("cursos_curriculo_detectados",[]) if norm(x)]

COURSE_EQUIVALENCE_GROUPS=[
    ("administracao","administracao de empresas"),
    ("ciencias contabeis","contabilidade"),
    ("gestao de recursos humanos","gestao de rh","recursos humanos"),
    ("analise e desenvolvimento de sistemas","ads","tecnologia da informacao",
     "sistemas de informacao","ciencia da computacao","engenharia de software"),
    ("publicidade e propaganda","publicidade","comunicacao social"),
]

def course_equivalents(course):
    """Retorna nomes acadêmicos equivalentes sem presumir uma formação fixa."""
    value=semantic_norm(course)
    aliases={value} if value else set()
    for group in COURSE_EQUIVALENCE_GROUPS:
        normalized={semantic_norm(item) for item in group}
        if value in normalized:aliases.update(normalized)
    return aliases

def course_matches_text(course,text):
    value=semantic_norm(text)
    aliases=course_equivalents(course)
    if any(re.search(r"(?<!\w)"+re.escape(alias)+r"(?!\w)",value) for alias in aliases if alias):
        return True
    significant=[token for token in semantic_norm(course).split() if len(token)>=4]
    # Para cursos compostos, palavras institucionais espalhadas pela descrição
    # não podem formar uma correspondência acadêmica por acidente.
    required=len(significant)
    return bool(significant and sum(token in value for token in significant)>=required)

def feedback_tokens(text):
    ignored={"vaga","estagio","assistente","auxiliar","analista","junior","pleno","senior","para","com","de","da","do","em"}
    return [x for x in re.findall(r"[a-z][a-z0-9+#.-]{3,}",norm(text)) if x not in ignored][:8]

def salary_and_benefits(job):
    text=clean((job.get("salario") or "")+" "+(job.get("descricao") or ""))
    # Prefer salary field, then explicit R$ values near salary/remuneration words.
    sal=(job.get("salario") or "").strip()
    if not sal:
        pats=re.findall(r"(?:sal[aá]rio|remunera[cç][aã]o|bolsa)[^\n.;:]{0,45}?(R\$\s*[\d\.]+(?:,\d{2})?(?:\s*(?:a|até|-)\s*R?\$?\s*[\d\.]+(?:,\d{2})?)?)",text,re.I)
        if pats: sal=pats[0].strip()
    benefits=[]
    mapping=[
      ("Vale-refeição",r"\b(?:vale[- ]?refei[cç][aã]o|vr)\b"),("Vale-alimentação",r"\b(?:vale[- ]?alimenta[cç][aã]o|va)\b"),
      ("Vale-transporte",r"\b(?:vale[- ]?transporte|vt)\b"),("Plano de saúde",r"\b(?:plano|assist[eê]ncia|conv[eê]nio)\s+(?:de\s+)?sa[uú]de\b"),
      ("Plano odontológico",r"\b(?:plano|assist[eê]ncia)\s+odontol[oó]gic[oa]\b"),("Seguro de vida",r"\bseguro de vida\b"),
      ("Auxílio home office",r"\b(?:aux[ií]lio|ajuda de custo)[^\n.;]{0,25}home office\b"),("Gympass/Wellhub",r"\b(?:gympass|wellhub)\b"),
      ("TotalPass",r"\btotalpass\b"),("PLR",r"\b(?:plr|participa[cç][aã]o nos lucros)\b"),
      ("Auxílio educação",r"\b(?:aux[ií]lio|bolsa)[^\n.;]{0,20}(?:educa[cç][aã]o|estudo|faculdade|curso)\b"),
      ("Day off",r"\bday off\b")]
    for label,pat in mapping:
        if re.search(pat,text,re.I):benefits.append(label)
    return sal or "Não informado", benefits

_HTTP_LOCAL=threading.local()

def session():
    """Reutiliza conexões HTTP dentro de cada thread de coleta."""
    s=getattr(_HTTP_LOCAL,"session",None)
    if s is None:
        s=requests.Session()
        s.headers.update({"User-Agent":UA,"Accept-Language":"pt-BR,pt;q=0.9,en;q=0.7"})
        _HTTP_LOCAL.session=s
    return s

def json_get(url,params=None,timeout=20):
    s=session(); r=s.get(url,params=params,timeout=timeout); r.raise_for_status(); return r.json()

def cached_source_json(cache_key,url,params=None,ttl=3600):
    """Evita repetir consultas de APIs publicas durante buscas consecutivas."""
    now=time.time()
    with SOURCE_API_CACHE_LOCK:
        cached=SOURCE_API_CACHE.get(cache_key)
        if cached and now-cached[0]<ttl:return cached[1]
    data=json_get(url,params)
    with SOURCE_API_CACHE_LOCK:SOURCE_API_CACHE[cache_key]=(now,data)
    return data

def cached_source_text(cache_key,url,ttl=3600):
    """Cache equivalente para feeds RSS/XML públicos."""
    now=time.time()
    with SOURCE_API_CACHE_LOCK:
        cached=SOURCE_API_CACHE.get(cache_key)
        if cached and now-cached[0]<ttl:return cached[1]
    response=session().get(url,timeout=20);response.raise_for_status();data=response.text
    with SOURCE_API_CACHE_LOCK:SOURCE_API_CACHE[cache_key]=(now,data)
    return data

def enrich_jobs_parallel(jobs,p,source_name,max_workers=4):
    """Completa detalhes em paralelo, mantendo quantidade e ordem dos resultados."""
    if not p.get("enriquecer_somente_se_necessario",True):return jobs
    targets=[(index,item) for index,item in enumerate(jobs)
             if item.get("url") and needs_enrichment(item)]
    if not targets:return jobs

    def enrich(target):
        index,item=target
        try:return index,merge_job_data(item,generic_job_from_url(item["url"],source_name),p)
        except Exception:return index,None

    with ThreadPoolExecutor(max_workers=min(max_workers,len(targets)),thread_name_prefix="detalhes") as executor:
        for index,enriched in executor.map(enrich,targets):
            if enriched is not None:jobs[index]=enriched
    return jobs

def fetch_gupy(p):
    """Public endpoint used by Gupy's employability portal."""
    out=[]; seen=set()
    base="https://employability-portal.gupy.io/api/v1/jobs"
    state_code=str(p.get("estado_local") or "").strip().upper()
    state_name=BRAZIL_STATE_NAMES.get(state_code,state_code)
    # A consulta estadual prioriza vagas locais, mas não substitui a descoberta
    # nacional: cadastros sem UF estruturada e modalidades inconsistentes só
    # aparecem no escopo amplo. A filtragem geográfica acontece depois da coleta.
    scopes=[]
    if state_name:scopes.append(({"state":state_name},(0,100)))
    scopes.append(({},(0,100)))
    accepts_remote=bool(p.get("aceitar_remoto",p.get("aceita_remoto",True)))
    if accepts_remote:scopes.append(({"workplaceType":"remote"},(0,100)))
    queries=list(dict.fromkeys(p.get("consultas_gupy",p.get("consultas_br",[]))))
    max_results=max(1,int(p.get("max_resultados_por_fonte",500)))
    # Com muitos cargos, um termo amplo (por exemplo "jurídico") não pode consumir
    # sozinho todo o limite da fonte. Divide a capacidade entre os cargos e os
    # escopos para preservar variedade de áreas, cidades e modalidades.
    balanced=len(queries)>3
    query_budget=max(6,(max_results+len(queries)-1)//max(1,len(queries))) if balanced else max_results
    for query_index,q in enumerate(queries):
        query_added=0
        active_scopes=len(scopes)
        scope_budget=max(1,(query_budget+active_scopes-1)//active_scopes)
        for scope_index,(scope,offsets) in enumerate(scopes):
            scope_added=0
            if balanced:offsets=(0,)
            for offset in offsets:
                try:
                    params={"jobName":q,"limit":100,"offset":offset};params.update(scope)
                    d=json_get(base,params)
                except Exception:
                    break
                items=d.get("data") or d.get("results") or d.get("jobs") or []
                if isinstance(items,dict): items=items.get("data",[]) or items.get("results",[])
                if not items: break
                selected_items=items
                if balanced and len(items)>scope_budget:
                    # Mantém os primeiros resultados e alterna diariamente uma
                    # segunda janela, acumulada pelo cache persistente.
                    head_count=max(1,(scope_budget+1)//2)
                    tail_count=max(0,scope_budget-head_count)
                    selected_items=list(items[:head_count])
                    if tail_count:
                        available=max(1,len(items)-head_count-tail_count+1)
                        rotation=(date.today().toordinal()+query_index*17+scope_index*31)%available
                        start=head_count+rotation
                        selected_items.extend(items[start:start+tail_count])
                for j in selected_items:
                    url=j.get("jobUrl") or j.get("url") or j.get("applyUrl") or ""
                    jid=j.get("id") or j.get("jobId")
                    if not url and jid:
                        url=f"https://portal.gupy.io/job-search"
                    key=url or f"gupy:{jid}:{j.get('name')}"
                    if key in seen:continue
                    seen.add(key)
                    wt_raw=j.get("workplaceType") or j.get("workplaceTypeLabel") or ""
                    wt=norm(wt_raw)
                    city=j.get("city") or ((j.get("address") or {}).get("city") if isinstance(j.get("address"),dict) else "")
                    state=j.get("state") or ((j.get("address") or {}).get("state") if isinstance(j.get("address"),dict) else "")
                    loc=", ".join(x for x in [city,state] if x) or j.get("location") or "Brasil"
                    item={
                        "titulo":j.get("name") or j.get("title") or "",
                        "empresa":j.get("careerPageName") or j.get("companyName") or j.get("company") or "Não informado",
                        "local":loc,
                        "descricao":clean(j.get("description") or j.get("descriptionHtml") or j.get("requirements") or ""),
                        "url":url,"fonte":"Gupy",
                        "data_publicacao":str(j.get("publishedDate") or j.get("publicationDate") or "")[:10],
                        "salario":str(j.get("salary") or ""),
                        "remote":("remote" in wt or "remot" in wt),
                        "workplace_type":wt,
                        "workplace_type_raw":str(wt_raw),
                        "workplace_source":"structured" if wt else "",
                        "structured_location":j.get("address") or {"city":city,"state":state,"display":loc},
                        "applicant_location_requirements":j.get("applicantLocationRequirements") or "",
                        "source_brazil":True
                    }
                    if not early_date_allowed(item,p):
                        continue
                    out.append(item)
                    query_added+=1;scope_added+=1
                    if len(out)>=max_results:
                        return enrich_jobs_parallel(out,p,"Gupy")
                    if balanced and (query_added>=query_budget or scope_added>=scope_budget):break
                if balanced and (query_added>=query_budget or scope_added>=scope_budget):break
                if len(items)<100:break
            if balanced and query_added>=query_budget:break
    return enrich_jobs_parallel(out,p,"Gupy")

def linkedin_search_html(q,remote=False,location="Brazil",start=0):
    params={"keywords":q,"location":location,"start":start}
    if remote: params["f_WT"]="2"
    url="https://www.linkedin.com/jobs-guest/jobs/api/seeMoreJobPostings/search?"+urllib.parse.urlencode(params)
    s=session(); r=s.get(url,timeout=20); r.raise_for_status()
    return r.text

def fetch_linkedin(p):
    out=[];seen=set()
    terms=p.get("consultas_linkedin") or ["assistente jurídico","estágio jurídico",
           "assistente administrativo","atendimento","suporte técnico"]
    cities=p.get("cidades_presencial") or p.get("cidades_presencial_hibrido") or []
    state=(p.get("estado_local") or "").strip()
    state_label=BRAZIL_STATE_NAMES.get(state.upper(),state)
    # A consulta externa recebe uma chave canônica. Assim, "São Paulo" e
    # "sao paulo" produzem exatamente a mesma busca; a grafia do usuário é preservada no perfil.
    state_search=norm(state_label).title()
    # Empregos comuns preservam a estratégia econômica da Beta (cidade principal).
    # Em modo somente estágios, porém, cada cidade configurada precisa ser
    # consultada: a busca estadual do LinkedIn frequentemente devolve resultados
    # de outros estados e omite oportunidades locais. O filtro geográfico interno
    # continua sendo a autoridade final.
    main_city=norm(cities[0]).title() if cities and norm(cities[0]) else ""
    local_query=", ".join(x for x in [main_city,state_search,"Brazil"] if x) or "Brazil"
    internship_only=internship_search_mode(p)=="somente"
    accepts_remote=bool(p.get("aceitar_remoto",p.get("aceita_remoto",True)))
    internship_locations=[]
    if internship_only:
        for city in cities:
            city_search=norm(city).title()
            if city_search:
                internship_locations.append(", ".join(
                    x for x in (city_search,state_search,"Brazil") if x))
        if state_search:internship_locations.append(", ".join((state_search,"Brazil")))
        internship_locations=list(dict.fromkeys(internship_locations)) or [local_query]
    for q in terms:
        if internship_only:
            # Locais primeiro: vagas remotas não ocupam as primeiras posições
            # que serão enriquecidas com descrição e modalidade.
            searches=[(False,loc,(0,)) for loc in internship_locations]
            if accepts_remote:searches.append((True,"Brazil",(0,25)))
        else:
            searches=[(True,"Brazil",(0,25)),(False,local_query,(0,25))]
        for remote,loc,starts in searches:
            for start in starts:
                try: html=linkedin_search_html(q,remote,loc,start)
                except Exception: continue
                soup=BeautifulSoup(html,"html.parser")
                cards=soup.select("li")
                if not cards:break
                for card in cards:
                    a=card.select_one("a.base-card__full-link") or card.find("a",href=True)
                    title=card.select_one(".base-search-card__title")
                    company=card.select_one(".base-search-card__subtitle")
                    location=card.select_one(".job-search-card__location")
                    if not a or not title: continue
                    url=(a.get("href") or "").split("?")[0]
                    if not url or url in seen:continue
                    seen.add(url)
                    out.append({
                        "titulo":title.get_text(" ",strip=True),
                        "empresa":company.get_text(" ",strip=True) if company else "Não informado",
                        "local":location.get_text(" ",strip=True) if location else loc,
                        "descricao":"","url":url,"fonte":"LinkedIn",
                        "data_publicacao":"","salario":"","remote":False,"workplace_type":"",
                        "search_remote_hint":bool(remote),"source_brazil":True
                    })
                    if len(out)>=500:return out
                time.sleep(.25)
    # O cartão público não confirma a modalidade. Enriquece antes do filtro
    # geográfico para não perder vagas realmente remotas por falta de detalhes.
    initial_limit=max(0,int(p.get("enriquecimento_inicial_linkedin",60)))
    if initial_limit<=0:initial_limit=60  # compatibilidade com perfis anteriores
    enriched=enrich_jobs_parallel(out[:initial_limit],p,"LinkedIn",max_workers=6)
    out[:len(enriched)]=enriched
    return out



def source_rank(name,p=None):
    name=norm(name)
    order=["gupy","linkedin","ciee/google","nube/google","iel/google","eureca/google",
           "companhia de estagios/google","indeed/google","google","remotive","jobicy","himalayas","remote landers",
           "remote game jobs","work with indies","hitmarker/google","vagas em games/google"]
    if p:
        order=[norm(x) for x in p.get("fonte_prioridade",order)]
    for i,x in enumerate(order):
        if x and x in name:return i
    return 99

def field_missing(v):
    return norm(v) in ("","nao informado","não informado","brasil","brazil")

def needs_enrichment(job):
    return (
        field_missing(job.get("empresa")) or
        field_missing(job.get("local")) or
        len((job.get("descricao") or "").strip()) < 120 or
        not job.get("workplace_type")
    )

def merge_job_data(a,b,p):
    ra,rb=source_rank(a.get("fonte",""),p),source_rank(b.get("fonte",""),p)
    primary,other=(a,b) if ra<=rb else (b,a)
    out=dict(primary)

    for k in ("titulo","empresa","local","descricao","data_publicacao","valid_through","salario","workplace_type",
              "workplace_type_raw","workplace_source","structured_location","structured_location_json",
              "applicant_location_requirements"):
        pv=out.get(k)
        ov=other.get(k)
        if (field_missing(pv) or (k=="descricao" and len(str(pv or ""))<len(str(ov or "")))) and ov:
            out[k]=ov

    out["remote"]=bool(primary.get("remote") or other.get("remote"))
    out["source_brazil"]=bool(primary.get("source_brazil") or other.get("source_brazil"))

    fontes=[]
    for obj in (a,b):
        raw=obj.get("fontes_encontradas") or obj.get("fonte") or ""
        for f in str(raw).split(" + "):
            if f and f not in fontes:fontes.append(f)
    out["fontes_encontradas"]=" + ".join(fontes)
    out["fonte"]=primary.get("fonte") or other.get("fonte") or ""
    out["url"]=primary.get("url") or other.get("url") or ""
    return out

def dedupe_multisource(jobs,p):
    by_fp={};ordered=[]
    for j in jobs:
        fp=job_fingerprint(j)
        if not fp:
            j["fontes_encontradas"]=j.get("fonte","")
            ordered.append(j);continue
        if fp in by_fp:
            merged=merge_job_data(by_fp[fp],j,p)
            by_fp[fp].clear();by_fp[fp].update(merged)
        else:
            j["fontes_encontradas"]=j.get("fonte","")
            by_fp[fp]=j;ordered.append(j)
    return ordered

def early_date_allowed(job,p):
    ok,_=vacancy_date_ok(job,p)
    return ok

def generic_job_from_url(url,fonte="Google"):
    global DETAIL_CACHE_DIRTY
    ttl=int(load_profile().get("cache_detalhes_horas",24))*3600
    cached=DETAIL_CACHE.get(url)
    if isinstance(cached,dict) and time.time()-cached.get("ts",0)<ttl:
        j=dict(cached.get("job") or {})
        if j:
            j["fonte"]=fonte
            return j

    s=session();original_url=url
    request_url=url
    if norm(fonte)=="linkedin":
        match=re.search(r"(?:currentJobId=|/jobs/view/(?:[^/?]*-)?)(\d{5,})",url)
        if match:
            request_url=f"https://www.linkedin.com/jobs-guest/jobs/api/jobPosting/{match.group(1)}"
    try:
        r=s.get(request_url,timeout=15,allow_redirects=True);r.raise_for_status()
    except Exception:
        if request_url==url:raise
        r=s.get(url,timeout=15,allow_redirects=True);r.raise_for_status()
    soup=BeautifulSoup(r.text,"html.parser")
    job={}

    # Modalidade explícita exibida no topo de páginas do LinkedIn.
    # Procuramos somente elementos cujo texto seja exatamente o selo da modalidade,
    # evitando confundir benefícios como "auxílio home office" com trabalho remoto.
    explicit_page_workplace=""
    if norm(fonte)=="linkedin":
        try:
            exact=[]
            for el in soup.find_all(["span","li","div"]):
                txt=clean(el.get_text(" ",strip=True))
                nt=norm(txt)
                if nt in ("presencial","on-site","onsite"):
                    exact.append("onsite")
                elif nt in ("hibrido","híbrido","hybrid"):
                    exact.append("hybrid")
                elif nt in ("remoto","remota","remote"):
                    exact.append("remote")
            # Presencial/híbrido têm prioridade se a página contiver termos conflitantes.
            if "onsite" in exact: explicit_page_workplace="onsite"
            elif "hybrid" in exact: explicit_page_workplace="hybrid"
            elif "remote" in exact: explicit_page_workplace="remote"
        except Exception:
            pass

    # 1) JSON-LD JobPosting é a fonte preferida.
    for tag in soup.find_all("script",type="application/ld+json"):
        try:
            d=json.loads(tag.string or "{}")
            candidates=[]
            if isinstance(d,list):candidates=d
            elif isinstance(d,dict):
                candidates=[d]
                if isinstance(d.get("@graph"),list):candidates+=d["@graph"]
            for x in candidates:
                if isinstance(x,dict) and x.get("@type")=="JobPosting":
                    job=x;break
            if job:break
        except Exception:pass

    title=job.get("title") if job else ""
    desc=clean(job.get("description","")) if job else ""
    company=""
    if job and isinstance(job.get("hiringOrganization"),dict):
        company=job["hiringOrganization"].get("name","")

    loc=""
    remote=False
    workplace_type=""
    if job:
        jlt_raw=job.get("jobLocationType") or ""
        jlt=norm(jlt_raw)
        remote=("telecommute" in jlt or "remote" in jlt)
        if remote:workplace_type="remote"
        elif "hybrid" in jlt or "hibrid" in jlt:workplace_type="hybrid"
        elif "onsite" in jlt or "on-site" in jlt or "presencial" in jlt:workplace_type="onsite"

        jl=job.get("jobLocation")
        if isinstance(jl,list):jl=jl[0] if jl else None
        if isinstance(jl,dict):
            addr=jl.get("address") or {}
            if isinstance(addr,dict):
                loc=", ".join(str(addr.get(k,"")) for k in
                    ("addressLocality","addressRegion","addressCountry") if addr.get(k))
        req=job.get("applicantLocationRequirements")
        if req and not loc:loc=clean(req)

    # No LinkedIn, o selo visível de modalidade é mais específico que a pista da busca.
    if explicit_page_workplace:
        workplace_type=explicit_page_workplace
        remote=(explicit_page_workplace=="remote")

    # 2) Fallbacks de metadados/HTML.
    if not title:
        og=soup.find("meta",property="og:title")
        title=(og.get("content","") if og else "") or (soup.title.get_text(" ",strip=True) if soup.title else "")
    if not desc:
        main=soup.find("main") or soup.body
        desc=main.get_text(" ",strip=True)[:12000] if main else ""
    if not company:
        company_meta=soup.find("meta",{"name":"author"})
        company=company_meta.get("content","") if company_meta else "Não informado"

    salary=""
    if job and job.get("baseSalary"):
        salary=clean(json.dumps(job.get("baseSalary"),ensure_ascii=False))

    workplace_source="page_badge" if explicit_page_workplace else "structured" if workplace_type else ""
    result={
        "titulo":title[:220],"empresa":company[:160],"local":loc or "Não informado",
        "descricao":desc,"url":original_url,"fonte":fonte,
        "data_publicacao":str(job.get("datePosted","") if job else "")[:10],
        "valid_through":str(job.get("validThrough","") if job else "")[:25],
        "salario":salary,"remote":remote,"workplace_type":workplace_type,"workplace_source":workplace_source,
        "workplace_type_raw":str(job.get("jobLocationType") or explicit_page_workplace or "") if job else explicit_page_workplace,
        "structured_location_json":compact_json(job.get("jobLocation") if job else ""),
        "applicant_location_requirements":compact_json(job.get("applicantLocationRequirements") if job else ""),
        "source_brazil":False
    }

    # Salva tanto pela URL original quanto pela URL final.
    payload={"ts":time.time(),"job":result}
    with CACHE_LOCK:
        DETAIL_CACHE[url]=payload
        DETAIL_CACHE[r.url]=payload
        DETAIL_CACHE_DIRTY+=1
        # Evita cache crescer indefinidamente.
        if len(DETAIL_CACHE)>1500:
            oldest=sorted(DETAIL_CACHE.items(),key=lambda kv:kv[1].get("ts",0))[:300]
            for k,_ in oldest:DETAIL_CACHE.pop(k,None)
    flush_detail_cache()
    return result

def google_urls(query,n=10):
    # Low-volume Google discovery. Google can occasionally throttle automated queries.
    from googlesearch import search
    try:
        return list(search(query,num_results=n,lang="pt",sleep_interval=1.5))
    except TypeError:
        return list(search(query,num=n,stop=n,pause=1.5))

def fetch_google(p,site=None,source_name="Google",max_queries=12,max_results=7,max_jobs=60):
    out=[];seen=set()
    cv_queries=p.get("consultas_gupy") or p.get("consultas_br") or []
    queries=p.get("consultas_google") or [f'"{q}" vagas Brasil' for q in cv_queries[:6]]
    if not queries:return out
    if site:
        queries=[f"site:{site} "+q for q in queries]
    queries=queries[:max_queries]
    # Google é complementar: menos URLs, mais qualidade, sem repetir dezenas de sinônimos.
    for q in queries:
        try:urls=google_urls(q,max_results)
        except Exception:continue
        pending=[]
        for u in urls:
            base=u.split("?")[0]
            if base in seen:continue
            seen.add(base)
            pending.append(u)

        def fetch_detail(u):
            try:return generic_job_from_url(u,source_name)
            except Exception:return None

        with ThreadPoolExecutor(max_workers=min(5,len(pending)) or 1,thread_name_prefix="google-detalhes") as executor:
            details=executor.map(fetch_detail,pending)
        for j in details:
            if not j:continue
            if not early_date_allowed(j,p):continue
            out.append(j)
            if len(out)>=max_jobs:return out
    return out

def fetch_internship_portals(p):
    """Descobre vagas em portais oficiais sem acessar áreas autenticadas."""
    portals=[
        ("portal.ciee.org.br/quero-uma-vaga","CIEE/Google"),
        ("nube.com.br/estudantes","Nube/Google"),
        ("ielcarreiras.com.br/oportunidades","IEL/Google"),
        ("eureca.me","Eureca/Google"),
        ("ciadeestagios.com.br/vagas","Companhia de Estágios/Google"),
    ]
    jobs=[]
    for domain,source in portals:
        try:jobs += fetch_google(p,domain,source,max_queries=6,max_results=5,max_jobs=30)
        except Exception:LOGGER.exception("Falha ao consultar portal público de estágio: %s",domain)
    return dedupe_multisource(jobs,p)

def fetch_remotive(p):
    out=[]
    try:
        d=json_get("https://remotive.com/api/remote-jobs",{"search":"Brazil"})
        for j in d.get("jobs",[]):
            out.append({"titulo":j.get("title",""),"empresa":j.get("company_name",""),
             "local":j.get("candidate_required_location",""),"descricao":clean(j.get("description","")),
             "url":j.get("url",""),"fonte":"Remotive","data_publicacao":str(j.get("publication_date",""))[:10],
             "salario":j.get("salary","") or "","remote":True,"workplace_type":"remote",
             "workplace_type_raw":"remote","workplace_source":"source_api",
             "structured_location":{"candidate_required_location":j.get("candidate_required_location","")},
             "applicant_location_requirements":j.get("candidate_required_location",""),"source_brazil":False})
    except Exception:
        LOGGER.exception("Falha ao buscar vagas na Remotive")
    return out

def fetch_jobicy(p):
    """Vagas remotas internacionais da API publica Jobicy."""
    out=[];seen=set()
    try:
        tags=[clean(tag).strip() for tag in p.get("consultas_ingles",[]) if 3<=len(clean(tag).strip())<=50][:6]
        requests_to_make=[({"count":50,"tag":tag},tag) for tag in tags] or [({"count":200},"latest")]
        items=[]
        def query(request_data):
            params,key=request_data
            return cached_source_json("jobicy-"+norm(key),"https://jobicy.com/api/v2/remote-jobs",params)
        with ThreadPoolExecutor(max_workers=min(4,len(requests_to_make))) as executor:
            responses=executor.map(query,requests_to_make)
        for data in responses:
            items.extend(data.get("jobs",[]))
        for item in items:
            canonical=item.get("url","") or ""
            if not canonical or canonical in seen:continue
            seen.add(canonical)
            geo=clean(item.get("jobGeo") or "")
            salary_parts=[item.get("salaryMin"),item.get("salaryMax")]
            salary=" - ".join(str(x) for x in salary_parts if x not in (None,""))
            if salary:
                salary=" ".join(x for x in [salary,item.get("salaryCurrency"),item.get("salaryPeriod")] if x)
            job={"titulo":item.get("jobTitle","") or "","empresa":item.get("companyName","") or "",
                 "local":geo or "Não informado","descricao":clean(item.get("jobDescription") or item.get("jobExcerpt") or ""),
                 "url":canonical,"fonte":"Jobicy","data_publicacao":str(item.get("pubDate","") or "")[:10],
                 "salario":salary,"remote":True,"workplace_type":"remote","workplace_type_raw":"remote",
                 "workplace_source":"source_api","structured_location":{"candidate_required_location":geo},
                 "applicant_location_requirements":geo,"source_brazil":False}
            if early_date_allowed(job,p):out.append(job)
    except Exception:
        LOGGER.exception("Falha ao buscar vagas na Jobicy")
    return out

def epoch_date(value):
    try:
        number=float(value)
        if number>100000000000:number/=1000
        return datetime.fromtimestamp(number,timezone.utc).date().isoformat()
    except Exception:return str(value or "")[:10]

def fetch_himalayas(p):
    """Pesquisa direcionada na API pública Himalayas, incluindo vagas Worldwide."""
    out=[];seen=set()
    try:
        queries=[clean(q).strip() for q in p.get("consultas_ingles",[]) if clean(q).strip()][:6]
        if not queries:queries=["customer support","administrative assistant","IT support"]
        def query(term):
            params={"q":term,"country":"BR","sort":"recent"}
            return cached_source_json("himalayas-br-"+norm(term),
                                      "https://himalayas.app/jobs/api/search",params,ttl=86400)
        with ThreadPoolExecutor(max_workers=min(4,len(queries))) as executor:
            responses=executor.map(query,queries)
        for data in responses:
            for item in data.get("jobs",[]):
                url=item.get("applicationLink") or item.get("guid") or ""
                if not url or url in seen:continue
                seen.add(url)
                restrictions=item.get("locationRestrictions") or []
                names=[]
                for restriction in restrictions:
                    name=restriction.get("name") if isinstance(restriction,dict) else restriction
                    if name:names.append(str(name))
                location=", ".join(names) if names else "Worldwide"
                salary_values=[item.get("minSalary"),item.get("maxSalary")]
                salary=" - ".join(str(x) for x in salary_values if x not in (None,""))
                if salary:salary=" ".join(x for x in [salary,item.get("currency"),item.get("salaryPeriod")] if x)
                job={"titulo":item.get("title","") or "","empresa":item.get("companyName","") or "",
                     "local":location,"descricao":clean(item.get("description") or item.get("excerpt") or ""),
                     "url":url,"fonte":"Himalayas","data_publicacao":epoch_date(item.get("pubDate")),
                     "valid_through":epoch_date(item.get("expiryDate")),"salario":salary,"remote":True,
                     "workplace_type":"remote","workplace_type_raw":"remote","workplace_source":"source_api",
                     "structured_location":{"candidate_required_location":names},
                     "applicant_location_requirements":names or "Worldwide","source_brazil":False}
                if early_date_allowed(job,p):out.append(job)
    except Exception:
        LOGGER.exception("Falha ao buscar vagas na Himalayas")
    return out

def fetch_remote_landers(p):
    """Lê vagas remotas com links de origem ATS na API pública Remote Landers."""
    out=[]
    try:
        data=cached_source_json("remote-landers-latest","https://remotelanders.com/api/jobs",
                                {"limit":100,"page":1},ttl=600)
        for item in data.get("jobs",[]):
            location=clean(item.get("location") or "") or "Não informado"
            metadata=" | ".join(str(x) for x in [item.get("category"),item.get("level"),item.get("type")]
                                if x)
            subtags=", ".join(map(str,item.get("subtags") or []))
            description=" | ".join(x for x in [metadata,subtags] if x)
            job={"titulo":item.get("title","") or "","empresa":item.get("company","") or "",
                 "local":location,"descricao":description,"url":item.get("url") or item.get("applyUrl") or "",
                 "fonte":"Remote Landers","data_publicacao":str(item.get("postedDate","") or "")[:10],
                 "salario":item.get("salary","") or "","remote":True,"workplace_type":"remote",
                 "workplace_type_raw":"remote","workplace_source":"source_api",
                 "structured_location":{"candidate_required_location":location},
                 "applicant_location_requirements":location,"source_brazil":False}
            if job["url"] and early_date_allowed(job,p):out.append(job)
    except Exception:
        LOGGER.exception("Falha ao buscar vagas na Remote Landers")
    return out

def fetch_games_rss(p,url,source_name):
    """Converte feeds públicos especializados em games para o formato interno."""
    out=[]
    try:
        root=ElementTree.fromstring(cached_source_text("rss-"+norm(source_name),url))
        for item in root.findall(".//item"):
            def raw_value(tag):
                node=item.find(tag);return str(node.text or "").strip() if node is not None else ""
            title=clean(raw_value("title"));link=raw_value("link") or raw_value("guid")
            description=clean(raw_value("description"))
            if not title or not link:continue
            company=title.split(" is hiring ",1)[0].strip() if " is hiring " in title else "Não informado"
            role=title.split(" is hiring ",1)[1].strip() if " is hiring " in title else title
            location="Remote"
            match=re.search(r"\bto work from\s+(.+?)(?:\s*[|—-]|$)",role,re.I)
            if match:
                location=match.group(1).strip();role=role[:match.start()].strip()
            role=re.sub(r"\s*\(Remote Job\)\s*$","",role,flags=re.I).strip()
            role=re.sub(r"^(?:a|an)\s+","",role,flags=re.I).strip()
            published=raw_value("pubDate")
            try:published=parsedate_to_datetime(published).date().isoformat()
            except Exception:pass
            job={"titulo":role,"empresa":company,"local":location,"descricao":description,
                 "url":link,"fonte":source_name,"data_publicacao":published,"salario":"",
                 "remote":True,"workplace_type":"remote","workplace_type_raw":"remote",
                 "workplace_source":"source_feed","structured_location":{"candidate_required_location":location},
                 "applicant_location_requirements":location,"source_brazil":False}
            if early_date_allowed(job,p):out.append(job)
    except Exception:
        LOGGER.exception("Falha ao buscar vagas em %s",source_name)
    return out

def fetch_remote_game_jobs(p):
    return fetch_games_rss(p,"https://remotegamejobs.com/feed.rss","Remote Game Jobs")

def fetch_work_with_indies(p):
    return fetch_games_rss(p,"https://www.workwithindies.com/careers/rss.xml","Work With Indies")

def job_fingerprint(job):
    title=norm(job.get("titulo",""))
    company=norm(job.get("empresa",""))
    loc=norm(job.get("local",""))
    # Remove apenas ruídos de apresentação. Senioridade é preservada para não
    # fundir, por exemplo, uma vaga Júnior com outra Pleno da mesma empresa.
    title=re.sub(r"\b(remoto|remote|hibrido|presencial|home office)\b","",title)
    title=re.sub(r"\s*[-–—|/]\s*$","",title)
    title=re.sub(r"\s+"," ",title).strip()
    company=re.sub(r"\b(s\.?a\.?|ltda\.?|eireli|inc\.?|llc|corp\.?|corporation)\b\.?$","",company).strip(" -.,")
    if company and company not in ("nao informado","confidencial","confidential"):
        return f"{title}|{company}"
    if loc and loc!="nao informado":
        return f"{title}|{loc}"
    return ""

def is_intern(job):
    x=norm(job.get("titulo","")+" "+job.get("descricao",""))
    return bool(re.search(r"\b(estagio|estagiari[oa]s?|interns?|internship)\b",x))

def is_apprentice(job):
    title=norm(job.get("titulo","") or "")
    return bool(re.search(r"\b(jovem aprendiz|menor aprendiz|aprendiz)\b",title))

def legal_evidence(job,p):
    title=semantic_norm(job.get("titulo",""))
    desc=semantic_norm(job.get("descricao",""))
    text=title+" "+desc
    titles=[norm(x) for x in p.get("cargos_juridicos_amplos",[])]

    if any(x and x in title for x in titles):
        return True

    legal_terms=[
        "juridico","juridica","legal operations","legal ops","paralegal",
        "controladoria juridica","contencioso","processual","processo judicial",
        "processos judiciais","prazo processual","prazos processuais","publicacoes",
        "intimacoes","peticoes","recursos","diligencias","protocolo judicial",
        "acompanhamento processual","andamento processual","carteira processual",
        "pje","eproc","esaj","e-saj","projudi","tribunal","foro","comarca",
        "contratos juridicos","departamento juridico","escritorio de advocacia",
        "advocacia","compliance juridico","regulatorio juridico"
    ]
    hits=sum(1 for x in legal_terms if x in text)
    return hits>=2 or any(x in title for x in [
        "jurid","paralegal","legal ops","legal operations","contencioso",
        "processual","controladoria","controller jurid"
    ])


def job_category(job,p):
    if is_intern(job):
        cs=internship_course_status(job,p)
        if cs=="OK_PERFIL":return "Estágio compatível"
        if cs=="REVISAR":return "Estágio — verificar curso"
        return "Estágio fora da área"
    if legal_evidence(job,p):return "Jurídico"
    text=semantic_norm((job.get("titulo","") or "")+" "+(job.get("descricao","") or ""))
    if any(x in text for x in ["suporte n1","suporte tecnico","help desk","service desk",
                               "analista de suporte","implantacao","tecnologia da informacao","software"]):
        return "Geral — Tecnologia/Suporte"
    return "Geral"

def location_decision(job,p):
    loc=norm(job.get("local",""))
    title=norm(job.get("titulo",""))
    desc=norm(job.get("descricao",""))
    workplace=norm(job.get("workplace_type",""))

    allowed=[norm(x) for x in p.get("cidades_presencial",p.get("cidades_presencial_hibrido",[]))]
    default_uf=norm(p.get("estado_local","")).upper()
    uf_names={"AC":"acre","AL":"alagoas","AP":"amapa","AM":"amazonas","BA":"bahia","CE":"ceara",
              "DF":"distrito federal","ES":"espirito santo","GO":"goias","MA":"maranhao","MT":"mato grosso",
              "MS":"mato grosso do sul","MG":"minas gerais","PA":"para","PB":"paraiba","PR":"parana",
              "PE":"pernambuco","PI":"piaui","RJ":"rio de janeiro","RN":"rio grande do norte",
              "RS":"rio grande do sul","RO":"rondonia","RR":"roraima","SC":"santa catarina",
              "SP":"sao paulo","SE":"sergipe","TO":"tocantins"}
    location_parts=[part.strip() for part in re.split(r"\s*[,/|]\s*|\s+-\s+",loc) if part.strip()]
    published_city=location_parts[0] if location_parts else ""
    published_region=" ".join(location_parts[1:])
    def allowed_city(entry):
        bits=[part.strip() for part in entry.rsplit("/",1)]
        city=bits[0];uf=(bits[1].upper() if len(bits)==2 else default_uf)
        if published_city!=city:return False
        if not uf:return True
        return bool(re.search(rf"\b{re.escape(uf.lower())}\b",published_region) or
                    re.search(rf"\b{re.escape(uf_names.get(uf,''))}\b",published_region))
    local_allowed=any(city and allowed_city(city) for city in allowed)

    generic_locations={
        "","brasil","brazil","remoto","remota","remote","home office","anywhere",
        "latam","latin america","america latina","worldwide","nao informado","não informado"
    }
    loc_generic=loc in generic_locations

    uf_names=[
        "acre","alagoas","amapa","amazonas","bahia","ceara","distrito federal","espirito santo",
        "goias","maranhao","mato grosso","mato grosso do sul","minas gerais","para","paraiba",
        "parana","pernambuco","piaui","rio de janeiro","rio grande do norte","rio grande do sul",
        "rondonia","roraima","santa catarina","sao paulo","sergipe","tocantins"
    ]
    uf_abbr=r"(?:ac|al|ap|am|ba|ce|df|es|go|ma|mt|ms|mg|pa|pb|pr|pe|pi|rj|rn|rs|ro|rr|sc|sp|se|to)"
    explicit_br_location=(
        bool(re.search(rf"\b[a-zà-ÿ][a-zà-ÿ .'-]+[,/ -]\s*{uf_abbr}\b",loc))
        or any(re.search(rf"\b{re.escape(state)}\b",loc) for state in uf_names)
    )

    scope_text=loc+" "+desc[:4500]
    remote_scope_allowed=(
        explicit_br_location or local_allowed or norm(job.get("fonte",""))=="gupy" or
        bool(re.search(r"\b(brasil|brazil|latam|latin america|america latina|worldwide|anywhere)\b",loc)) or
        bool(re.search(
            r"\b(candidat[oa]s? (?:do|de|no|em) brasil|candidates? (?:in|from) brazil|"
            r"work from brazil|trabalh(?:ar|e) (?:do|no|em) brasil|"
            r"(?:open to|available (?:in|for)) (?:brazil|latam|latin america)|"
            r"latam|latin america|america latina|worldwide|work from anywhere|trabalhe de qualquer lugar)\b",
            scope_text))
    )

    def verify(evidence="dados insuficientes",confidence="Baixa"):
        return {"mode":"Verificar modelo","ok":True,"confidence":confidence,"evidence":evidence}

    def confirmed_remote(evidence,confidence):
        if remote_scope_allowed:
            return {"mode":"Remoto Brasil — confirmado","ok":True,"confidence":confidence,"evidence":evidence}
        if loc_generic:
            return verify("remoto confirmado; disponibilidade para o Brasil não informada","Baixa")
        return {"mode":"Remoto sem disponibilidade para o Brasil","ok":False,
                "confidence":confidence,"evidence":evidence+"; localização incompatível"}

    def confirmed_local(kind,evidence,confidence):
        if local_allowed:
            return {"mode":f"{kind} — confirmado","ok":True,"confidence":confidence,"evidence":evidence}
        if not loc or loc_generic:
            return verify(f"{kind.lower()} confirmado; local não informado","Baixa")
        return {"mode":f"{kind} fora da região","ok":False,"confidence":confidence,"evidence":evidence}

    # Evidências estruturadas são soberanas.
    # search_remote_hint e o booleano legado remote são deliberadamente ignorados:
    # eles podem vir apenas do filtro usado na pesquisa e não provam a modalidade da vaga.
    sr=any(x in workplace for x in ["remote","remoto","telecommute"])
    sh=any(x in workplace for x in ["hybrid","hibrid"])
    so=any(x in workplace for x in ["onsite","on-site","presencial"])

    if sum(bool(x) for x in (sr,sh,so))>1:
        return verify("workplace_type conflitante")

    if sh:
        return confirmed_local("Híbrido","workplace_type","Alta")
    if so:
        return confirmed_local("Presencial","workplace_type","Alta")
    if sr:
        return confirmed_remote("workplace_type/localização","Alta")

    # Título e local publicados pela vaga são evidência explícita, mas nunca usamos a consulta.
    tl=title+" "+loc
    explicit_hybrid=bool(re.search(r"\b(hibrid[oa]|hybrid)\b",tl))
    explicit_onsite=bool(re.search(r"\b(presencial|on[- ]?site|onsite)\b",tl))
    explicit_remote=bool(re.search(r"\b(100% remoto|100% remota|fully remote|remote[- ]only|remoto|remota|remote)\b",tl))

    if explicit_hybrid or explicit_onsite:
        if explicit_remote:
            return verify("título/local conflitante")
        if explicit_hybrid:
            return confirmed_local("Híbrido","título/local","Alta")
        return confirmed_local("Presencial","título/local","Alta")

    if explicit_remote:
        return confirmed_remote("título/local","Alta")

    # Descrição: só frases de modalidade, nunca a palavra isolada "home office".
    head=desc[:4500]
    d_hybrid=bool(re.search(
        r"\b(modelo hibrid[oa]|regime hibrid[oa]|modalidade hibrid[oa]|trabalho hibrid[oa]|"
        r"atuacao hibrida|atuação híbrida|[1-5]\s+dias?\s+(?:por semana\s+)?(?:no|em)\s+escritorio)\b",head))
    d_onsite=bool(re.search(
        r"\b(100% presencial|modelo presencial|regime presencial|modalidade presencial|"
        r"trabalho presencial|atuacao presencial|atuação presencial|atuar presencialmente|"
        r"trabalho no local)\b",head))
    d_remote=bool(re.search(
        r"\b(100% remoto|100% remota|totalmente remoto|totalmente remota|fully remote|"
        r"remote[- ]only|modelo (?:de trabalho )?remoto|regime remoto|modalidade remota?|"
        r"trabalho 100% remoto|atuacao 100% remota?|atuação 100% remota?|"
        r"trabalhe de qualquer lugar|work from anywhere|trabalho remoto integral)\b",head))

    # Presencial/híbrido explícitos têm precedência. Contradição não vira remoto.
    if d_hybrid or d_onsite:
        if d_remote:
            return verify("descrição conflitante")
        if d_hybrid:
            return confirmed_local("Híbrido","descrição","Média")
        return confirmed_local("Presencial","descrição","Média")

    if d_remote:
        return confirmed_remote("descrição","Média")

    if not desc.strip():
        if local_allowed:return verify("local compatível informado; descrição e modalidade ausentes")
        if loc_generic:return verify("localização genérica; descrição e modalidade ausentes")
        return {"mode":"Local fora da região — remoto não confirmado","ok":False,"confidence":"Baixa",
                "evidence":"local externo; descrição e modalidade ausentes"}
    if local_allowed:return verify("local compatível informado; descrição não declara a modalidade","Média")
    if loc_generic:return verify("descrição não declara a modalidade; localização genérica")
    return {"mode":"Local fora da região — remoto não confirmado","ok":False,"confidence":"Média",
            "evidence":"local externo; descrição não confirma trabalho remoto"}

def excluded_location_reason(job,p):
    """Compara somente localização publicada/estruturada com exclusões explícitas."""
    excluded=[clean(value).strip() for value in p.get("localidades_excluidas",[]) if clean(value).strip()]
    if not excluded:return ""
    raw_parts=[job.get("local",""),job.get("applicant_location_requirements","")]
    structured=job.get("structured_location") or job.get("structured_location_json") or ""
    raw_parts.append(compact_json(structured))
    haystack=norm(" | ".join(str(value or "") for value in raw_parts))
    haystack=re.sub(r"\s*[,/|;]+\s*"," ",haystack)
    haystack=re.sub(r"\s+-\s+"," ",haystack)
    haystack=re.sub(r"\s+"," ",haystack).strip()
    if not haystack or haystack in ("nao informado","remote","remoto","worldwide","anywhere"):return ""
    for original in excluded:
        needle=norm(original)
        needle=re.sub(r"\s*[,/|;]+\s*"," ",needle)
        needle=re.sub(r"\s+-\s+"," ",needle)
        needle=re.sub(r"\s+"," ",needle).strip()
        if needle and re.search(rf"(?<!\w){re.escape(needle)}(?!\w)",haystack):return original
    return ""

def location_mode(job,p):
    d=location_decision(job,p)
    return d["mode"],d["ok"]

def closed_job_reason(text):
    t=norm(text)
    phrases=[
        "vaga encerrada","vaga finalizada","vaga fechada","processo seletivo encerrado",
        "processo seletivo finalizado","inscricoes encerradas","inscrições encerradas",
        "candidaturas encerradas","nao aceita mais candidaturas","não aceita mais candidaturas",
        "nao estamos mais aceitando candidaturas","não estamos mais aceitando candidaturas",
        "esta vaga nao esta mais disponivel","esta vaga não está mais disponível",
        "vaga indisponivel","vaga indisponível","oportunidade encerrada",
        "job is no longer available","job no longer available","no longer accepting applications",
        "applications are closed","applications closed","position has been filled",
        "this job has expired","job expired","posting has expired"
    ]
    for ph in phrases:
        if norm(ph) in t:return ph
    return ""

def expired_job_reason(job):
    """Aceita somente evidência explícita de encerramento ou data-limite vencida."""
    text=(job.get("titulo","") or "")+" "+(job.get("descricao","") or "")
    reason=closed_job_reason(text)
    if reason:return reason
    raw=str(job.get("valid_through","") or "").strip()
    if not raw:return ""
    try:
        limit=datetime.fromisoformat(raw.replace("Z","+00:00")).date()
    except Exception:
        try:limit=datetime.strptime(raw[:10],"%Y-%m-%d").date()
        except Exception:return ""
    return "data limite de candidatura encerrada" if limit<datetime.now().date() else ""

def requires_completed_higher_education(text):
    """
    Hard-discard only explicit completed-degree requirements.
    Do not discard 'cursando', 'completo ou cursando', or merely desirable/preferred degrees.
    """
    t=norm(text)
    if not t:return False,""
    # Explicit alternatives that the current profile can satisfy.
    alternatives=[
        "superior completo ou cursando","superior cursando ou completo",
        "graduacao completa ou cursando","graduacao cursando ou completa",
        "ensino superior completo ou em andamento","ensino superior em andamento ou completo",
        "superior em andamento","graduacao em andamento"
    ]
    if any(x in t for x in alternatives):
        return False,""

    patterns=[
        r"\bensino\s+superior\s+complet[oa]\b",
        r"\bsuperior\s+complet[oa]\b",
        r"\bgraduacao\s+complet[oa]\b",
        r"\bcurso\s+superior\s+complet[oa]\b",
        r"\bformacao\s+superior\s+complet[oa]\b",
        r"\bformacao\s+academica\s*:\s*superior\s+complet[oa]\b",
        r"\bbachelor(?:'s)?\s+degree\s+(?:is\s+)?required\b",
        r"\bcompleted\s+bachelor(?:'s)?\s+degree\b"
    ]
    for pat in patterns:
        m=re.search(pat,t)
        if not m:continue
        left=t[max(0,m.start()-90):m.start()]
        # If clearly optional/preferred, don't hard discard.
        if any(x in left for x in ["desejavel","diferencial","preferencial","preferivel","preferred","nice to have"]):
            continue
        return True,m.group(0)
    return False,""

def live_page_invalid_reason(page):
    """Check the opened vacancy page immediately before preparing an application."""
    try:
        body=page.locator("body").inner_text(timeout=8000)
    except:
        return ""
    reason=closed_job_reason(body)
    if reason:return "vaga encerrada/finalizada"
    req,phrase=requires_completed_higher_education(body)
    if req:return "exige ensino superior completo"
    if not pcd_vacancies_enabled(load_profile()):
        pcd_reason=pcd_exclusive_reason(body)
        if pcd_reason:return pcd_reason
    return ""

def pcd_exclusive_reason(text):
    """Descarta somente exclusividade PCD explícita, evitando textos genéricos de diversidade."""
    raw=text or ""
    t=norm(raw)
    # Analisa principalmente o começo do anúncio, onde título/cabeçalho/requisito afirmativo aparecem.
    head=t[:1800]
    patterns=[
        r"\bvaga\s+exclusiva\s+(?:para\s+)?pcd\b",
        r"\bvaga\s+exclusiva\s+(?:para\s+)?pessoas?\s+com\s+deficiencia\b",
        r"\boportunidade\s+exclusiva\s+(?:para\s+)?pcd\b",
        r"\bprocesso\s+seletivo\s+exclusivo\s+(?:para\s+)?pcd\b",
        r"\bvaga\s+afirmativa\s+(?:para\s+)?pcd\b",
        r"\boportunidade\s+afirmativa\s+(?:para\s+)?pcd\b",
        r"\bexclusiv[ao]\s+(?:para\s+)?pcd\b",
        r"\bsomente\s+(?:para\s+)?pcd\b",
        r"\bapenas\s+(?:para\s+)?pcd\b",
        r"\breservad[ao]\s+exclusivamente\s+(?:para\s+)?pcd\b",
        r"\bdestinad[ao]\s+exclusivamente\s+(?:a|para)\s+(?:pcd|pessoas?\s+com\s+deficiencia)\b"
    ]
    return "vaga exclusiva/afirmativa para PCD" if any(re.search(x,head) for x in patterns) else ""

def pcd_job_reason(job):
    """Identifica vaga direcionada a PCD sem confundir mensagens genéricas de inclusão."""
    title=norm(job.get("titulo","") or "")
    title_patterns=[
        r"(?:^|[\s\[(\-\u2013\u2014|/])pcd(?:$|[\s\])\-\u2013\u2014|/])",
        r"\b(?:vaga|oportunidade)\s+(?:para\s+)?pessoas?\s+com\s+deficiencia\b",
        r"\b(?:exclusiva|afirmativa)\s+(?:para\s+)?pessoas?\s+com\s+deficiencia\b",
    ]
    if any(re.search(pattern,title) for pattern in title_patterns):
        return "vaga direcionada/exclusiva para PCD"
    text=(job.get("titulo","") or "")+" "+(job.get("descricao","") or "")
    exclusive=pcd_exclusive_reason(text)
    if exclusive:return exclusive
    body=norm(job.get("descricao","") or "")[:6000]
    pcd=r"(?:pcd|pessoas?\s+com\s+deficiencia)"
    acceptance_patterns=[
        rf"\b(?:vaga|oportunidade|posicao).{{0,45}}(?:aberta|disponivel|destinada).{{0,35}}{pcd}\b",
        rf"\b(?:aceitamos|aceita|incluimos|inclui|incentivamos).{{0,70}}{pcd}\b",
        rf"\b{pcd}.{{0,70}}(?:bem[- ]?vind[ao]s?|podem\s+se\s+candidatar|candidaturas?|sao\s+aceit[ao]s?)\b",
        rf"\b(?:tambem|inclusive).{{0,35}}(?:para\s+)?{pcd}\b",
    ]
    if any(re.search(pattern,body) for pattern in acceptance_patterns):
        return "vaga aberta para PCD"
    return ""

def pcd_vacancies_enabled(p):
    if "buscar_vagas_pcd" in p:return bool(p.get("buscar_vagas_pcd"))
    return not bool(p.get("descartar_vagas_exclusivas_pcd",True))


def publication_age_days(job):
    """Calcula a idade da publicação; retorna None quando a fonte é ambígua."""
    raw=(job.get("data_publicacao","") or "").strip()
    if not raw:
        return None

    t=norm(raw)
    today=datetime.now().date()
    age=None

    if any(x in t for x in ["hoje","today","agora","just now"]):
        age=0
    elif any(x in t for x in ["ontem","yesterday"]):
        age=1
    else:
        m=re.search(r"(\d+)\s*(?:dia|dias|day|days)\b",t)
        if m:
            age=int(m.group(1))
        if age is None:
            m=re.search(r"(\d+)\s*(?:semana|semanas|week|weeks)\b",t)
            if m: age=int(m.group(1))*7
        if age is None:
            m=re.search(r"(\d+)\s*(?:mes|meses|month|months)\b",t)
            if m: age=int(m.group(1))*30

    if age is None:
        candidates=[
            "%Y-%m-%dT%H:%M:%S.%fZ","%Y-%m-%dT%H:%M:%SZ","%Y-%m-%dT%H:%M:%S",
            "%Y-%m-%d","%d/%m/%Y","%d-%m-%Y","%d.%m.%Y"
        ]
        cleaned=raw.strip()
        for fmt in candidates:
            try:
                dt=datetime.strptime(cleaned[:26] if "%f" in fmt else cleaned[:19] if "T" in fmt else cleaned[:10],fmt)
                age=max(0,(today-dt.date()).days)
                break
            except Exception:
                pass

    # Data que não conseguimos interpretar não elimina a vaga.
    if age is None:
        return None

    return age

def vacancy_date_ok(job,p):
    """Retorna (aceita, idade_em_dias). Datas desconhecidas são mantidas."""
    age=publication_age_days(job)
    if age is None:return True,None

    # `idade_maxima_dias` é o campo atual da tela; o nome mais longo é
    # mantido apenas para compatibilidade com perfis antigos.
    max_days=int(p.get("idade_maxima_dias",p.get("idade_maxima_vaga_dias",60)))
    return age<=max_days,age


def mandatory_blocker(job,p):
    """Barreiras objetivas; evita descartar por mera preferência/diferencial."""
    text=norm((job.get("titulo","") or "")+" "+(job.get("descricao","") or ""))
    title=norm(job.get("titulo","") or "")

    # Profissões regulamentadas ou inequivocamente ligadas a uma graduação não
    # podem ganhar relevância pela palavra genérica "assistente" ou "analista".
    profession_courses=(
        (r"\bassistente social\b","servico social","Serviço Social"),
        (r"\bpsicolog[oa]\b","psicologia","Psicologia"),
        (r"\benfermeir[oa]\b","enfermagem","Enfermagem"),
        (r"\bfarmaceutic[oa]\b","farmacia","Farmácia"),
        (r"\bfisioterapeuta\b","fisioterapia","Fisioterapia"),
        (r"\bnutricionista\b","nutricao","Nutrição"),
    )
    courses=profile_courses(p)
    for pattern,required_course,label in profession_courses:
        if re.search(pattern,title) and not any(course_matches_text(course,required_course) for course in courses):
            return f"profissão exige formação em {label}"

    optional=("desejavel","diferencial","preferencial","preferred","nice to have")
    def optional_near(pos):
        return any(x in text[max(0,pos-90):pos] for x in optional)

    # Advocacia/OAB.
    for pat in [r"\boab\s+ativa\b",r"\binscricao\s+(?:ativa\s+)?na\s+oab\b",
                r"\bregistro\s+(?:ativo\s+)?na\s+oab\b",r"\boab\s+obrigatoria\b"]:
        m=re.search(pat,text)
        if m and not optional_near(m.start()):
            return "OAB ativa/obrigatória"

    if re.search(r"\badvogad[oa]\b",title):
        return "cargo de advogado/OAB"

    # Senioridade inequívoca.
    senior=r"\b(senior|sr\.?|especialista|coordenador(?:a)?|gerente|supervisor(?:a)?|head|lead)\b"
    subordinate=bool(re.search(
        r"\b(assistente|auxiliar|secretari[oa])\s+(?:do|da|de)\s+"
        r"(coordenador(?:a)?|gerente|supervisor(?:a)?|head|lead)\b",title))
    if re.search(senior,title) and not subordinate:
        return "senioridade/cargo acima do perfil"

    # Experiência específica longa e explicitamente requerida.
    for m in re.finditer(r"(?:minimo|minima|pelo menos|at least|minimum of)?\s*(\d+)\s*(?:anos|years)\s+(?:de\s+)?experiencia",text):
        try:
            yrs=int(m.group(1))
            context=text[max(0,m.start()-80):m.end()+80]
            if yrs>20 or any(x in context for x in ["empresa","mercado","historia","fundada"]):
                continue
            if yrs>=int(p.get("descartar_experiencia_especifica_anos",5)) and not optional_near(m.start()):
                return f"exige {yrs}+ anos de experiência"
        except: pass

    # Compara a exigência explícita com o nível informado pelo candidato.
    required_english=required_english_level(text)
    declared_english=english_level(p)
    if declared_english=="Não informado" and any(
            skill in ("ingles fluente","ingles avancado","ingles nativo") for skill in profile_skills(p)):
        declared_english="Fluente"
    levels={"Não informado":0,"Básico":1,"Intermediário":2,"Fluente":3}
    if levels[declared_english]<levels[required_english]:
        if required_english=="Fluente":return "inglês fluente obrigatório"
        return f"exige inglês {required_english.lower()} obrigatório"

    return ""

def internship_course_status(job,p=None):
    """Retorna OK_PERFIL, FORA ou REVISAR sem presumir uma formação fixa."""
    text=semantic_norm((job.get("titulo","") or "")+" "+(job.get("descricao","") or ""))
    if not is_intern(job): return ""

    if p:
        courses=[norm(course) for course in internship_selected_areas(p)]
        if courses:
            # O próprio título pode declarar uma formação incompatível sem usar
            # palavras como "cursando" ou "graduação" na descrição.
            title=semantic_norm(job.get("titulo","") or "")
            if any(course_matches_text(course,title) for course in courses):
                return "OK_PERFIL"
            known_courses=[
                "administracao","ciencias contabeis","contabilidade","recursos humanos",
                "contabil","seguranca do trabalho","atracao e selecao","geoprocessamento",
                "marketing","publicidade","jornalismo","pedagogia","psicologia","farmacia",
                "enfermagem","arquitetura","engenharia civil","engenharia mecanica",
                "engenharia eletrica","design","economia","financas","logistica","nutricao",
                "biomedicina","medicina veterinaria","servico social"
            ]
            explicit_title=bool(re.search(
                r"\b(?:estagio|estagiari[oa])\s+(?:tecnico\s+)?(?:em|de|para|na area de|no setor de)?\s*"
                r"(?:"+"|".join(re.escape(course) for course in known_courses)+r")\b",title))
            if explicit_title:return "FORA"
            if any(course_matches_text(course,text) for course in courses):
                return "OK_PERFIL"
            if re.search(r"\b(cursando|curso|graduacao|formacao|estudante|semestre|periodo)\b",text):
                return "FORA"
            return "REVISAR"

    legal=["direito","curso de direito","graduacao em direito","bacharelado em direito"]
    ads=["analise e desenvolvimento de sistemas","sistemas de informacao","ciencia da computacao",
         "engenharia de software","tecnologia da informacao","computacao","desenvolvimento de sistemas",
         "ads"]
    outside=["administracao","ciencias contabeis","contabilidade","recursos humanos","gestao de rh",
             "marketing","publicidade","engenharia civil","engenharia mecanica","pedagogia",
             "psicologia","farmacia","enfermagem","arquitetura"]

    course_context=bool(re.search(r"\b(cursando|curso|graduacao|formacao|estudante|semestre|periodo)\b",text))
    if any(x in text for x in legal+ads) or re.search(r"(?<!\w)ti(?!\w)",text): return "OK_PERFIL"
    if course_context and any(x in text for x in outside): return "FORA"
    return "REVISAR"

def collection_confidence(job):
    company=norm(job.get("empresa",""))
    loc=norm(job.get("local",""))
    desc=(job.get("descricao","") or "").strip()
    if company not in ("","nao informado") and loc not in ("","nao informado") and len(desc)>=250:
        return "Alta"
    if len(desc)>=180 and (company not in ("","nao informado") or loc not in ("","nao informado")):
        return "Média"
    return "Baixa"


def decision_level(job,p,mode):
    if is_intern(job) and internship_course_status(job,p)=="REVISAR":
        return "REVISAR"
    ld=location_decision(job,p)
    if "conflitante" in norm(ld["mode"]):
        return "REVISAR"
    if ld["confidence"]=="Baixa":
        return "REVISAR"
    if "confirmar" in norm(ld["mode"]) or "nao confirmados" in norm(ld["mode"]):
        return "REVISAR"
    if collection_confidence(job)=="Baixa":
        return "REVISAR"
    return "APROVADA"

def hard_filter(job,p):
    date_ok,age_days=vacancy_date_ok(job,p)
    if not date_ok:
        return False,f"vaga antiga ({age_days} dias)" if age_days is not None else "vaga antiga"

    excluded=excluded_location_reason(job,p)
    if excluded:return False,f"localidade excluída pelo usuário — {excluded}"

    text=(job.get("titulo","") or "")+" "+(job.get("descricao","") or "")
    if p.get("descartar_vagas_encerradas",True):
        reason=expired_job_reason(job)
        if reason:return False,"vaga encerrada/finalizada"

    if not pcd_vacancies_enabled(p):
        reason=pcd_job_reason(job)
        if reason:return False,reason

    blocker=mandatory_blocker(job,p)
    if blocker:return False,blocker

    if p.get("descartar_superior_completo_obrigatorio",True):
        req,phrase=requires_completed_higher_education(text)
        if req:return False,"exige ensino superior completo"

    if is_apprentice(job) and not p.get("buscar_jovem_aprendiz",False):
        return False,"Jovem Aprendiz desativado pelo usuário"

    internship_mode=internship_search_mode(p)
    if internship_mode=="somente" and not is_intern(job):
        return False,"modo somente estágios"

    if is_intern(job) and internship_mode=="nao_buscar":
        return False,"estágio desativado pelo usuário"

    if is_intern(job):
        cs=internship_course_status(job,p)
        if cs=="FORA":
            text=norm((job.get("titulo","") or "")+" "+(job.get("descricao","") or ""))
            detected="outro curso"
            mapa=[
                ("ciencias contabeis","Ciências Contábeis"),("contabilidade","Contabilidade"),
                ("administracao","Administração"),("recursos humanos","Recursos Humanos"),
                ("marketing","Marketing"),("publicidade","Publicidade"),
                ("engenharia civil","Engenharia Civil"),("engenharia mecanica","Engenharia Mecânica"),
                ("psicologia","Psicologia"),("pedagogia","Pedagogia"),
                ("enfermagem","Enfermagem"),("arquitetura","Arquitetura")
            ]
            for token,label in mapa:
                if token in text:
                    detected=label;break
            return False,f"estágio fora da formação identificada — {detected}"

    mode,ok=location_mode(job,p)
    if not ok:return False,mode

    return True,mode

def strong_title_profile_match(title,p):
    """Reconhece família profissional no título sem depender da descrição da fonte."""
    value=semantic_norm(title)
    title_tokens=set(re.findall(r"[a-z0-9+#]{3,}",value))
    if not title_tokens:return False
    role_tokens={"assistente","auxiliar","analista","estagio","estagiario","trainee",
                 "consultor","atendente","suporte","support","assistant","intern","junior"}
    stop=role_tokens|{"para","com","the","and","das","dos","uma"}
    queries=list(p.get("consultas_gupy",[]))+list(p.get("consultas_linkedin",[]))
    for query in queries:
        normalized=semantic_norm(query)
        if len(normalized)>=5 and (normalized in value or value in normalized):return True
        query_tokens=set(re.findall(r"[a-z0-9+#]{3,}",normalized))
        common=title_tokens&query_tokens
        domain=common-stop
        if domain and common&role_tokens:return True
        if domain and len(common)>=2 and len(common)>=max(2,int(len(query_tokens)*.6)):return True
    return False

def score_job(job,p,cv):
    title=semantic_norm(job.get("titulo",""));desc=semantic_norm(job.get("descricao",""));allx=title+" "+desc
    parts=[];score=30
    incomplete_description=len((job.get("descricao","") or "").strip())<120
    strong_title=bool(incomplete_description and strong_title_profile_match(title,p))

    terms=profile_terms(p)
    th=sum(1 for x in terms if x in title)
    dh=sum(1 for x in terms if x in desc)
    cargo=min(25,th*10+dh*2)
    score+=cargo
    if cargo:parts.append(f"Cargo/área +{cargo}")

    if strong_title:
        score+=20;parts.append("Título alinhado ao perfil +20")

    skills=set(x for x in profile_skills(p) if x in allx)
    comp=min(22,len(skills)*3)
    score+=comp
    if comp:parts.append(f"Competências +{comp}")

    if p.get("perfil_inicio_carreira",False):
        entry_title=bool(re.search(
            r"\b(auxiliar|atendente|operador(?:a)? de caixa|repositor(?:a)?|estoquista|"
            r"recepcionista|vendedor(?:a)?|separador(?:a)?|conferente|assistente operacional|"
            r"ajudante|empacotador(?:a)?|promotor(?:a)? de vendas|jovem aprendiz|aprendiz)\b",title))
        if entry_title:
            score+=25;parts.append("Cargo de entrada +25")
        if re.search(r"\b(ensino medio|high school)\b",allx):
            score+=8;parts.append("Ensino médio compatível +8")
        if re.search(r"\b(sem experiencia|nao exige experiencia|primeiro emprego|no experience)\b",allx):
            score+=8;parts.append("Não exige experiência +8")

    feedback=p.get("feedback_preferencias",{})
    preference=sum(max(-2,min(2,int(feedback.get(token,0)))) for token in feedback_tokens(allx))
    preference=max(-6,min(6,preference))
    if preference:
        score+=preference;parts.append(f"Preferências {preference:+d}")

    area_hits=sum(1 for area in p.get("areas_curriculo_detectadas",[]) if norm(area) in allx)
    if area_hits:
        bonus=min(12,area_hits*6);score+=bonus;parts.append(f"Área do currículo +{bonus}")

    age=publication_age_days(job)
    freshness=5 if age is not None and age<=3 else 3 if age is not None and age<=7 else 1 if age is not None and age<=14 else 0
    if freshness:
        score+=freshness;parts.append(f"Publicação recente +{freshness}")

    if is_intern(job):
        cs=internship_course_status(job,p)
        if cs=="OK_PERFIL":
            score+=8;parts.append("Estágio compatível +8")

    ld=location_decision(job,p);mode=ld["mode"]
    if "Remoto Brasil" in mode:
        score+=10;parts.append("Remoto elegível +10")
    elif "Local permitido" in mode or "confirmado" in mode:
        score+=7;parts.append("Local/modalidade +7")
    if ld["confidence"]=="Baixa":
        score-=8;parts.append("Local/modalidade incertos -8")

    subordinate=bool(re.search(
        r"\b(assistente|auxiliar|secretari[oa])\s+(?:do|da|de)\s+"
        r"(coordenador(?:a)?|gerente|supervisor(?:a)?|head|lead)\b",title))
    if re.search(r"\b(pleno|senior|sr\.?|especialista|coordenador|coordenadora|gerente|lead|head)\b",title) and not subordinate:
        score-=25;parts.append("Senioridade -25")

    yrs=[int(x) for x in re.findall(r"(\d+)\+?\s*(?:anos|years)\s+(?:de\s+)?(?:experiencia|experience)",desc)]
    if yrs:
        req=max(yrs)
        if req>=5:
            score-=35;parts.append(f"Experiência {req}+ anos -35")
        elif req>=3:
            score-=15;parts.append(f"Experiência {req}+ anos -15")
        elif p.get("perfil_inicio_carreira",False) and req>=1:
            score-=12;parts.append(f"Experiência {req}+ ano(s) -12")

    conf=collection_confidence(job)
    if conf=="Baixa" and not strong_title:
        score-=8;parts.append("Dados incompletos -8")
    elif conf=="Baixa":
        parts.append("Descrição incompleta — revisar")
    elif conf=="Alta":
        score+=3;parts.append("Dados completos +3")

    score=max(0,min(100,score))
    label="Excelente" if score>=85 else "Boa" if score>=70 else "Possível" if score>=55 else "Incompatível"
    reason="Base +30" + ((" | "+" | ".join(parts)) if parts else "")
    return score,label,reason,mode

def curriculum_compatibility(job,p,cv):
    """Usa a mesma régua explicável da lista principal sem alterar a vaga."""
    return score_job(job,p,cv)[0]

def extract_requirements(job):
    text=semantic_norm(job.get("titulo","")+" "+job.get("descricao",""))
    req=[]
    if re.search(r"\bexcel\b",text): req.append(("Excel","excel"))
    if any(x in text for x in ["pacote office","microsoft office"]): req.append(("Pacote Office","office"))
    if any(x in text for x in ["atendimento ao cliente","customer service","customer support"]): req.append(("Atendimento ao cliente","atendimento"))
    if any(x in text for x in ["help desk","service desk","suporte tecnico","suporte técnico"]): req.append(("Suporte/Service Desk","suporte"))
    if any(x in text for x in ["juridico","jurídico","direito","legal operations"]): req.append(("Área jurídica","juridico"))
    if any(x in text for x in ["processos administrativos","rotinas administrativas","administrativo"]): req.append(("Rotinas administrativas","administrativo"))
    if any(x in text for x in ["ingles fluente","inglês fluente","fluent english","native english"]): req.append(("Inglês fluente","ingles_fluente"))
    yrs=[int(x) for x in re.findall(r"(\d+)\+?\s*(?:anos|years)\s+(?:de\s+)?(?:experiencia|experience)",text)]
    if yrs:req.append((f"{max(yrs)}+ anos de experiência específica",f"anos:{max(yrs)}"))
    if any(x in text for x in ["senior","sênior","sr.","tier iii","tier 3","lead","manager","gerente","coordenador"]):
        req.append(("Senioridade elevada","senioridade"))
    return req

def autofill_known_fields(page,p):
    d=p.get("dados_candidatura",{})
    candidates = {
        "nome": d.get("nome_completo",""),
        "name": d.get("nome_completo",""),
        "email": d.get("email",""),
        "e-mail": d.get("email",""),
        "telefone": d.get("telefone",""),
        "phone": d.get("telefone",""),
        "celular": d.get("telefone",""),
        "cidade": d.get("cidade",""),
        "city": d.get("cidade",""),
        "estado": d.get("estado",""),
        "state": d.get("estado",""),
        "country": d.get("pais",""),
        "pais": d.get("pais",""),
        "país": d.get("pais",""),
        "linkedin": d.get("linkedin",""),
        "portfolio": d.get("portfolio","")
    }
    filled=0
    for field in page.locator("input, textarea").all():
        try:
            typ=(field.get_attribute("type") or "text").lower()
            if typ in ("hidden","submit","button","checkbox","radio","file"):continue
            label=" ".join(filter(None,[
                field.get_attribute("name"), field.get_attribute("id"),
                field.get_attribute("placeholder"), field.get_attribute("aria-label")
            ]))
            ln=norm(label)
            value=""
            for k,v in candidates.items():
                if v and k in ln:
                    value=v;break
            if value and not field.input_value():
                field.fill(value);filled+=1
        except:pass
    return filled

def resume_attachment_path(p=None):
    p=p or load_profile()
    configured=p.get("arquivo_curriculo_original","")
    candidates=[]
    if configured:
        candidates.append(configured if os.path.isabs(configured) else os.path.join(BASE_DIR,configured))
    candidates.extend(os.path.join(BASE_DIR,"curriculo_original"+ext) for ext in (".pdf",".docx",".txt"))
    candidates.append(os.path.join(BASE_DIR,"curriculo.pdf"))
    return next((path for path in candidates if os.path.isfile(path)),"")


def attach_resume(page,p=None):
    resume=resume_attachment_path(p)
    if not resume:return 0
    count=0
    for field in page.locator('input[type="file"]').all():
        try:
            accept=norm(field.get_attribute("accept") or "")
            ext=os.path.splitext(resume)[1].lstrip(".")
            if not accept or ext in accept or "document" in accept or "*/*" in accept:
                field.set_input_files(resume);count+=1
        except:pass
    return count

def find_unknown_required_fields(page):
    unknown=[]
    for field in page.locator("input, textarea, select").all():
        try:
            if not field.is_visible():continue
            required=field.get_attribute("required") is not None or field.get_attribute("aria-required")=="true"
            if not required:continue
            tag=field.evaluate("(e)=>e.tagName.toLowerCase()")
            typ=(field.get_attribute("type") or "").lower()
            if typ in ("hidden","submit","button","file"):continue
            val=""
            if tag=="select":
                val=field.input_value()
            elif typ in ("checkbox","radio"):
                val="checked" if field.is_checked() else ""
            else:
                val=field.input_value()
            if val:continue
            desc=" ".join(filter(None,[field.get_attribute("name"),field.get_attribute("id"),
                                       field.get_attribute("placeholder"),field.get_attribute("aria-label")]))
            unknown.append(desc or "campo obrigatório não identificado")
        except:pass
    return unknown[:10]

def detect_browser_path(browser):
    browser=norm(browser)
    envs={k:os.environ.get(k,"") for k in ("LOCALAPPDATA","PROGRAMFILES","PROGRAMFILES(X86)")}
    candidates=[]
    if browser=="brave":
        candidates=[
            os.path.join(envs["LOCALAPPDATA"],"BraveSoftware","Brave-Browser","Application","brave.exe"),
            os.path.join(envs["PROGRAMFILES"],"BraveSoftware","Brave-Browser","Application","brave.exe"),
            os.path.join(envs["PROGRAMFILES(X86)"],"BraveSoftware","Brave-Browser","Application","brave.exe"),
        ]
    elif browser=="chrome":
        candidates=[
            os.path.join(envs["LOCALAPPDATA"],"Google","Chrome","Application","chrome.exe"),
            os.path.join(envs["PROGRAMFILES"],"Google","Chrome","Application","chrome.exe"),
            os.path.join(envs["PROGRAMFILES(X86)"],"Google","Chrome","Application","chrome.exe"),
        ]
    elif browser=="edge":
        candidates=[
            os.path.join(envs["PROGRAMFILES(X86)"],"Microsoft","Edge","Application","msedge.exe"),
            os.path.join(envs["PROGRAMFILES"],"Microsoft","Edge","Application","msedge.exe"),
            os.path.join(envs["LOCALAPPDATA"],"Microsoft","Edge","Application","msedge.exe"),
        ]
    for path in candidates:
        if path and os.path.exists(path):
            return path
    return ""

def get_browser_launch_settings(p):
    choice=norm(p.get("navegador_automacao","automatico"))
    aliases={"automático":"automatico","google chrome":"chrome","microsoft edge":"edge",
             "chromium interno":"chromium","navegador interno":"chromium"}
    choice=aliases.get(choice,choice)

    # Automático: Chrome -> Edge -> Brave -> Chromium do Playwright.
    if choice in ("automatico","auto",""):
        for name in ("chrome","edge","brave"):
            path=detect_browser_path(name)
            if path:
                return {"browser":"chromium","channel":None,"executable_path":path,
                        "profile_dir":os.path.join(BASE_DIR,"browser_profiles",name),
                        "label":name.title()}
        return {"browser":"chromium","channel":None,"executable_path":None,
                "profile_dir":os.path.join(BASE_DIR,"browser_profiles","chromium"),
                "label":"Chromium"}

    if choice in ("chrome","edge","brave"):
        path=detect_browser_path(choice)
        if not path:
            raise RuntimeError(f"{choice.title()} não foi encontrado neste computador. Escolha Automático ou outro navegador.")
        return {"browser":"chromium","channel":None,"executable_path":path,
                "profile_dir":os.path.join(BASE_DIR,"browser_profiles",choice),
                "label":choice.title()}

    if choice=="firefox":
        return {"browser":"firefox","channel":None,"executable_path":None,
                "profile_dir":os.path.join(BASE_DIR,"browser_profiles","firefox"),
                "label":"Firefox (Playwright)"}

    return {"browser":"chromium","channel":None,"executable_path":None,
            "profile_dir":os.path.join(BASE_DIR,"browser_profiles","chromium"),
            "label":"Chromium"}


def simple_location_mode(local, modalidade):
    loc=(local or "").strip()
    mode=norm(modalidade or "")

    if mode.startswith("remoto brasil") and "confirmado" in mode:
        simple_mode="Remoto"
        if not loc or norm(loc) in ("remoto","remota","remote","brasil","brazil","anywhere","worldwide"):
            return "Remoto"
        return f"{loc} • Remoto"

    if mode.startswith("remoto"):
        simple_mode="Verificar modelo"
    elif mode.startswith("presencial"):
        simple_mode="Presencial"
    elif mode.startswith("hibrid") or mode.startswith("híbr"):
        simple_mode="Híbrido"
    elif "conflit" in mode or "confirm" in mode or "verificar" in mode or not mode:
        simple_mode="Verificar modelo"
    else:
        simple_mode="Verificar modelo"

    if not loc or norm(loc) in ("nao informado","não informado",""):
        return simple_mode

    return f"{loc} • {simple_mode}"

def format_date_br(value):
    raw=str(value or "").strip()
    if not raw:return "Data não informada"
    try:return datetime.fromisoformat(raw.replace("Z","+00:00")).strftime("%d/%m/%Y")
    except Exception:pass
    for fmt in ("%d/%m/%Y","%d-%m-%Y","%Y/%m/%d"):
        try:return datetime.strptime(raw[:10],fmt).strftime("%d/%m/%Y")
        except Exception:pass
    return raw[:10]


def jobs_query(view="todas",search=""):
    sql=("SELECT id,score,titulo,empresa,local,modalidade,COALESCE(NULLIF(categoria,''),'Geral'),status,"
         "COALESCE(selecionada_lote,0),COALESCE(data_publicacao,'') "
         "FROM vagas")
    clauses=[];params=[]
    if view=="candidaturas":
        clauses.append("status IN ('Candidatado','Entrevista','Rejeitado')")
    else:
        clauses.extend(["status='Nova'","COALESCE(selecionada_lote,0)=0","decisao IN ('APROVADA','REVISAR')"])
        if view=="recomendadas":
            clauses.append("decisao='APROVADA'")
        elif view=="revisar":
            clauses.append("decisao='REVISAR'")
        elif view=="estagio":
            clauses.append("(categoria LIKE 'Estágio%' OR lower(titulo) LIKE '%estágio%' OR lower(titulo) LIKE '%estagio%')")
        elif view=="aprendiz":
            clauses.append("(lower(titulo) LIKE '%jovem aprendiz%' OR lower(titulo) LIKE '%menor aprendiz%' OR lower(titulo) LIKE '%aprendiz%')")
        elif view=="pcd":
            clauses.append("""(lower(titulo) LIKE '%pcd%'
                OR lower(descricao) LIKE '%para pcd%'
                OR lower(descricao) LIKE '%vaga%pcd%'
                OR lower(descricao) LIKE '%pessoa com defici%'
                OR lower(descricao) LIKE '%pessoas com defici%')""")
        elif view in ("remoto","home_office"):
            clauses.append("modalidade LIKE 'Remoto%confirmado%'")
            clauses.append("COALESCE(categoria,'') NOT LIKE 'Estágio%'")
        elif view=="presencial":
            clauses.append("(modalidade LIKE 'Presencial%' OR modalidade LIKE 'Híbrido%')")
            clauses.append("COALESCE(categoria,'') NOT LIKE 'Estágio%'")
    if search.strip():
        clauses.append("(titulo LIKE ? OR empresa LIKE ? OR descricao LIKE ? OR local LIKE ? OR modalidade LIKE ? OR fonte LIKE ?)")
        value="%"+search.strip()+"%";params.extend([value]*6)
    if clauses:sql+=" WHERE "+" AND ".join(clauses)
    sql+=" ORDER BY CASE decisao WHEN 'APROVADA' THEN 0 ELSE 1 END,score DESC,id DESC"
    return sql,params

class App(tk.Tk):
    def __init__(self):
        super().__init__();self.title(f"{APP_TITLE} {APP_VERSION}")
        self.instance_socket=socket.socket(socket.AF_INET,socket.SOCK_STREAM)
        try:self.instance_socket.bind(("127.0.0.1",47832))
        except OSError:
            messagebox.showinfo(APP_TITLE,"O aplicativo já está aberto.")
            self.destroy();raise SystemExit
        try:backup_database()
        except Exception:LOGGER.exception("Não foi possível criar o backup automático")
        self.geometry("1420x820")
        self.minsize(1100,700)
        self.configure(bg="#141922");self.geometry("1450x820");self.minsize(1150,680)
        if requests is None or BeautifulSoup is None:
            messagebox.showerror("Dependências","Execute iniciar.bat para instalar requests e beautifulsoup4.")
        self.p=load_profile();self.cv=read_cv();self.conn=connect_database(DB_PATH);self.current=None
        if self.cv.strip() and int(self.p.get("perfil_curriculo_versao",0) or 0)<4:
            adapt_profile_to_cv(self.p,self.cv);save_json_file(PROFILE_PATH,self.p)
        self.search_running=False;self.open_windows={};self.closing=False;self.last_source_counts={}
        self.shutdown_event=threading.Event();self.worker_threads=set();self.worker_lock=threading.Lock()
        self.db();self.load_feedback_profile();self.migrate_v19();self.migrate_v23();self.migrate_v24();self.migrate_v25();self.migrate_v26();self.migrate_v27();self.ui()
        self.apply_pcd_preference(pcd_vacancies_enabled(self.p))
        self.apply_internship_mode(internship_search_mode(self.p))
        self.apply_apprentice_preference(bool(self.p.get("buscar_jovem_aprendiz",False)))
        self.apply_international_preference(international_search_enabled(self.p))
        self.apply_excluded_locations()
        self.separate_outside_region_jobs();self.refresh()
        self.protocol("WM_DELETE_WINDOW",self.close_app)
        self.after(250,self.first_run_privacy_flow)

    def db(self):
        self.conn.execute("""CREATE TABLE IF NOT EXISTS vagas(
        id INTEGER PRIMARY KEY AUTOINCREMENT,titulo TEXT,empresa TEXT,local TEXT,modalidade TEXT,descricao TEXT,
        url TEXT UNIQUE,fonte TEXT,data_publicacao TEXT,salario TEXT,score INTEGER,classificacao TEXT,motivo TEXT,
        status TEXT DEFAULT 'Nova',criada_em TEXT DEFAULT CURRENT_TIMESTAMP)""")
        self.conn.execute("""CREATE TABLE IF NOT EXISTS descartadas(
        id INTEGER PRIMARY KEY AUTOINCREMENT,titulo TEXT,empresa TEXT,local TEXT,descricao TEXT,
        url TEXT UNIQUE,fonte TEXT,data_publicacao TEXT,salario TEXT,motivo_descarte TEXT,
        descartada_em TEXT DEFAULT CURRENT_TIMESTAMP)""")
        for col,ddl in [("workplace_type","TEXT DEFAULT ''"),
                        ("applicant_location_requirements","TEXT DEFAULT ''"),
                        ("structured_location_json","TEXT DEFAULT ''")]:
            try:self.conn.execute(f"ALTER TABLE descartadas ADD COLUMN {col} {ddl}")
            except sqlite3.OperationalError:pass
        for col,ddl in [
            ("selecionada_lote","INTEGER DEFAULT 0"),
            ("tentativas_envio","INTEGER DEFAULT 0"),
            ("ultimo_resultado","TEXT DEFAULT ''"),
            ("workplace_type","TEXT DEFAULT ''"),
            ("location_confidence","TEXT DEFAULT ''"),
            ("location_evidence","TEXT DEFAULT ''"),
            ("workplace_type_raw","TEXT DEFAULT ''"),
            ("workplace_source","TEXT DEFAULT ''"),
            ("structured_location_json","TEXT DEFAULT ''"),
            ("applicant_location_requirements","TEXT DEFAULT ''"),
            ("remote_eligible_brazil","INTEGER DEFAULT NULL"),
            ("modality_checked_at","TEXT DEFAULT ''"),
            ("description_status","TEXT DEFAULT ''"),
            ("description_attempts","INTEGER DEFAULT 0"),
            ("description_last_error","TEXT DEFAULT ''"),
            ("description_last_attempt_at","TEXT DEFAULT ''"),
            ("description_next_retry_at","TEXT DEFAULT ''"),
            ("description_source","TEXT DEFAULT ''"),
            ("candidatura_em","TEXT DEFAULT ''")
        ]:
            try:self.conn.execute(f"ALTER TABLE vagas ADD COLUMN {col} {ddl}")
            except sqlite3.OperationalError:pass
        try:self.conn.execute("ALTER TABLE vagas ADD COLUMN categoria TEXT DEFAULT ''")
        except sqlite3.OperationalError:pass
        try:self.conn.execute("ALTER TABLE vagas ADD COLUMN decisao TEXT DEFAULT 'APROVADA'")
        except sqlite3.OperationalError:pass
        try:self.conn.execute("ALTER TABLE vagas ADD COLUMN confianca TEXT DEFAULT 'Média'")
        except sqlite3.OperationalError:pass
        try:self.conn.execute("ALTER TABLE vagas ADD COLUMN fontes TEXT DEFAULT ''")
        except sqlite3.OperationalError:pass
        try:self.conn.execute("ALTER TABLE vagas ADD COLUMN fingerprint TEXT DEFAULT ''")
        except sqlite3.OperationalError:pass
        try:self.conn.execute("CREATE INDEX IF NOT EXISTS idx_vagas_fingerprint ON vagas(fingerprint)")
        except sqlite3.OperationalError:pass
        self.conn.execute("""CREATE TABLE IF NOT EXISTS preferencias_feedback(
            termo TEXT PRIMARY KEY,positivo INTEGER DEFAULT 0,negativo INTEGER DEFAULT 0)""")
        self.conn.execute("CREATE TABLE IF NOT EXISTS app_meta(chave TEXT PRIMARY KEY, valor TEXT)")
        try:
            rows=self.conn.execute("SELECT id,titulo,empresa,local FROM vagas WHERE COALESCE(fingerprint,'')=''").fetchall()
            for vid,t,e,l in rows:
                fp=job_fingerprint({"titulo":t,"empresa":e,"local":l})
                if fp:self.conn.execute("UPDATE vagas SET fingerprint=? WHERE id=?",(fp,vid))
        except Exception:
            LOGGER.exception("Falha ao preencher fingerprints durante a migration")
        self.conn.commit()

    def start_worker(self,target,*args,name="trabalho"):
        def wrapped():
            try:target(*args)
            finally:
                with self.worker_lock:self.worker_threads.discard(threading.current_thread())
        worker=threading.Thread(target=wrapped,name=name,daemon=False)
        with self.worker_lock:self.worker_threads.add(worker)
        worker.start();return worker

    def ui_call(self,callback):
        if self.__dict__.get("closing",False):return False
        try:self.after(0,callback);return True
        except tk.TclError:return False

    def shutdown_requested(self):
        event=self.__dict__.get("shutdown_event")
        return bool(event and event.is_set())

    def close_app(self):
        if self.closing:return
        with self.worker_lock:active=[worker for worker in self.worker_threads if worker.is_alive()]
        if active:
            if not messagebox.askyesno("Encerrar com segurança",
                "Há uma tarefa em andamento. Deseja cancelá-la e aguardar a finalização segura do aplicativo?"):
                return
            self.closing=True;self.shutdown_event.set()
            try:self.info.set("Encerrando com segurança…")
            except Exception:pass
            self.after(100,self.wait_workers_before_close);return
        self.finalize_close()

    def wait_workers_before_close(self):
        with self.worker_lock:active=[worker for worker in self.worker_threads if worker.is_alive()]
        if active:
            self.after(100,self.wait_workers_before_close);return
        self.finalize_close()

    def finalize_close(self):
        try:flush_detail_cache(force=True)
        except Exception:LOGGER.exception("Falha ao persistir cache no encerramento")
        try:self.conn.close()
        except Exception:pass
        try:self.instance_socket.close()
        except Exception:pass
        self.destroy()

    def meta_get(self,key,default=""):
        row=self.conn.execute("SELECT valor FROM app_meta WHERE chave=?",(key,)).fetchone()
        return row[0] if row else default

    def meta_set(self,key,value):
        with self.conn:
            self.conn.execute("INSERT OR REPLACE INTO app_meta(chave,valor) VALUES(?,?)",(key,str(value)))

    def privacy_notice_accepted(self):
        return (self.meta_get("aviso_privacidade_versao")==PRIVACY_NOTICE_VERSION and
                self.meta_get("termos_uso_versao")==TERMS_OF_USE_VERSION)

    def pcd_consent_valid(self):
        return (self.p.get("consentimento_pcd_versao")==PCD_CONSENT_VERSION
                and bool(self.p.get("consentimento_pcd_em")))

    def request_pcd_consent(self,parent=None):
        accepted=messagebox.askyesno(
            "Preferência opcional para vagas PCD",
            "Ativar este filtro registra neste computador que o usuário deseja visualizar vagas "
            "direcionadas ou abertas para PCD. Essa preferência pode permitir uma inferência "
            "relacionada à saúde.\n\nO dado não é enviado ao desenvolvedor nem usado para enviar "
            "candidaturas automaticamente. Ele pode ser removido desativando a opção ou usando "
            "Limpar tudo.\n\nDeseja ativar?",parent=parent)
        if accepted:
            self.p["consentimento_pcd_versao"]=PCD_CONSENT_VERSION
            self.p["consentimento_pcd_em"]=datetime.now(timezone.utc).isoformat(timespec="seconds")
        else:
            self.p["buscar_vagas_pcd"]=False
            self.p["descartar_vagas_exclusivas_pcd"]=True
            self.p["consentimento_pcd_versao"]=""
            self.p["consentimento_pcd_em"]=""
        save_json_file(PROFILE_PATH,self.p)
        return accepted

    def first_run_privacy_flow(self):
        if not self.privacy_notice_accepted():
            self.show_privacy_notice(required=True)
            return
        self.continue_after_privacy()

    def continue_after_privacy(self):
        if self.p.get("buscar_vagas_pcd",False) and not self.pcd_consent_valid():
            if not self.request_pcd_consent(self):self.apply_pcd_preference(False)
        if not self.cv.strip():self.after(150,self.edit_profile)

    def show_privacy_notice(self,required=False):
        win,created=self.managed_window("privacidade","Privacidade e Termos de Uso", "760x680",modal=True)
        if not created:return
        body=ttk.Frame(win,padding=18);body.pack(fill="both",expand=True)
        ttk.Label(body,text="Privacidade e Termos de Uso",font=("Segoe UI",17,"bold")).pack(anchor="w")
        ttk.Label(body,text=f"Versão do aplicativo: {APP_VERSION}",style="Muted.Panel.TLabel").pack(anchor="w",pady=(2,10))
        text_box=tk.Text(body,wrap="word",bg=self.colors["white"],fg=self.colors["ink"],
                         insertbackground=self.colors["ink"],relief="flat",padx=12,pady=12)
        combined_notice=PRIVACY_NOTICE+"\n\n"+("="*68)+"\n\n"+TERMS_OF_USE
        text_box.insert("1.0",combined_notice);text_box.configure(state="disabled");text_box.pack(fill="both",expand=True)
        buttons=ttk.Frame(body);buttons.pack(fill="x",pady=(12,0))
        if required:
            def accept():
                self.meta_set("aviso_privacidade_versao",PRIVACY_NOTICE_VERSION)
                self.meta_set("aviso_privacidade_aceito_em",datetime.now(timezone.utc).isoformat(timespec="seconds"))
                self.meta_set("termos_uso_versao",TERMS_OF_USE_VERSION)
                self.meta_set("termos_uso_aceito_em",datetime.now(timezone.utc).isoformat(timespec="seconds"))
                win._managed_close();self.continue_after_privacy()
            ttk.Button(buttons,text="Sair",style="Danger.TButton",command=self.close_app).pack(side="left")
            ttk.Button(buttons,text="Aceito e continuar",style="Primary.TButton",command=accept).pack(side="right")
        else:
            ttk.Button(buttons,text="Fechar",style="Primary.TButton",command=win._managed_close).pack(side="right")

    def show_credits(self):
        win,created=self.managed_window("creditos","Créditos e apoio", "780x650")
        if not created:return
        body=ttk.Frame(win,padding=22);body.pack(fill="both",expand=True)
        ttk.Label(body,text="Créditos",font=("Segoe UI",18,"bold")).pack(anchor="w")

        content=ttk.Frame(body);content.pack(fill="both",expand=True,pady=(14,0))
        thanks=ttk.Frame(content);thanks.pack(side="left",fill="both",expand=True,padx=(0,22))
        message=("Agradeço imensamente a todos que apoiaram a ideia e testaram, "
                 "me ajudando do início ao fim.\n\n"
                 "Buiuz\nComedorDeTui\nTuieba\nArigher\nEnderionvel\nKamyh\nFolkss\nAkamui")
        ttk.Label(thanks,text=message,wraplength=340,justify="left").pack(anchor="w",pady=(0,18))
        contact=ttk.LabelFrame(thanks,text="Contato",padding=12);contact.pack(fill="x",anchor="w")
        ttk.Label(contact,text=AUTHOR_EMAIL).pack(anchor="w",pady=(0,8))
        contact_buttons=ttk.Frame(contact);contact_buttons.pack(fill="x")
        ttk.Button(contact_buttons,text="LinkedIn",style="Soft.TButton",
                   command=lambda:webbrowser.open(AUTHOR_LINKEDIN)).pack(side="left",padx=(0,6))
        ttk.Button(contact_buttons,text="GitHub",style="Soft.TButton",
                   command=lambda:webbrowser.open(AUTHOR_GITHUB)).pack(side="left")
        def copy_email():
            win.clipboard_clear();win.clipboard_append(AUTHOR_EMAIL);win.update()
            email_status.set("E-mail copiado")
        email_status=tk.StringVar(value="Copiar e-mail")
        ttk.Button(contact,textvariable=email_status,command=copy_email).pack(fill="x",pady=(8,0))

        support=ttk.LabelFrame(content,padding=16)
        support.pack(side="right",fill="y")
        ttk.Label(support,
                  text="A cada R$ 7 doados, eu tomo uma cervejinha enquanto desejo a todos boa sorte na busca por uma nova oportunidade. 🍺",
                  wraplength=280,justify="center").pack(pady=(0,10))
        qr_path=resource_path("pix_qr_only.png")
        try:
            qr=tk.PhotoImage(file=qr_path)
            factor=max(1,max(qr.width(),qr.height())//220)
            if factor>1:qr=qr.subsample(factor,factor)
            win.qr_image=qr
            qr_frame=tk.Frame(support,bg="#2a3545",padx=7,pady=7)
            qr_frame.pack(pady=5)
            tk.Label(qr_frame,image=qr,bg="#2a3545",borderwidth=0).pack()
        except Exception:
            LOGGER.exception("Não foi possível exibir o QR Code PIX")
            ttk.Label(support,text="QR Code indisponível nesta instalação.").pack(pady=12)
        ttk.Label(support,text=f"Chave PIX: {DONATION_PIX}",font=("Segoe UI",10,"bold")).pack(pady=(8,5))
        def copy_pix():
            win.clipboard_clear();win.clipboard_append(DONATION_PIX);win.update()
            copy_status.set("Chave PIX copiada")
        copy_status=tk.StringVar(value="Copiar chave PIX")
        ttk.Button(support,textvariable=copy_status,style="Primary.TButton",command=copy_pix).pack(fill="x")
        ttk.Label(support,text="Contribuição totalmente opcional.",
                  style="Muted.Panel.TLabel").pack(pady=(9,0))

        ttk.Button(body,text="Fechar",command=win._managed_close).pack(anchor="e",pady=(14,0))

    def load_feedback_profile(self):
        profile = self.__dict__.get("p") if hasattr(self, "__dict__") else None
        if profile is None:
            return
        try:
            profile["feedback_preferencias"]={term:int(pos or 0)-int(neg or 0)
                for term,pos,neg in self.conn.execute("SELECT termo,positivo,negativo FROM preferencias_feedback")}
        except Exception:
            profile["feedback_preferencias"]={}

    def record_feedback(self,vid,positive):
        row=self.conn.execute("SELECT titulo FROM vagas WHERE id=?",(vid,)).fetchone()
        if not row:return
        column="positivo" if positive else "negativo"
        for token in feedback_tokens(row[0]):
            self.conn.execute("INSERT OR IGNORE INTO preferencias_feedback(termo) VALUES(?)",(token,))
            self.conn.execute(f"UPDATE preferencias_feedback SET {column}={column}+1 WHERE termo=?",(token,))
        self.conn.commit()
        self.load_feedback_profile()

    def migrate_v19(self):
        """Corrige classificações antigas do LinkedIn sem apagar vagas/candidaturas."""
        try:
            self.conn.execute("CREATE TABLE IF NOT EXISTS app_meta(chave TEXT PRIMARY KEY, valor TEXT)")
            done=self.conn.execute("SELECT valor FROM app_meta WHERE chave='migracao_v19' ").fetchone()
            if done:return
            # Versões anteriores usavam o filtro da pesquisa remota como se fosse prova da modalidade.
            # Essas vagas passam para 'a confirmar' até serem coletadas novamente com evidência real.
            self.conn.execute("""UPDATE vagas
                               SET modalidade='Local/modalidade não confirmados', decisao='REVISAR'
                               WHERE fonte='LinkedIn' AND status='Nova'
                                 AND modalidade LIKE 'Remoto%'""")
            self.conn.execute("INSERT OR REPLACE INTO app_meta(chave,valor) VALUES('migracao_v19','1')")
            self.conn.commit()
        except Exception:
            LOGGER.exception("Falha na migration v19")

    def migrate_v23(self):
        """Reclassifica modalidade sem apagar vagas nem alterar candidaturas."""
        try:
            self.conn.execute("CREATE TABLE IF NOT EXISTS app_meta(chave TEXT PRIMARY KEY, valor TEXT)")
            done=self.conn.execute("SELECT valor FROM app_meta WHERE chave='migracao_v23_modalidade'").fetchone()
            if done:return
            rows=self.conn.execute("""SELECT id,titulo,local,descricao,fonte,COALESCE(workplace_type,'')
                                      FROM vagas""").fetchall()
            for vid,title,local,description,source,workplace_type in rows:
                job={"titulo":title or "","local":local or "","descricao":description or "",
                     "fonte":source or "","workplace_type":workplace_type or ""}
                ld=location_decision(job,self.p)
                decision="REVISAR" if not ld["ok"] else decision_level(job,self.p,ld["mode"])
                self.conn.execute("""UPDATE vagas SET modalidade=?,decisao=?,location_confidence=?,
                                      location_evidence=? WHERE id=?""",
                                  (ld["mode"],decision,ld["confidence"],ld["evidence"],vid))
            self.conn.execute("INSERT OR REPLACE INTO app_meta(chave,valor) VALUES('migracao_v23_modalidade','1')")
            self.conn.commit()
        except Exception:
            self.conn.rollback()
            LOGGER.exception("Falha na migration de modalidade v23")

    def migrate_v24(self):
        """Preserva dados estruturados existentes sem remover ou recriar tabelas."""
        try:
            self.conn.execute("CREATE TABLE IF NOT EXISTS app_meta(chave TEXT PRIMARY KEY, valor TEXT)")
            done=self.conn.execute("SELECT valor FROM app_meta WHERE chave='migracao_v24_dados_estruturados'").fetchone()
            if done:return
            rows=self.conn.execute("""SELECT id,local,modalidade,COALESCE(workplace_type,''),
                                      COALESCE(workplace_type_raw,''),COALESCE(workplace_source,''),
                                      COALESCE(structured_location_json,''),remote_eligible_brazil,
                                      COALESCE(modality_checked_at,'') FROM vagas""").fetchall()
            checked=datetime.now(timezone.utc).isoformat(timespec="seconds")
            for vid,local,mode,workplace,raw,source,structured,eligible,checked_at in rows:
                raw=raw or workplace or ""
                source=source or ("legacy" if workplace else "")
                structured=structured or compact_json({"display":local or ""})
                if eligible is None:
                    normalized_mode=norm(mode)
                    if normalized_mode.startswith("remoto brasil") and "confirmado" in normalized_mode:eligible=1
                    elif normalized_mode.startswith("remoto sem"):eligible=0
                self.conn.execute("""UPDATE vagas SET workplace_type_raw=?,workplace_source=?,
                                      structured_location_json=?,remote_eligible_brazil=?,modality_checked_at=?
                                      WHERE id=?""",
                                  (raw,source,structured,eligible,checked_at or checked,vid))
            self.conn.execute("INSERT OR REPLACE INTO app_meta(chave,valor) VALUES('migracao_v24_dados_estruturados','1')")
            self.conn.commit()
        except Exception:
            self.conn.rollback()
            LOGGER.exception("Falha na migration de dados estruturados v24")

    def migrate_v25(self):
        """Classifica descrições existentes; migration aditiva e idempotente."""
        try:
            self.conn.execute("CREATE TABLE IF NOT EXISTS app_meta(chave TEXT PRIMARY KEY, valor TEXT)")
            done=self.conn.execute("SELECT valor FROM app_meta WHERE chave='migracao_v25_descricoes'").fetchone()
            if done:return
            self.conn.execute("""UPDATE vagas SET description_status=CASE
                WHEN LENGTH(TRIM(COALESCE(descricao,'')))>=120 THEN 'disponivel'
                WHEN fonte='LinkedIn' THEN 'pendente'
                ELSE 'indisponivel' END
                WHERE COALESCE(description_status,'')=''""")
            self.conn.execute("""UPDATE vagas SET description_source=fonte
                                 WHERE COALESCE(description_source,'')='' AND LENGTH(TRIM(COALESCE(descricao,'')))>=120""")
            self.conn.execute("INSERT OR REPLACE INTO app_meta(chave,valor) VALUES('migracao_v25_descricoes','1')")
            self.conn.commit()
        except Exception:
            self.conn.rollback();LOGGER.exception("Falha na migration v25 de descrições")

    def migrate_v26(self):
        """Adiciona a data da candidatura preservando registros antigos."""
        try:
            self.conn.execute("CREATE TABLE IF NOT EXISTS app_meta(chave TEXT PRIMARY KEY, valor TEXT)")
            done=self.conn.execute("SELECT valor FROM app_meta WHERE chave='migracao_v26_data_candidatura'").fetchone()
            if done:return
            self.conn.execute("""UPDATE vagas SET candidatura_em=COALESCE(NULLIF(criada_em,''),CURRENT_TIMESTAMP)
                                 WHERE status IN ('Candidatado','Entrevista','Rejeitado')
                                   AND COALESCE(candidatura_em,'')=''""")
            self.conn.execute("INSERT OR REPLACE INTO app_meta(chave,valor) VALUES('migracao_v26_data_candidatura','1')")
            self.conn.commit()
        except Exception:
            self.conn.rollback();LOGGER.exception("Falha na migration v26 de candidatura")

    def migrate_v27(self):
        """Reconhece remoto estruturado da Gupy Brasil sem alterar status ou histórico."""
        try:
            self.conn.execute("CREATE TABLE IF NOT EXISTS app_meta(chave TEXT PRIMARY KEY, valor TEXT)")
            done=self.conn.execute("SELECT valor FROM app_meta WHERE chave='migracao_v27_remoto_gupy'").fetchone()
            if done:return
            rows=self.conn.execute("""SELECT id,titulo,empresa,local,descricao,fonte,COALESCE(workplace_type,'')
                                      FROM vagas WHERE lower(fonte)='gupy' AND lower(workplace_type)='remote'""").fetchall()
            checked=datetime.now(timezone.utc).isoformat(timespec="seconds")
            for vid,title,company,local,description,source,workplace in rows:
                job={"titulo":title or "","empresa":company or "","local":local or "",
                     "descricao":description or "","fonte":source or "","workplace_type":workplace or ""}
                ld=location_decision(job,self.p)
                self.conn.execute("""UPDATE vagas SET modalidade=?,decisao=?,location_confidence=?,
                                      location_evidence=?,remote_eligible_brazil=?,modality_checked_at=? WHERE id=?""",
                                  (ld["mode"],decision_level(job,self.p,ld["mode"]),ld["confidence"],
                                   ld["evidence"],1 if ld["mode"].startswith("Remoto Brasil") else 0,checked,vid))
            self.conn.execute("INSERT OR REPLACE INTO app_meta(chave,valor) VALUES('migracao_v27_remoto_gupy','1')")
            self.conn.commit()
        except Exception:
            self.conn.rollback();LOGGER.exception("Falha na migration v27 de remoto Gupy")


    def ui(self):
        style=ttk.Style(self)
        try: style.theme_use("clam")
        except Exception: pass
        colors={"bg":"#141922","panel":"#1b2230","white":"#202938","ink":"#e6eaf0",
                "muted":"#aab4c3","blue":"#5b8def","blue_dark":"#b8d0ff","blue_soft":"#304a70",
                "green":"#63c7a0","green_soft":"#203c35","danger":"#e58a96","line":"#344154"}
        self.colors=colors
        self.configure(bg=colors["bg"])
        self.ui_images={}
        def ui_photo(name,max_width,max_height):
            try:
                original=tk.PhotoImage(file=resource_path(name))
                factor=max(1,(original.width()+max_width-1)//max_width,
                           (original.height()+max_height-1)//max_height)
                image=original.subsample(factor,factor) if factor>1 else original
                self.ui_images[name]=image
                return image
            except Exception:
                LOGGER.exception("Não foi possível carregar o recurso visual %s",name)
                return None
        style.configure(".",font=("Segoe UI",10),background=colors["bg"],foreground=colors["ink"])
        style.configure("TFrame",background=colors["bg"])
        style.configure("TLabel",background=colors["bg"],foreground=colors["ink"])
        style.configure("Panel.TFrame",background=colors["panel"])
        style.configure("Panel.TLabel",background=colors["panel"],foreground=colors["ink"])
        style.configure("Muted.Panel.TLabel",background=colors["panel"],foreground=colors["muted"])
        style.configure("TButton",font=("Segoe UI",10,"bold"),padding=(15,10),borderwidth=0,
                        background="#273244",foreground=colors["ink"])
        style.map("TButton",background=[("active","#344258"),("disabled","#1b2230")],
                  foreground=[("disabled","#758195")])
        style.configure("TEntry",fieldbackground=colors["white"],foreground=colors["ink"],bordercolor=colors["line"])
        style.configure("TCombobox",fieldbackground=colors["white"],background=colors["panel"],
                        foreground=colors["ink"],arrowcolor=colors["ink"])
        style.map("TCombobox",fieldbackground=[("readonly",colors["white"])],foreground=[("readonly",colors["ink"])])
        style.configure("TCheckbutton",background=colors["bg"],foreground=colors["ink"],
                        focuscolor=colors["bg"],indicatorbackground=colors["white"],
                        indicatorforeground=colors["blue_dark"])
        style.map("TCheckbutton",
                  background=[("active",colors["bg"]),("pressed",colors["bg"]),
                              ("selected",colors["bg"]),("disabled",colors["bg"])],
                  foreground=[("active",colors["ink"]),("pressed",colors["ink"]),
                              ("selected",colors["ink"]),("disabled",colors["muted"])],
                  indicatorbackground=[("active",colors["white"]),("selected",colors["blue"]),
                                       ("disabled",colors["panel"])])
        style.configure("TLabelframe",background=colors["bg"],foreground=colors["ink"],bordercolor=colors["line"])
        style.configure("TLabelframe.Label",background=colors["bg"],foreground=colors["ink"])
        style.configure("Primary.TButton",background=colors["blue"],foreground="white",padding=(20,12))
        style.map("Primary.TButton",background=[("active","#709cf1"),("disabled","#33445f")])
        style.configure("Soft.TButton",background=colors["blue_soft"],foreground=colors["blue_dark"])
        style.map("Soft.TButton",background=[("active","#3b5a86")])
        style.configure("Danger.TButton",background="#4a2d36",foreground=colors["danger"])
        style.map("Danger.TButton",background=[("active","#5c3742")])
        style.configure("Summary.TButton",background=colors["panel"],foreground=colors["muted"],
                        font=("Segoe UI",10),padding=(6,5),anchor="w",borderwidth=0)
        style.map("Summary.TButton",background=[("active",colors["blue_soft"])],foreground=[("active",colors["blue_dark"])])
        style.configure("Filter.TRadiobutton",background=colors["panel"],foreground=colors["ink"],
                        font=("Segoe UI",11,"bold"),padding=(8,9))
        style.map("Filter.TRadiobutton",background=[("selected",colors["blue_soft"]),("active","#273244")],
                  foreground=[("selected",colors["blue_dark"])])
        style.configure("Treeview",font=("Segoe UI",10),rowheight=58,background=colors["white"],
                        fieldbackground=colors["white"],foreground=colors["ink"],borderwidth=0)
        style.configure("Treeview.Heading",font=("Segoe UI",9,"bold"),padding=(10,11),
                        background="#242e3d",foreground=colors["muted"],relief="flat")
        style.map("Treeview",background=[("selected",colors["blue_soft"])],foreground=[("selected",colors["ink"])])
        style.configure("Status.Horizontal.TProgressbar",background=colors["blue"],troughcolor=colors["panel"],borderwidth=0)

        header=tk.Frame(self,bg=colors["panel"],highlightbackground=colors["line"],highlightthickness=0)
        header.pack(fill="x")
        brand=tk.Frame(header,bg=colors["panel"]);brand.pack(side="left",padx=(18,20),pady=10)
        header_mascot=ui_photo("mascote_busca_ui.png",88,88)
        if header_mascot:
            tk.Label(brand,image=header_mascot,bg=colors["panel"],borderwidth=0).pack(side="left",padx=(0,10))
        brand_text=tk.Frame(brand,bg=colors["panel"]);brand_text.pack(side="left",anchor="center")
        brand_line=tk.Frame(brand_text,bg=colors["panel"]);brand_line.pack(anchor="w")
        tk.Label(brand_line,text="Tô no Corre",font=("Segoe UI",23,"bold"),bg=colors["panel"],fg=colors["ink"]).pack(side="left")
        tk.Label(brand_line,text="BETA",font=("Segoe UI",8,"bold"),bg=colors["blue_soft"],fg=colors["blue_dark"],
                 padx=7,pady=3).pack(side="left",padx=(9,0),pady=(5,0))
        tk.Label(brand_text,text="Vagas trabalhando por você.",font=("Segoe UI",10),bg=colors["panel"],fg=colors["muted"]).pack(anchor="w")
        nav=ttk.Frame(header,style="Panel.TFrame");nav.pack(side="right",padx=20,pady=16)
        self.search_button=ttk.Button(nav,text="Buscar vagas",style="Primary.TButton",command=lambda:self.start_source("all"))
        self.search_button.pack(side="left",padx=4)
        ttk.Button(nav,text="Limpar pesquisa",style="Soft.TButton",command=self.clear_search).pack(side="left",padx=4)
        ttk.Button(nav,text="Fila",style="Soft.TButton",command=self.show_batch).pack(side="left",padx=4)
        ttk.Button(nav,text="Candidatar",style="Soft.TButton",command=self.start_batch).pack(side="left",padx=4)
        ttk.Button(nav,text="Minhas candidaturas",command=self.show_applications).pack(side="left",padx=4)
        ttk.Button(nav,text="Configurações",command=self.edit_profile).pack(side="left",padx=4)
        ttk.Button(nav,text="Créditos",command=self.show_credits).pack(side="left",padx=4)

        workspace=ttk.Frame(self,padding=(18,16,18,8));workspace.pack(fill="both",expand=True)
        sidebar=ttk.Frame(workspace,style="Panel.TFrame",padding=(14,16));sidebar.pack(side="left",fill="y",padx=(0,14))
        sidebar_controls=ttk.Frame(sidebar,style="Panel.TFrame")
        sidebar_controls.pack(side="top",fill="x")
        ttk.Label(sidebar_controls,text="MOSTRAR VAGAS",style="Muted.Panel.TLabel",font=("Segoe UI",9,"bold")).pack(anchor="w",padx=6,pady=(0,8))
        self.view_mode=tk.StringVar(value="todas")
        for text,value in [("Todas","todas"),("Presencial","presencial"),("Remoto","remoto"),
                           ("Estágios","estagio"),("Jovem Aprendiz","aprendiz"),("Vagas PCD","pcd")]:
            ttk.Radiobutton(sidebar_controls,text=text,value=value,variable=self.view_mode,command=self.refresh,
                            style="Filter.TRadiobutton",width=17).pack(fill="x",pady=2)
        ttk.Separator(sidebar_controls).pack(fill="x",pady=14)
        ttk.Label(sidebar_controls,text="RESUMO",style="Muted.Panel.TLabel",font=("Segoe UI",9,"bold")).pack(anchor="w",padx=6,pady=(0,7))
        self.stat_rec=tk.StringVar(value="Recomendadas: 0")
        self.stat_rev=tk.StringVar(value="Vale conferir: 0")
        self.stat_app=tk.StringVar(value="Candidaturas: 0")
        self.stat_out=tk.StringVar(value="Fora do perfil: 0")
        self.stat_model_pending=tk.StringVar(value="Modelo a confirmar: 0")
        self.stat_location_excluded=tk.StringVar(value="Localidades excluídas: 0")
        self.stat_discarded=tk.StringVar(value="Descartadas: 0")
        for var,command in [(self.stat_rec,lambda:self.set_view("recomendadas")),
                            (self.stat_rev,lambda:self.set_view("revisar")),
                            (self.stat_app,self.show_applications),(self.stat_out,self.show_discarded),
                            (self.stat_model_pending,self.show_model_unconfirmed),
                            (self.stat_location_excluded,self.show_location_excluded),
                            (self.stat_discarded,self.show_manually_discarded)]:
            ttk.Button(sidebar_controls,textvariable=var,style="Summary.TButton",command=command).pack(fill="x",pady=1)

        sidebar_mascot_large=ui_photo("mascote_sidebar_large_ui.png",190,190)
        sidebar_mascot=ui_photo("mascote_sidebar_ui.png",168,168)
        sidebar_mascot_medium=ui_photo("mascote_sidebar_medium_ui.png",112,112)
        sidebar_mascot_small=ui_photo("mascote_sidebar_small_ui.png",80,80)
        sidebar_mascot_label=tk.Label(sidebar,bg=colors["panel"],borderwidth=0)
        sidebar_mascot_label.pack(side="bottom")

        sidebar_mascot_resize_job=None
        sidebar_mascot_current=None

        def apply_sidebar_mascot_size():
            nonlocal sidebar_mascot_resize_job,sidebar_mascot_current
            sidebar_mascot_resize_job=None
            available=sidebar.winfo_height()-sidebar_controls.winfo_reqheight()-4
            choices=((190,sidebar_mascot_large),(168,sidebar_mascot),
                     (112,sidebar_mascot_medium),(80,sidebar_mascot_small))
            selected=next((image for size,image in choices if image and available>=size),None)
            if selected is sidebar_mascot_current:return
            sidebar_mascot_current=selected
            sidebar_mascot_label.configure(image=selected or "")
            sidebar_mascot_label.pack_configure(pady=(4,0) if selected else 0)

        def resize_sidebar_mascot(event):
            nonlocal sidebar_mascot_resize_job
            if event.widget is not self:return
            if sidebar_mascot_resize_job:
                try:self.after_cancel(sidebar_mascot_resize_job)
                except tk.TclError:pass
            # Maximizar/restaurar emite vários eventos antes que os painéis
            # internos recebam sua geometria definitiva.
            sidebar_mascot_resize_job=self.after(120,apply_sidebar_mascot_size)

        self.bind("<Configure>",resize_sidebar_mascot,add="+")
        self.after_idle(apply_sidebar_mascot_size)
        main=ttk.Frame(workspace);main.pack(side="left",fill="both",expand=True)
        tools=ttk.Frame(main);tools.pack(fill="x",pady=(0,10))
        titlebox=ttk.Frame(tools);titlebox.pack(side="left")
        ttk.Label(titlebox,text="Vagas para você",font=("Segoe UI",16,"bold")).pack(anchor="w")
        ttk.Label(titlebox,text="Escolha uma vaga para ver os detalhes.",foreground=colors["muted"]).pack(anchor="w")
        self.search_activity=tk.Frame(tools,bg=colors["blue_soft"],highlightthickness=0)
        self.search_spinner=tk.StringVar(value="◐")
        tk.Label(self.search_activity,textvariable=self.search_spinner,font=("Segoe UI",16,"bold"),
                 bg=colors["blue_soft"],fg=colors["blue_dark"]).pack(side="left",padx=(10,5),pady=7)
        tk.Label(self.search_activity,text="Buscando vagas para você…",font=("Segoe UI",10,"bold"),
                 bg=colors["blue_soft"],fg=colors["blue_dark"]).pack(side="left",padx=(0,10),pady=7)
        self.search_animation_job=None;self.search_animation_frame=0
        self.q=tk.StringVar()
        # Refina localmente qualquer aba sem repetir a consulta às fontes.
        # O pequeno atraso evita reconstruir uma lista grande mais de uma vez
        # quando várias teclas chegam juntas.
        self.filter_refresh_job=None
        def schedule_filter_refresh(*_args):
            if self.filter_refresh_job:
                try:self.after_cancel(self.filter_refresh_job)
                except tk.TclError:pass
            self.filter_refresh_job=self.after(180,self.refresh)
        self.q.trace_add("write",schedule_filter_refresh)
        self.batch_selection={r[0] for r in self.conn.execute(
            "SELECT id FROM vagas WHERE selecionada_lote=1 AND status='Nova'").fetchall()}
        self.queue_action_text=tk.StringVar(value="Incluir na fila")
        ttk.Button(tools,textvariable=self.queue_action_text,style="Primary.TButton",
                   command=self.apply_batch_selection).pack(side="right",padx=(0,10))
        filter_box=ttk.Frame(tools)
        filter_box.pack(side="right",padx=(12,14))
        ttk.Label(filter_box,text="Filtrar vagas:").pack(side="left",padx=(0,6))
        filter_entry=ttk.Entry(filter_box,textvariable=self.q,width=32)
        filter_entry.pack(side="left")
        filter_entry.bind("<Escape>",lambda _event:self.q.set(""))
        self.update_queue_action()

        self.empty_state=tk.Frame(main,bg=colors["panel"],highlightbackground=colors["line"],highlightthickness=1)
        empty_text=tk.Frame(self.empty_state,bg=colors["panel"]);empty_text.pack(side="left",padx=20,pady=15)
        tk.Label(empty_text,text="Comece pelo seu currículo",font=("Segoe UI",13,"bold"),
                 bg=colors["panel"],fg=colors["ink"]).pack(anchor="w")
        tk.Label(empty_text,text="Carregue um PDF, DOCX ou TXT. Os dados permanecem neste computador.",
                 font=("Segoe UI",10),bg=colors["panel"],fg=colors["muted"]).pack(anchor="w",pady=(3,0))
        ttk.Button(self.empty_state,text="Carregar currículo",style="Primary.TButton",
                   command=self.edit_profile).pack(side="left",padx=(0,18),pady=13)
        initial_mascots=ui_photo("mascotes_iniciais_ui.png",310,205)
        if initial_mascots:
            tk.Label(self.empty_state,image=initial_mascots,bg=colors["panel"],
                     borderwidth=0).pack(side="right",padx=12,pady=8)

        pane=ttk.Panedwindow(main,orient="horizontal");pane.pack(fill="both",expand=True)
        self.main_pane=pane
        if not self.cv.strip():self.empty_state.pack(fill="x",pady=(0,10),before=pane)
        left=ttk.Frame(pane,style="Panel.TFrame",padding=1)
        right=ttk.Frame(pane,style="Panel.TFrame",padding=(16,14))
        pane.add(left,weight=3);pane.add(right,weight=2)

        cols=("batch","score","titulo","empresa","local","data")
        self.tree=ttk.Treeview(left,columns=cols,show="headings",selectmode="browse")
        self.job_heading_labels={"batch":"FILA","score":"COMPAT.","titulo":"VAGA","empresa":"EMPRESA","local":"LOCAL / MODELO","data":"PUBLICAÇÃO"}
        self.job_sort_state={"column":"","reverse":False}
        for col,title,width,anchor in [("batch","FILA",58,"center"),("score","COMPAT.",88,"center"),("titulo","VAGA",275,"w"),
                                       ("empresa","EMPRESA",165,"w"),("local","LOCAL / MODELO",215,"w"),("data","PUBLICAÇÃO",105,"center")]:
            self.tree.heading(col,text=title,command=lambda c=col:self.sort_jobs(c))
            self.tree.column(col,width=width,anchor=anchor,minwidth=50 if col=="batch" else 70)
        tree_scroll=ttk.Scrollbar(left,orient="vertical",command=self.tree.yview);self.tree.configure(yscrollcommand=tree_scroll.set)
        tree_scroll.pack(side="right",fill="y");self.tree.pack(side="left",fill="both",expand=True)
        self.tree.tag_configure("great",foreground="#74d3a7")
        self.tree.tag_configure("possible",foreground=colors["ink"])
        self.tree.bind("<Button-1>",self.toggle_batch_checkbox,add="+")
        self.tree.bind("<<TreeviewSelect>>",self.select);self.tree.bind("<Double-1>",lambda e:self.open_job())

        detail_head=ttk.Frame(right,style="Panel.TFrame");detail_head.pack(fill="x")
        detail_head.columnconfigure(0,weight=1);detail_head.columnconfigure(1,weight=1)
        detail_caption=ttk.Label(detail_head,text="DETALHES DA VAGA",style="Muted.Panel.TLabel",
                                 font=("Segoe UI",9,"bold"))
        detail_caption.grid(row=0,column=0,columnspan=2,sticky="w",pady=(0,5))
        open_job_button=ttk.Button(detail_head,text="Ver vaga",command=self.open_job)
        discard_job_button=ttk.Button(detail_head,text="Descartar",style="Danger.TButton",
                                      command=self.discard_current)
        open_job_button.grid(row=1,column=0,sticky="ew",padx=(0,3))
        discard_job_button.grid(row=1,column=1,sticky="ew",padx=(3,0))
        self.tv=tk.StringVar(value="Selecione uma vaga")
        detail_title=ttk.Label(right,textvariable=self.tv,style="Panel.TLabel",font=("Segoe UI",15,"bold"),wraplength=480)
        detail_title.pack(anchor="w",pady=(8,3))
        self.meta=tk.StringVar(value="Escolha uma vaga na lista ao lado.")
        detail_meta=ttk.Label(right,textvariable=self.meta,style="Muted.Panel.TLabel",wraplength=480)
        detail_meta.pack(anchor="w")
        self.data_quality=tk.StringVar(value="")
        detail_quality=ttk.Label(right,textvariable=self.data_quality,style="Muted.Panel.TLabel",wraplength=480)
        detail_quality.pack(anchor="w",pady=(2,12))

        detail_layout={"width":None,"stacked":None}
        def resize_detail_panel(event):
            # No Windows, <Configure> também pode ser repetido enquanto a janela
            # é arrastada, mesmo sem alteração da largura do painel. Reaplicar
            # wrap e grid em cada evento provoca relayouts visíveis e movimento
            # engasgado. Ignora eventos geometricamente idênticos.
            width=max(1,int(event.width))
            if detail_layout["width"]==width:return
            detail_layout["width"]=width
            available=max(80,width-32)
            detail_title.configure(wraplength=available)
            detail_meta.configure(wraplength=available)
            detail_quality.configure(wraplength=available)
            stacked=width<245
            if detail_layout["stacked"]==stacked:return
            detail_layout["stacked"]=stacked
            if stacked:
                open_job_button.grid_configure(row=1,column=0,columnspan=2,padx=0,pady=(0,4))
                discard_job_button.grid_configure(row=2,column=0,columnspan=2,padx=0)
            else:
                open_job_button.grid_configure(row=1,column=0,columnspan=1,padx=(0,3),pady=0)
                discard_job_button.grid_configure(row=1,column=1,columnspan=1,padx=(3,0))
        right.bind("<Configure>",resize_detail_panel,add="+")

        def detail_section(label,height,expand=False):
            ttk.Label(right,text=label,style="Panel.TLabel",font=("Segoe UI",10,"bold")).pack(anchor="w",pady=(3,4))
            box=tk.Text(right,height=height,wrap="word",padx=10,pady=9,font=("Segoe UI",10),
                        bg=colors["white"],fg=colors["ink"],relief="flat",highlightthickness=1,
                        highlightbackground=colors["line"],selectbackground=colors["blue_soft"])
            box.pack(fill="both" if expand else "x",expand=expand,pady=(0,9));box.configure(state="disabled")
            return box
        self.desc_box=detail_section("Descrição",9,True)
        self.req_box=detail_section("Requisitos",5)
        self.pay_box=detail_section("Salário e benefícios",4)

        status=tk.Frame(self,bg=colors["panel"],highlightbackground=colors["line"],highlightthickness=1)
        status.pack(fill="x",side="bottom")
        self.info=tk.StringVar(value="Pronto para buscar vagas")
        tk.Label(status,textvariable=self.info,font=("Segoe UI",9),bg=colors["panel"],fg=colors["muted"]).pack(side="left",padx=18,pady=7)
        tk.Label(status,text=f"{APP_TITLE} • {APP_VERSION}",font=("Segoe UI",9),
                 bg=colors["panel"],fg=colors["muted"]).pack(side="right",padx=(4,18),pady=7)
        tk.Label(status,text="Dados armazenados localmente",font=("Segoe UI",9),
                 bg=colors["panel"],fg=colors["muted"]).pack(side="right",padx=(8,10),pady=7)
        self.prog=ttk.Progressbar(status,mode="indeterminate",length=120,style="Status.Horizontal.TProgressbar")
        self.prog.pack(side="right",padx=(18,4),pady=9)

    def set_view(self,view):
        self.view_mode.set(view);self.refresh()

    def clear_search(self):
        total=self.conn.execute("""SELECT COUNT(*) FROM vagas
            WHERE status='Nova' AND COALESCE(selecionada_lote,0)=0""").fetchone()[0]
        if not total:
            self.info.set("A pesquisa já está vazia. A fila e o histórico foram mantidos.");return 0
        if not messagebox.askyesno("Limpar pesquisa",
            f"Retirar {total} vaga(s) da pesquisa atual e começar novamente do zero?\n\n"
            "Fila, candidaturas e vagas fora do perfil serão mantidas."):
            return 0
        with self.conn:
            self.conn.execute("""UPDATE vagas SET status='Pesquisa limpa',selecionada_lote=0
                                 WHERE status='Nova' AND COALESCE(selecionada_lote,0)=0""")
        self.current=None;self.q.set("");self.view_mode.set("todas");self.refresh()
        self.tv.set("Selecione uma vaga");self.meta.set("Faça uma nova busca para carregar vagas.");self.data_quality.set("")
        for box in (self.desc_box,self.req_box,self.pay_box):
            box.configure(state="normal");box.delete("1.0","end");box.configure(state="disabled")
        self.info.set(f"Pesquisa limpa: {total} vaga(s) arquivada(s). Você já pode buscar novamente.")
        return total

    def sort_jobs(self,column,reverse=None):
        if reverse is None:
            reverse=(not self.job_sort_state["reverse"]) if self.job_sort_state["column"]==column else False
        self.job_sort_state={"column":column,"reverse":reverse}
        index=list(self.tree["columns"]).index(column)
        def key(iid):
            value=self.tree.item(iid,"values")[index]
            if column=="score":
                try:return int(str(value).replace("%","").strip())
                except ValueError:return -1
            if column=="batch":return 1 if str(value)=="☑" else 0
            if column=="data":
                try:return datetime.strptime(str(value),"%d/%m/%Y")
                except ValueError:return datetime.min
            return norm(value)
        items=list(self.tree.get_children(""));items.sort(key=key,reverse=reverse)
        for position,iid in enumerate(items):self.tree.move(iid,"",position)
        for col,label in self.job_heading_labels.items():
            marker=(" ▼" if reverse else " ▲") if col==column else ""
            self.tree.heading(col,text=label+marker,command=lambda c=col:self.sort_jobs(c))

    def managed_window(self,key,title,geometry,modal=False):
        current=self.open_windows.get(key)
        if current is not None:
            try:
                if current.winfo_exists():
                    current.deiconify();current.lift();current.focus_force()
                    return current,False
            except tk.TclError:pass
        win=tk.Toplevel(self);win.title(title);win.transient(self)
        win.configure(bg=getattr(self,"colors",{}).get("bg","#141922"))
        width,height=(int(x) for x in geometry.lower().split("x")[:2])
        self.center_window(win,width,height)
        self.open_windows[key]=win
        if modal:win.grab_set()
        def close():
            self.open_windows.pop(key,None)
            try:
                if modal:win.grab_release()
            except tk.TclError:pass
            win.destroy()
        win.protocol("WM_DELETE_WINDOW",close);win._managed_close=close
        return win,True

    def center_window(self,win,width,height):
        self.update_idletasks()
        x=self.winfo_rootx()+max(0,(self.winfo_width()-width)//2)
        y=self.winfo_rooty()+max(0,(self.winfo_height()-height)//2)
        x=max(0,min(x,win.winfo_screenwidth()-width))
        y=max(0,min(y,win.winfo_screenheight()-height))
        win.geometry(f"{width}x{height}+{x}+{y}")

    def update_queue_action(self):
        if "queue_action_text" not in self.__dict__:return
        count=len(self.batch_selection)
        self.queue_action_text.set(f"Incluir na fila ({count})" if count else "Incluir na fila")

    def toggle_batch_checkbox(self,event):
        if self.tree.identify_region(event.x,event.y)!="cell" or self.tree.identify_column(event.x)!="#1":return
        iid=self.tree.identify_row(event.y)
        if not iid:return "break"
        vid=int(iid)
        if vid in self.batch_selection:self.batch_selection.remove(vid)
        else:self.batch_selection.add(vid)
        values=list(self.tree.item(iid,"values"));values[0]="☑" if vid in self.batch_selection else "☐"
        self.tree.item(iid,values=values)
        self.update_queue_action()
        return "break"

    def apply_batch_selection(self):
        selected=sorted(self.batch_selection)
        with self.conn:
            self.conn.execute("UPDATE vagas SET selecionada_lote=0")
            self.conn.executemany("UPDATE vagas SET selecionada_lote=1 WHERE id=? AND status='Nova'",
                                  ((vid,) for vid in selected))
        self.batch_selection={r[0] for r in self.conn.execute(
            "SELECT id FROM vagas WHERE selecionada_lote=1 AND status='Nova'").fetchall()}
        total=len(self.batch_selection);self.update_queue_action()
        for vid in selected:self.record_feedback(vid,True)
        self.info.set(f"Fila atualizada: {total} vaga(s).")
        self.refresh()

    def sync_batch_selection(self):
        self.batch_selection={r[0] for r in self.conn.execute(
            "SELECT id FROM vagas WHERE selecionada_lote=1 AND status='Nova'").fetchall()}
        self.update_queue_action();self.refresh()

    def clear_batch(self,ask=True):
        total=self.conn.execute("SELECT COUNT(*) FROM vagas WHERE selecionada_lote=1").fetchone()[0]
        if not total:
            self.info.set("A fila já está vazia.")
            return 0
        if ask and not messagebox.askyesno("Limpar fila",f"Remover {total} vaga(s) da fila?\n\nAs vagas não serão excluídas."):
            return 0
        with self.conn:self.conn.execute("UPDATE vagas SET selecionada_lote=0")
        self.sync_batch_selection();self.info.set("Fila limpa. As vagas foram mantidas.")
        return total

    def remove_from_batch(self,vid):
        with self.conn:self.conn.execute("UPDATE vagas SET selecionada_lote=0 WHERE id=?",(vid,))
        if "batch_selection" in self.__dict__:self.batch_selection.discard(vid)
        if "queue_action_text" in self.__dict__:self.update_queue_action()
        self.refresh();self.info.set("Vaga removida da fila. Ela continua na lista de vagas.")
        return True

    def mark_application_completed(self,vid):
        with self.conn:
            self.conn.execute("""UPDATE vagas SET status='Candidatado',selecionada_lote=0,
                                 ultimo_resultado='Candidatura confirmada pelo usuário',
                                 candidatura_em=?
                                 WHERE id=?""",(datetime.now().isoformat(timespec="seconds"),vid))

    def start_source(self,src):
        if self.search_running:
            self.info.set("A busca atual ainda está em andamento.")
            return
        self.cv=read_cv()
        if not self.cv.strip():
            messagebox.showinfo("Adicione seu currículo","Carregue seu currículo para o aplicativo preparar uma busca compatível com o seu perfil.")
            self.edit_profile();return
        self.p=load_profile()
        # Reconstrói as consultas a cada busca para aplicar imediatamente as
        # preferências atuais e melhorias do mecanismo, sem exigir novo upload.
        adapt_profile_to_cv(self.p,self.cv);save_json_file(PROFILE_PATH,self.p)
        self.reactivate_searchable_jobs()
        self.apply_publication_age_preference()
        self.separate_outside_region_jobs()
        self.apply_pcd_preference(pcd_vacancies_enabled(self.p))
        self.apply_internship_mode(internship_search_mode(self.p))
        self.apply_apprentice_preference(bool(self.p.get("buscar_jovem_aprendiz",False)))
        self.apply_international_preference(international_search_enabled(self.p))
        self.apply_excluded_locations()
        self.search_running=True
        self.last_source_counts={}
        self.search_button.configure(state="disabled")
        self.start_search_animation()
        self.info.set(f"Buscando {src}...");self.prog.start(10)
        self.shutdown_event.clear()
        self.start_worker(self.run_source,src,name="busca-vagas")

    def collect(self,src):
        entry_mode=bool(self.p.get("buscar_vagas_inicio_carreira",False))
        internship_only=internship_search_mode(self.p)=="somente"
        internship_mode=internship_search_mode(self.p)
        google_limit=32 if internship_only else (20 if entry_mode else 12)
        google_jobs=120 if internship_only else (100 if entry_mode else 60)
        def scoped_profile(kind):
            profile=dict(self.p)
            suffix="estagio" if kind=="estagios" else "vagas"
            wants_internships=(kind=="estagios")
            def scoped_queries(provider):
                key=f"consultas_{suffix}_{provider}"
                if key in self.p:return list(self.p.get(key,[]))
                legacy_key="consultas_"+provider
                return [query for query in self.p.get(legacy_key,[])
                        if is_internship_query(query)==wants_internships]
            profile["consultas_gupy"]=scoped_queries("gupy")
            profile["consultas_linkedin"]=scoped_queries("linkedin")
            profile["consultas_google"]=scoped_queries("google")
            profile["consultas_br"]=list(profile["consultas_gupy"])
            profile["modo_estagios"]="somente" if kind=="estagios" else "nao_buscar"
            profile["buscar_estagios"]=(kind=="estagios")
            return profile
        regular_profile=scoped_profile("vagas")
        internship_profile=scoped_profile("estagios")
        scopes=([("Estágios",internship_profile)] if internship_mode=="somente" else
                [("Vagas",regular_profile),("Estágios",internship_profile)] if internship_mode=="incluir" else
                [("Vagas",regular_profile)])
        scopes=[item for item in scopes if item[0]=="Vagas" or any(item[1].get(key) for key in
                ("consultas_gupy","consultas_linkedin","consultas_google"))]
        def cached(name,profile,fn,hours=1,limit=600):
            return persistent_source_fetch(name,profile,fn,cooldown_seconds=hours*3600,max_cached=limit)
        def collect_scoped(provider,fetcher,hours=1,limit=600):
            jobs=[]
            for label,profile in scopes:
                jobs += cached(f"{provider} / {label}",profile,lambda p=profile:fetcher(p),hours,limit)
            return dedupe_multisource(jobs,self.p)
        if src=="gupy":return collect_scoped("Gupy",fetch_gupy,1)
        if src=="linkedin":return collect_scoped("LinkedIn",fetch_linkedin,6)
        if src=="google":return collect_scoped("Google",lambda p:fetch_google(
            p,max_queries=32 if internship_search_mode(p)=="somente" else google_limit,
            max_jobs=120 if internship_search_mode(p)=="somente" else google_jobs),3,300)
        if src=="indeed":return collect_scoped("Indeed via Google",lambda p:fetch_google(
            p,"br.indeed.com","Indeed/Google",max_queries=32 if internship_search_mode(p)=="somente" else google_limit,
            max_jobs=120 if internship_search_mode(p)=="somente" else google_jobs),3,300)
        if src=="remote":return dedupe_multisource(fetch_remotive(self.p),self.p)
        if src=="all":
            jobs=[]
            sources={}
            for label,profile in scopes:
                per_google_limit=32 if label=="Estágios" else google_limit
                per_google_jobs=120 if label=="Estágios" else google_jobs
                sources.update({
                f"Gupy / {label}":lambda p=profile,l=label:cached(
                    f"Gupy / {l}",p,lambda:fetch_gupy(p),1),
                f"LinkedIn / {label}":lambda p=profile,l=label:cached(
                    f"LinkedIn / {l}",p,lambda:fetch_linkedin(p),6),
                f"Google / {label}":lambda p=profile,l=label,g=per_google_limit,j=per_google_jobs:cached(
                    f"Google / {l}",p,lambda:fetch_google(p,max_queries=g,max_jobs=j),3,300),
                f"Indeed via Google / {label}":lambda p=profile,l=label,g=per_google_limit,j=per_google_jobs:cached(
                    f"Indeed via Google / {l}",p,lambda:fetch_google(
                        p,"br.indeed.com","Indeed/Google",max_queries=g,max_jobs=j),3,300),
                })
            sources.update({
                "Vagas em Games via Google":lambda:cached("Vagas em Games via Google",self.p,lambda:fetch_google(
                    self.p,"vagasemgames.com.br/vagas","Vagas em Games/Google",max_queries=4,max_results=5,max_jobs=20),3,120)
            })
            if internship_mode!="nao_buscar" and internship_profile.get("consultas_google"):
                sources["Portais de estágio"]=lambda:cached(
                    "Portais de estágio",internship_profile,
                    lambda:fetch_internship_portals(internship_profile),6,300)
            if international_search_enabled(self.p):
                sources.update({
                "Remotive":lambda:cached("Remotive",self.p,lambda:fetch_remotive(self.p),3,300),
                "Jobicy":lambda:cached("Jobicy",self.p,lambda:fetch_jobicy(self.p),3,300),
                "Himalayas":lambda:cached("Himalayas",self.p,lambda:fetch_himalayas(self.p),3,300),
                "Remote Landers":lambda:cached("Remote Landers",self.p,lambda:fetch_remote_landers(self.p),3,300),
                "Remote Game Jobs":lambda:cached("Remote Game Jobs",self.p,lambda:fetch_remote_game_jobs(self.p),3,300),
                "Work With Indies":lambda:cached("Work With Indies",self.p,lambda:fetch_work_with_indies(self.p),3,300),
                "Hitmarker via Google":lambda:cached("Hitmarker via Google",self.p,lambda:fetch_google(
                    self.p,"hitmarker.net/jobs","Hitmarker/Google",max_queries=4,max_results=5,max_jobs=20),3,120)
                })
            source_count=len(sources)
            source_counts=self.__dict__.setdefault("last_source_counts",{})
            self.ui_call(lambda count=source_count:self.info.set(f"Consultando {count} fontes em paralelo..."))
            with ThreadPoolExecutor(max_workers=source_count,thread_name_prefix="busca") as executor:
                futures={executor.submit(fn):name for name,fn in sources.items()}
                for future in as_completed(futures):
                    if self.shutdown_requested():
                        for pending in futures:pending.cancel()
                        break
                    name=futures[future]
                    try:
                        found=future.result() or [];jobs+=found
                        source_counts[name]=len(found)
                        self.ui_call(lambda n=name,total=len(found):self.info.set(f"{n} concluída: {total} vaga(s)"))
                    except Exception:LOGGER.exception("Falha na fonte %s",name)
            return dedupe_multisource(jobs,self.p)
        return []

    def save_discarded(self,j,reason):
        vals=(j.get("titulo",""),j.get("empresa",""),j.get("local",""),j.get("descricao",""),
              j.get("url",""),j.get("fonte",""),j.get("data_publicacao",""),j.get("salario",""),reason,
              j.get("workplace_type","") or "",compact_json(j.get("applicant_location_requirements")),
              compact_json(j.get("structured_location_json") or j.get("structured_location")))
        try:
            self.conn.execute("""INSERT INTO descartadas
                (titulo,empresa,local,descricao,url,fonte,data_publicacao,salario,motivo_descarte,
                 workplace_type,applicant_location_requirements,structured_location_json)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",vals)
        except sqlite3.IntegrityError:
            self.conn.execute("""UPDATE descartadas SET titulo=?,empresa=?,local=?,descricao=?,fonte=?,
                data_publicacao=?,salario=?,motivo_descarte=CASE
                    WHEN motivo_descarte='Descartada pelo usuário' THEN motivo_descarte ELSE ? END,
                workplace_type=?,applicant_location_requirements=?,structured_location_json=?,
                descartada_em=CURRENT_TIMESTAMP WHERE url=?""",
                (j.get("titulo",""),j.get("empresa",""),j.get("local",""),j.get("descricao",""),
                 j.get("fonte",""),j.get("data_publicacao",""),j.get("salario",""),reason,
                 j.get("workplace_type","") or "",compact_json(j.get("applicant_location_requirements")),
                 compact_json(j.get("structured_location_json") or j.get("structured_location")),j.get("url","")))

    def restore_discarded_record(self,did,refresh=True):
        row=self.conn.execute("""SELECT titulo,empresa,local,descricao,url,fonte,data_publicacao,salario,
                                  COALESCE(workplace_type,''),COALESCE(applicant_location_requirements,''),
                                  COALESCE(structured_location_json,'')
                                 FROM descartadas WHERE id=?""",(did,)).fetchone()
        if not row:return False
        t,e,l,d,u,f,dt,sal,workplace,requirements,structured=row
        job={"titulo":t or "","empresa":e or "","local":l or "","descricao":d or "",
             "url":u or "","fonte":f or "","data_publicacao":dt or "","salario":sal or "",
             "workplace_type":workplace or "","applicant_location_requirements":requirements or "",
             "structured_location_json":structured or "",
             "source_brazil":f in ("Gupy","LinkedIn","Google","Indeed/Google")}
        score,label,reason,mode=score_job(job,load_profile(),read_cv())
        with self.conn:
            existing=self.conn.execute("SELECT id FROM vagas WHERE url=?",(u,)).fetchone()
            if existing:
                self.conn.execute("""UPDATE vagas SET status='Nova',selecionada_lote=0,
                                     ultimo_resultado='' WHERE id=?""",(existing[0],))
            else:
                self.conn.execute("""INSERT INTO vagas(titulo,empresa,local,modalidade,descricao,url,fonte,
                    data_publicacao,salario,score,classificacao,motivo,status,workplace_type,
                    applicant_location_requirements,structured_location_json)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?, 'Nova',?,?,?)""",
                    (t,e,l,mode,d,u,f,dt,sal,score,label,reason,workplace,requirements,structured))
            self.conn.execute("DELETE FROM descartadas WHERE id=?",(did,))
        if refresh:self.refresh()
        return True

    def reactivate_searchable_jobs(self):
        """Toda nova pesquisa reapresenta vagas salvas, exceto descartadas pelo usuário."""
        with self.conn:
            cursor=self.conn.execute("""UPDATE vagas SET status='Nova',selecionada_lote=0,
                ultimo_resultado=CASE WHEN status='Arquivada' THEN '' ELSE ultimo_resultado END
                WHERE status IN ('Pesquisa limpa','Arquivada')""")
        restored=max(0,cursor.rowcount);self.refresh()
        return restored

    def separate_outside_region_jobs(self):
        """Move vagas locais externas para Fora do perfil, preservando fila e histórico."""
        rows=self.conn.execute("""SELECT id,titulo,empresa,local,descricao,url,fonte,data_publicacao,salario,
                                  COALESCE(workplace_type,'') FROM vagas
                                  WHERE status='Nova' AND COALESCE(selecionada_lote,0)=0""").fetchall()
        moved=0
        for vid,titulo,empresa,local,descricao,url,fonte,pub,salario,workplace in rows:
            job={"titulo":titulo or "","empresa":empresa or "","local":local or "",
                 "descricao":descricao or "","url":url or "","fonte":fonte or "",
                 "data_publicacao":pub or "","salario":salario or "","workplace_type":workplace or ""}
            ld=location_decision(job,self.p)
            if ld["ok"] or "fora da regiao" not in norm(ld["mode"]):continue
            self.save_discarded(job,ld["mode"]);moved+=1
            self.conn.execute("UPDATE vagas SET status='Ignorada',selecionada_lote=0,decisao='REVISAR' WHERE id=?",(vid,))
        self.conn.commit()
        if moved:self.refresh()
        return moved

    def apply_internship_preference(self,enabled):
        """Separa estágios quando desativados e restaura apenas os filtrados por essa preferência."""
        if enabled:
            ids=[r[0] for r in self.conn.execute(
                """SELECT id FROM descartadas
                   WHERE motivo_descarte='estágio desativado pelo usuário'
                      OR motivo_descarte LIKE 'estágio fora da formação identificada%'""")]
            restored=sum(1 for did in ids if self.restore_discarded_record(did,refresh=False))
            rows=self.conn.execute("""SELECT id,titulo,empresa,local,descricao,url,fonte,data_publicacao,salario
                                      FROM vagas WHERE status='Nova' AND COALESCE(selecionada_lote,0)=0""").fetchall()
            moved=0
            for vid,titulo,empresa,local,descricao,url,fonte,pub,salario in rows:
                job={"titulo":titulo or "","empresa":empresa or "","local":local or "",
                     "descricao":descricao or "","url":url or "","fonte":fonte or "",
                     "data_publicacao":pub or "","salario":salario or ""}
                if not is_intern(job) or internship_course_status(job,self.p)!="FORA":continue
                self.save_discarded(job,"estágio fora da formação identificada — outro curso")
                self.conn.execute("UPDATE vagas SET status='Ignorada',selecionada_lote=0 WHERE id=?",(vid,));moved+=1
            self.conn.commit()
            if restored or moved:self.refresh()
            return moved
        rows=self.conn.execute("""SELECT id,titulo,empresa,local,descricao,url,fonte,data_publicacao,salario
                                  FROM vagas WHERE status='Nova' AND COALESCE(selecionada_lote,0)=0""").fetchall()
        moved=0
        for vid,titulo,empresa,local,descricao,url,fonte,pub,salario in rows:
            job={"titulo":titulo or "","empresa":empresa or "","local":local or "","descricao":descricao or "",
                 "url":url or "","fonte":fonte or "","data_publicacao":pub or "","salario":salario or ""}
            if not is_intern(job):continue
            self.save_discarded(job,"estágio desativado pelo usuário")
            self.conn.execute("UPDATE vagas SET status='Ignorada',selecionada_lote=0 WHERE id=?",(vid,));moved+=1
        self.conn.commit()
        if moved:self.refresh()
        return moved

    def apply_internship_mode(self,mode):
        mode=mode if mode in INTERNSHIP_MODE_LABELS else internship_search_mode(self.p)
        self.p["modo_estagios"]=mode
        self.p["buscar_estagios"]=mode!="nao_buscar"
        changed=self.apply_internship_preference(mode!="nao_buscar")
        reason="modo somente estágios"
        if mode!="somente":
            ids=[row[0] for row in self.conn.execute(
                "SELECT id FROM descartadas WHERE motivo_descarte=?",(reason,))]
            restored=sum(1 for did in ids if self.restore_discarded_record(did,refresh=False))
            if restored:self.refresh()
            return changed+restored
        rows=self.conn.execute("""SELECT id,titulo,empresa,local,descricao,url,fonte,data_publicacao,salario
                                  FROM vagas WHERE status='Nova' AND COALESCE(selecionada_lote,0)=0""").fetchall()
        moved=0
        for vid,titulo,empresa,local,descricao,url,fonte,pub,salario in rows:
            job={"titulo":titulo or "","empresa":empresa or "","local":local or "",
                 "descricao":descricao or "","url":url or "","fonte":fonte or "",
                 "data_publicacao":pub or "","salario":salario or ""}
            if is_intern(job):continue
            self.save_discarded(job,reason)
            self.conn.execute("UPDATE vagas SET status='Ignorada',selecionada_lote=0 WHERE id=?",(vid,));moved+=1
        self.conn.commit()
        if moved:self.refresh()
        return changed+moved

    def apply_international_preference(self,enabled):
        """Oculta resultados de fontes internacionais sem afetar vagas brasileiras."""
        reason="vagas internacionais desativadas pelo usuário"
        if enabled:
            ids=[row[0] for row in self.conn.execute(
                "SELECT id FROM descartadas WHERE motivo_descarte=?",(reason,))]
            restored=sum(1 for did in ids if self.restore_discarded_record(did,refresh=False))
            if restored:self.refresh()
            return restored
        international=("Remotive","Jobicy","Himalayas","Remote Landers","Remote Game Jobs",
                       "Work With Indies","Hitmarker/Google")
        placeholders=",".join("?" for _ in international)
        rows=self.conn.execute(f"""SELECT id,titulo,empresa,local,descricao,url,fonte,data_publicacao,salario
            FROM vagas WHERE status='Nova' AND COALESCE(selecionada_lote,0)=0
              AND fonte IN ({placeholders})""",international).fetchall()
        moved=0
        for vid,titulo,empresa,local,descricao,url,fonte,pub,salario in rows:
            job={"titulo":titulo or "","empresa":empresa or "","local":local or "","descricao":descricao or "",
                 "url":url or "","fonte":fonte or "","data_publicacao":pub or "","salario":salario or ""}
            self.save_discarded(job,reason)
            self.conn.execute("UPDATE vagas SET status='Ignorada',selecionada_lote=0 WHERE id=?",(vid,));moved+=1
        self.conn.commit()
        if moved:self.refresh()
        return moved

    def apply_apprentice_preference(self,enabled):
        """Separa aprendizagem quando desativada e permite restauração sem apagar vagas."""
        reason="Jovem Aprendiz desativado pelo usuário"
        if enabled:
            ids=[row[0] for row in self.conn.execute(
                "SELECT id FROM descartadas WHERE motivo_descarte=?",(reason,))]
            restored=sum(1 for did in ids if self.restore_discarded_record(did,refresh=False))
            if restored:self.refresh()
            return restored
        rows=self.conn.execute("""SELECT id,titulo,empresa,local,descricao,url,fonte,data_publicacao,salario
                                  FROM vagas WHERE status='Nova' AND COALESCE(selecionada_lote,0)=0""").fetchall()
        moved=0
        for vid,titulo,empresa,local,descricao,url,fonte,pub,salario in rows:
            job={"titulo":titulo or "","empresa":empresa or "","local":local or "","descricao":descricao or "",
                 "url":url or "","fonte":fonte or "","data_publicacao":pub or "","salario":salario or ""}
            if not is_apprentice(job):continue
            self.save_discarded(job,reason)
            self.conn.execute("UPDATE vagas SET status='Ignorada',selecionada_lote=0 WHERE id=?",(vid,));moved+=1
        self.conn.commit()
        if moved:self.refresh()
        return moved

    def apply_pcd_preference(self,enabled):
        """Separa vagas PCD quando desativadas e restaura somente as filtradas por essa opção."""
        reasons=("vaga exclusiva/afirmativa para PCD","vaga direcionada/exclusiva para PCD")
        if enabled:
            placeholders=",".join("?" for _ in reasons)
            ids=[row[0] for row in self.conn.execute(
                f"SELECT id FROM descartadas WHERE motivo_descarte IN ({placeholders})",reasons)]
            restored=sum(1 for did in ids if self.restore_discarded_record(did,refresh=False))
            if restored:self.refresh()
            return restored
        rows=self.conn.execute("""SELECT id,titulo,empresa,local,descricao,url,fonte,data_publicacao,salario
                                  FROM vagas WHERE status='Nova' AND COALESCE(selecionada_lote,0)=0""").fetchall()
        moved=0
        for vid,titulo,empresa,local,descricao,url,fonte,pub,salario in rows:
            job={"titulo":titulo or "","empresa":empresa or "","local":local or "",
                 "descricao":descricao or "","url":url or "","fonte":fonte or "",
                 "data_publicacao":pub or "","salario":salario or ""}
            reason=pcd_job_reason(job)
            if not reason:continue
            self.save_discarded(job,reason)
            self.conn.execute("UPDATE vagas SET status='Ignorada',selecionada_lote=0 WHERE id=?",(vid,));moved+=1
        self.conn.commit()
        if moved:self.refresh()
        return moved

    def apply_excluded_locations(self):
        """Move/restaura somente vagas afetadas pela lista de localidades excluídas."""
        prefix="localidade excluída pelo usuário — "
        rows=self.conn.execute("""SELECT id,titulo,empresa,local,descricao,url,fonte,data_publicacao,salario,
                                  COALESCE(applicant_location_requirements,''),COALESCE(structured_location_json,'')
                                  FROM descartadas WHERE motivo_descarte LIKE ?""",(prefix+"%",)).fetchall()
        restored=0
        for did,title,company,location,description,url,source,published,salary,requirements,structured in rows:
            job={"titulo":title or "","empresa":company or "","local":location or "",
                 "descricao":description or "","url":url or "","fonte":source or "",
                 "data_publicacao":published or "","salario":salary or "",
                 "applicant_location_requirements":requirements or "","structured_location_json":structured or ""}
            if not excluded_location_reason(job,self.p):
                restored+=bool(self.restore_discarded_record(did,refresh=False))

        active=self.conn.execute("""SELECT id,titulo,empresa,local,descricao,url,fonte,data_publicacao,salario,
                                     COALESCE(applicant_location_requirements,''),COALESCE(structured_location_json,'')
                                     FROM vagas WHERE status='Nova' AND COALESCE(selecionada_lote,0)=0""").fetchall()
        moved=0
        for vid,title,company,location,description,url,source,published,salary,requirements,structured in active:
            job={"titulo":title or "","empresa":company or "","local":location or "",
                 "descricao":description or "","url":url or "","fonte":source or "",
                 "data_publicacao":published or "","salario":salary or "",
                 "applicant_location_requirements":requirements or "","structured_location_json":structured or ""}
            excluded=excluded_location_reason(job,self.p)
            if not excluded:continue
            self.save_discarded(job,prefix+excluded)
            self.conn.execute("UPDATE vagas SET status='Ignorada',selecionada_lote=0 WHERE id=?",(vid,));moved+=1
        self.conn.commit()
        if restored or moved:self.refresh()
        return moved,restored

    def apply_publication_age_preference(self):
        """Move/restaura vagas datadas conforme o período, preservando fila e histórico."""
        prefix="vaga antiga ("
        rows=self.conn.execute("""SELECT id,titulo,empresa,local,descricao,url,fonte,data_publicacao,salario,
                                         COALESCE(workplace_type,''),
                                         COALESCE(applicant_location_requirements,''),
                                         COALESCE(structured_location_json,'')
                                  FROM descartadas WHERE motivo_descarte LIKE ?""",(prefix+"%",)).fetchall()
        restored=0
        for did,title,company,location,description,url,source,published,salary,workplace,requirements,structured in rows:
            job={"titulo":title or "","empresa":company or "","local":location or "",
                 "descricao":description or "","url":url or "","fonte":source or "",
                 "data_publicacao":published or "","salario":salary or "",
                 "workplace_type":workplace or "","applicant_location_requirements":requirements or "",
                 "structured_location_json":structured or ""}
            if vacancy_date_ok(job,self.p)[0]:
                restored+=bool(self.restore_discarded_record(did,refresh=False))

        active=self.conn.execute("""SELECT id,titulo,empresa,local,descricao,url,fonte,data_publicacao,salario,
                                       COALESCE(workplace_type,''),
                                       COALESCE(applicant_location_requirements,''),
                                       COALESCE(structured_location_json,'')
                                    FROM vagas WHERE status='Nova' AND COALESCE(selecionada_lote,0)=0""").fetchall()
        moved=0
        for vid,title,company,location,description,url,source,published,salary,workplace,requirements,structured in active:
            job={"titulo":title or "","empresa":company or "","local":location or "",
                 "descricao":description or "","url":url or "","fonte":source or "",
                 "data_publicacao":published or "","salario":salary or "",
                 "workplace_type":workplace or "","applicant_location_requirements":requirements or "",
                 "structured_location_json":structured or ""}
            allowed,age=vacancy_date_ok(job,self.p)
            if allowed:continue
            reason=f"vaga antiga ({age} dias)"
            self.save_discarded(job,reason)
            self.conn.execute("UPDATE vagas SET status='Ignorada',selecionada_lote=0 WHERE id=?",(vid,));moved+=1
        self.conn.commit()
        if restored or moved:self.refresh()
        return moved,restored

    def show_manually_discarded(self):
        self.show_discarded(manual_only=True)

    def show_location_excluded(self):
        self.show_discarded(location_only=True)

    def show_model_unconfirmed(self):
        self.show_discarded(model_unconfirmed=True)

    def show_discarded(self,manual_only=False,location_only=False,model_unconfirmed=False):
        window_key=("modelo_confirmar" if model_unconfirmed else "localidades_excluidas" if location_only
                    else "descartadas_manuais" if manual_only else "fora_perfil")
        window_title=("Vagas com modelo a confirmar" if model_unconfirmed else
                      "Vagas excluídas por localidade" if location_only else
                      "Vagas descartadas pelo usuário" if manual_only else "Vagas fora do perfil")
        win,created=self.managed_window(window_key,window_title,"1250x650")
        if not created:return

        top=ttk.Frame(win,padding=8);top.pack(fill="x")
        ttk.Label(top,text="Filtro:").pack(side="left")
        q=tk.StringVar()
        ent=ttk.Entry(top,textvariable=q,width=35);ent.pack(side="left",padx=5)
        ttk.Label(top,text="Compatibilidade:").pack(side="left",padx=(16,5))
        compatibility_filter=tk.StringVar(value="Todas")
        compatibility_box=ttk.Combobox(top,textvariable=compatibility_filter,state="readonly",width=12,
                                       values=["Todas","50% ou mais","70% ou mais","85% ou mais"])
        compatibility_box.pack(side="left")

        pane=ttk.Panedwindow(win,orient="horizontal");pane.pack(fill="both",expand=True,padx=8,pady=(0,8))
        left,right=ttk.Frame(pane),ttk.Frame(pane);pane.add(left,weight=3);pane.add(right,weight=2)

        cols=("compat","titulo","empresa","local","fonte","motivo")
        tr=ttk.Treeview(left,columns=cols,show="headings")
        for c,t,w in [
            ("compat","Compat.",75),("titulo","Título",260),("empresa","Empresa",160),("local","Local",150),
            ("fonte","Fonte",90),("motivo","Motivo do descarte",270)
        ]:
            tr.heading(c,text=t);tr.column(c,width=w,anchor="center" if c=="compat" else "w")
        tr.pack(fill="both",expand=True)

        title=tk.StringVar(value="Selecione uma vaga descartada")
        ttk.Label(right,textvariable=title,font=("Segoe UI",12,"bold"),wraplength=450).pack(anchor="w",pady=(5,8))
        meta=tk.StringVar();ttk.Label(right,textvariable=meta,wraplength=450).pack(anchor="w")
        c=getattr(self,"colors",{"white":"#202938","ink":"#e6eaf0","line":"#344154","blue_soft":"#304a70"})
        desc=tk.Text(right,wrap="word",bg=c["white"],fg=c["ink"],insertbackground=c["ink"],
                     selectbackground=c["blue_soft"],relief="flat",highlightthickness=1,highlightbackground=c["line"])
        desc.pack(fill="both",expand=True,pady=8);desc.configure(state="disabled")
        current={"url":""};compatibility_scores={};visible_ids=[]

        def load():
            for x in tr.get_children():tr.delete(x)
            s=q.get().strip()
            sql="""SELECT id,titulo,empresa,local,fonte,motivo_descarte,descricao,url,data_publicacao,salario
                   FROM descartadas"""
            clauses=[];params=[]
            if manual_only:
                clauses.append("motivo_descarte = ?")
                params.append("Descartada pelo usuário")
            elif model_unconfirmed:
                clauses.append("motivo_descarte LIKE ?")
                params.append("%remoto não confirmado%")
            elif location_only:
                clauses.append("motivo_descarte LIKE ?")
                params.append("localidade excluída pelo usuário — %")
            else:
                clauses.append("motivo_descarte != ? AND motivo_descarte NOT LIKE ?")
                params.extend(["Descartada pelo usuário","localidade excluída pelo usuário — %"])
            if s:
                clauses.append("(titulo LIKE ? OR empresa LIKE ? OR local LIKE ? OR motivo_descarte LIKE ?)")
                like="%"+s+"%";params.extend([like,like,like,like])
            if clauses:sql+=" WHERE "+" AND ".join(clauses)
            sql+=" ORDER BY descartada_em DESC"
            minimum={"Todas":0,"50% ou mais":50,"70% ou mais":70,"85% ou mais":85}.get(compatibility_filter.get(),0)
            compatibility_scores.clear();visible=[];visible_ids.clear()
            for r in self.conn.execute(sql,params):
                did,title_text,company,local,source,discard_reason,description,url,published,salary=r
                job={"titulo":title_text or "","empresa":company or "","local":local or "",
                     "descricao":description or "","url":url or "","fonte":source or "",
                     "data_publicacao":published or "","salario":salary or ""}
                score=curriculum_compatibility(job,self.p,self.cv);compatibility_scores[did]=score
                if score>=minimum:visible.append((did,score,title_text,company,local,source,discard_reason))
            visible.sort(key=lambda item:(-item[1],item[0]))
            for did,score,title_text,company,local,source,discard_reason in visible:
                visible_ids.append(did)
                tr.insert("","end",iid=str(did),values=(f"{score}%",title_text,company,local,source,discard_reason))

        def selected(_=None):
            sel=tr.selection()
            if not sel:return
            row=self.conn.execute("""SELECT titulo,empresa,local,fonte,motivo_descarte,descricao,url,
                                     data_publicacao,salario FROM descartadas WHERE id=?""",(int(sel[0]),)).fetchone()
            if not row:return
            t,e,l,f,m,d,u,dt,sal=row
            title.set(t)
            score=compatibility_scores.get(int(sel[0]),curriculum_compatibility(
                {"titulo":t or "","empresa":e or "","local":l or "","descricao":d or ""},self.p,self.cv))
            meta.set(f"Compatibilidade com o currículo: {score}%\\n{e} | {l} | {f} | {dt or 'data n/i'} | {sal or 'salário n/i'}\\nMotivo: {m}")
            current["url"]=u or ""
            desc.configure(state="normal");desc.delete("1.0","end");desc.insert("1.0",d or "Descrição não coletada.")
            desc.configure(state="disabled")

        def open_selected():
            if current["url"]:webbrowser.open(current["url"])

        def restore_selected():
            sel=tr.selection()
            if not sel:return
            if self.restore_discarded_record(int(sel[0])):load()

        def restore_all():
            if manual_only:
                ids=[r[0] for r in self.conn.execute(
                    "SELECT id FROM descartadas WHERE motivo_descarte=? ORDER BY descartada_em DESC",
                    ("Descartada pelo usuário",))]
            elif location_only:
                ids=[r[0] for r in self.conn.execute(
                    "SELECT id FROM descartadas WHERE motivo_descarte LIKE ? ORDER BY descartada_em DESC",
                    ("localidade excluída pelo usuário — %",))]
            else:
                ids=[r[0] for r in self.conn.execute(
                    "SELECT id FROM descartadas WHERE motivo_descarte!=? AND motivo_descarte NOT LIKE ? ORDER BY descartada_em DESC",
                    ("Descartada pelo usuário","localidade excluída pelo usuário — %"))]
            if not ids:messagebox.showinfo("Restaurar todas","Não há vagas descartadas para restaurar.",parent=win);return
            if not messagebox.askyesno("Restaurar vagas",
                f"Devolver todas as {len(ids)} vaga(s) deste diretório à lista principal?\n\nEsta ação não considera o filtro exibido. Confira o local: algumas podem ser presenciais ou híbridas fora da sua região.",parent=win):return
            restored=0
            for did in ids:
                if self.restore_discarded_record(did,refresh=False):restored+=1
            self.refresh();load();messagebox.showinfo("Vagas restauradas",f"{restored} vaga(s) voltaram para a lista.",parent=win)

        btn=ttk.Frame(right);btn.pack(fill="x",pady=(0,5))
        ttk.Button(btn,text="Abrir vaga",command=open_selected).pack(side="left")
        ttk.Button(btn,text="Restaurar manualmente",command=restore_selected).pack(side="left",padx=5)
        ttk.Button(btn,text="Restaurar todas",style="Primary.TButton",command=restore_all).pack(side="left",padx=5)

        ent.bind("<KeyRelease>",lambda e:load())
        compatibility_box.bind("<<ComboboxSelected>>",lambda e:load())
        tr.bind("<<TreeviewSelect>>",selected)
        load()

    def recalculate_all_jobs(self):
        """Reaplica o perfil atual sem alterar status, fila ou histórico do usuário."""
        rows=self.conn.execute("""SELECT id,titulo,empresa,local,descricao,url,fonte,data_publicacao,
                                  salario,COALESCE(workplace_type,'') FROM vagas""").fetchall()
        for vid,titulo,empresa,local,descricao,url,fonte,pub,salario,workplace in rows:
            job={"titulo":titulo or "","empresa":empresa or "","local":local or "",
                 "descricao":descricao or "","url":url or "","fonte":fonte or "",
                 "data_publicacao":pub or "","salario":salario or "","workplace_type":workplace or ""}
            score,label,reason,mode=score_job(job,self.p,self.cv)
            ld=location_decision(job,self.p)
            self.conn.execute("""UPDATE vagas SET score=?,classificacao=?,motivo=?,modalidade=?,categoria=?,
                                  decisao=?,confianca=?,location_confidence=?,location_evidence=? WHERE id=?""",
                              (score,label,reason,mode,job_category(job,self.p),decision_level(job,self.p,mode),
                               collection_confidence(job),ld["confidence"],ld["evidence"],vid))
        self.conn.commit();self.separate_outside_region_jobs();return len(rows)

    def recalculate_with_notice(self,parent=None):
        self.p=load_profile();self.load_feedback_profile();self.cv=read_cv()
        total=self.recalculate_all_jobs();self.refresh()
        messagebox.showinfo("Vagas reavaliadas",f"{total} vaga(s) foram reavaliadas com o currículo e as preferências atuais.",parent=parent)

    def restore_auto_filtered_remote(self,vid):
        row=self.conn.execute("""SELECT titulo,empresa,local,descricao,url,fonte,COALESCE(workplace_type,'')
                                 FROM vagas WHERE id=?""",(vid,)).fetchone()
        if not row:return False
        title,company,local,description,url,source,workplace=row
        job={"titulo":title or "","empresa":company or "","local":local or "","descricao":description or "",
             "url":url or "","fonte":source or "","workplace_type":workplace or ""}
        ld=location_decision(job,self.p)
        if not ld["mode"].startswith("Remoto Brasil") or "confirmado" not in ld["mode"]:return False
        auto=self.conn.execute("""SELECT id FROM descartadas WHERE url=?
                                  AND motivo_descarte='Local fora da região — remoto não confirmado'""",(url,)).fetchone()
        if not auto:return False
        self.conn.execute("""UPDATE vagas SET status='Nova',modalidade=?,location_confidence=?,
                              location_evidence=? WHERE id=? AND status='Ignorada'""",
                          (ld["mode"],ld["confidence"],ld["evidence"],vid))
        self.conn.execute("DELETE FROM descartadas WHERE id=?",(auto[0],))
        return True

    def update_pending_descriptions(self,limit=40,notify=True):
        if getattr(self,"description_update_running",False):
            if notify:messagebox.showinfo("Descrições","A atualização já está em andamento.")
            return
        self.description_update_running=True
        self.info.set("Atualizando descrições pendentes...")
        self.start_worker(self._update_pending_descriptions_worker,limit,notify,name="atualizar-descricoes")

    def _update_pending_descriptions_worker(self,limit,notify=True):
        updated=failed=0
        try:
            rows=self.conn.execute("""SELECT id,url FROM vagas
                WHERE fonte='LinkedIn' AND LENGTH(TRIM(COALESCE(descricao,'')))<120
                  AND COALESCE(description_attempts,0)<4
                  AND (COALESCE(description_next_retry_at,'')='' OR description_next_retry_at<=?)
                ORDER BY COALESCE(description_attempts,0),id DESC LIMIT ?""",
                (datetime.now(timezone.utc).isoformat(timespec="seconds"),int(limit))).fetchall()
            for vid,url in rows:
                if self.shutdown_requested():break
                attempted=datetime.now(timezone.utc).isoformat(timespec="seconds")
                try:
                    detail=generic_job_from_url(url,"LinkedIn")
                    desc=(detail.get("descricao") or "").strip()
                    if len(desc)<120:raise RuntimeError("A página não disponibilizou uma descrição completa")
                    self.conn.execute("""UPDATE vagas SET descricao=?,description_status='disponivel',
                        description_attempts=description_attempts+1,description_last_error='',
                        description_last_attempt_at=?,description_next_retry_at='',description_source='LinkedIn detalhe',
                        workplace_type=CASE WHEN ?!='' THEN ? ELSE workplace_type END,
                        workplace_type_raw=CASE WHEN ?!='' THEN ? ELSE workplace_type_raw END,
                        workplace_source=CASE WHEN ?!='' THEN ? ELSE workplace_source END
                        WHERE id=?""",(desc,attempted,
                            detail.get("workplace_type","") or "",detail.get("workplace_type","") or "",
                            detail.get("workplace_type_raw","") or "",detail.get("workplace_type_raw","") or "",
                            detail.get("workplace_source","") or "",detail.get("workplace_source","") or "",vid))
                    self.restore_auto_filtered_remote(vid);updated+=1
                except Exception as error:
                    retry=(datetime.now(timezone.utc)+timedelta(hours=12)).isoformat(timespec="seconds")
                    self.conn.execute("""UPDATE vagas SET description_status='pendente',
                        description_attempts=description_attempts+1,description_last_error=?,
                        description_last_attempt_at=?,description_next_retry_at=? WHERE id=?""",
                        (str(error)[:300],attempted,retry,vid));failed+=1
                self.conn.commit();time.sleep(.8)
            if updated:self.recalculate_all_jobs()
            self.ui_call(lambda:self._finish_description_update(updated,failed,notify=notify))
        except Exception as error:
            LOGGER.exception("Falha ao atualizar descrições pendentes")
            self.ui_call(lambda:self._finish_description_update(updated,failed,str(error),notify))

    def _finish_description_update(self,updated,failed,error="",notify=True):
        self.description_update_running=False;self.refresh()
        self.info.set(f"Descrições: {updated} atualizadas, {failed} ainda pendentes")
        if notify:
            if error:messagebox.showwarning("Descrições",f"A atualização terminou com erro: {error}")
            else:messagebox.showinfo("Descrições",f"{updated} descrição(ões) atualizada(s).\n{failed} continuará(ão) pendente(s) para nova tentativa.")

    def run_source(self,src):
        try:
            self.p=load_profile();self.load_feedback_profile();self.cv=read_cv();jobs=self.collect(src)
            if self.shutdown_requested():return
            a=u=f=0;seen=set();seen_fp=set()
            for j in jobs:
                if self.shutdown_requested():
                    self.conn.rollback();return
                url=j.get("url") or ""
                if not url:continue
                key=url.split("?")[0]
                if key in seen:continue
                fp=job_fingerprint(j)
                if fp and fp in seen_fp:continue
                seen.add(key)
                if fp:seen_fp.add(fp)
                j["url"]=key
                ok,why=hard_filter(j,self.p)
                if not ok:
                    self.save_discarded(j,why)
                    fp=job_fingerprint(j)
                    self.conn.execute("""UPDATE vagas SET status='Ignorada',selecionada_lote=0,decisao='REVISAR'
                                         WHERE status='Nova' AND COALESCE(selecionada_lote,0)=0
                                           AND (url=? OR (?!='' AND fingerprint=?))""",(j["url"],fp,fp))
                    f+=1;continue
                score,label,reason,mode=score_job(j,self.p,self.cv)
                location_info=location_decision(j,self.p)
                structured=structured_persistence(j,location_info)
                # Se já havia sido descartada em uma busca anterior e agora passou, remove do histórico ativo.
                self.conn.execute("""DELETE FROM descartadas WHERE url=?
                                     AND motivo_descarte!='Descartada pelo usuário'""",(j["url"],))
                category=job_category(j,self.p)
                decisao=decision_level(j,self.p,mode)
                confianca=collection_confidence(j)
                fp=job_fingerprint(j)
                fontes=j.get("fontes_encontradas") or j.get("fonte","")

                existing=None
                if fp:
                    existing=self.conn.execute("SELECT id,fonte,fontes FROM vagas WHERE fingerprint=? ORDER BY id LIMIT 1",(fp,)).fetchone()

                if existing:
                    vid,old_source,old_fontes=existing
                    all_sources=[]
                    for raw in (old_fontes,old_source,fontes,j.get("fonte","")):
                        for source_name in str(raw or "").split(" + "):
                            if source_name and source_name not in all_sources:all_sources.append(source_name)
                    self.conn.execute("""UPDATE vagas SET titulo=?,empresa=?,local=?,modalidade=?,descricao=?,
                        fonte=?,fontes=?,data_publicacao=?,salario=?,score=?,classificacao=?,motivo=?,
                        categoria=?,decisao=?,confianca=?,fingerprint=?,workplace_type=?,
                        location_confidence=?,location_evidence=?,workplace_type_raw=?,workplace_source=?,
                        structured_location_json=?,applicant_location_requirements=?,remote_eligible_brazil=?,
                        modality_checked_at=?,status=CASE WHEN status='Pesquisa limpa' THEN 'Nova' ELSE status END WHERE id=?""",
                        (j.get("titulo",""),j.get("empresa",""),j.get("local",""),mode,j.get("descricao",""),
                         j.get("fonte","")," + ".join(all_sources),j.get("data_publicacao",""),j.get("salario",""),
                         score,label,reason,category,decisao,confianca,fp,j.get("workplace_type",""),
                         location_info["confidence"],location_info["evidence"],*structured,vid))
                    u+=1
                else:
                    vals=(j.get("titulo",""),j.get("empresa",""),j.get("local",""),mode,j.get("descricao",""),
                           j["url"],j.get("fonte",""),fontes,j.get("data_publicacao",""),j.get("salario",""),
                           score,label,reason,category,decisao,confianca,fp,j.get("workplace_type",""),
                           location_info["confidence"],location_info["evidence"],*structured)
                    try:
                        self.conn.execute("""INSERT INTO vagas(titulo,empresa,local,modalidade,descricao,url,fonte,fontes,
                        data_publicacao,salario,score,classificacao,motivo,categoria,decisao,confianca,fingerprint,
                        workplace_type,location_confidence,location_evidence,workplace_type_raw,workplace_source,
                        structured_location_json,applicant_location_requirements,remote_eligible_brazil,
                        modality_checked_at,status)
                        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?, 'Nova')""",vals);a+=1
                    except sqlite3.IntegrityError:
                        self.conn.execute("""UPDATE vagas SET titulo=?,empresa=?,local=?,modalidade=?,descricao=?,
                            fonte=?,fontes=?,data_publicacao=?,salario=?,score=?,classificacao=?,motivo=?,
                            categoria=?,decisao=?,confianca=?,fingerprint=?,workplace_type=?,
                            location_confidence=?,location_evidence=?,workplace_type_raw=?,workplace_source=?,
                            structured_location_json=?,applicant_location_requirements=?,remote_eligible_brazil=?,
                            modality_checked_at=?,status=CASE WHEN status='Pesquisa limpa' THEN 'Nova' ELSE status END WHERE url=?""",
                            (j.get("titulo",""),j.get("empresa",""),j.get("local",""),mode,j.get("descricao",""),
                             j.get("fonte",""),fontes,j.get("data_publicacao",""),j.get("salario",""),score,label,
                             reason,category,decisao,confianca,fp,j.get("workplace_type",""),
                             location_info["confidence"],location_info["evidence"],*structured,j["url"]));u+=1
            self.conn.execute("""UPDATE vagas SET description_status=CASE
                WHEN LENGTH(TRIM(COALESCE(descricao,'')))>=120 THEN 'disponivel'
                WHEN fonte='LinkedIn' THEN 'pendente' ELSE 'indisponivel' END,
                description_source=CASE WHEN LENGTH(TRIM(COALESCE(descricao,'')))>=120
                                        THEN fonte ELSE COALESCE(description_source,'') END
                WHERE COALESCE(description_status,'')='' OR description_status!='disponivel'""")
            self.conn.commit()
            flush_detail_cache(force=True)
            self.ui_call(lambda:self.finish(a,u,f,len(jobs),src))
        except Exception as e:
            LOGGER.exception("Falha durante a busca: %s",src)
            error_msg=str(e)
            self.ui_call(lambda msg=error_msg:self.fail(msg))

    def finish(self,a,u,f,total,src):
        self.search_running=False;self.search_button.configure(state="normal")
        self.stop_search_animation()
        self.p=load_profile()
        self.apply_pcd_preference(pcd_vacancies_enabled(self.p))
        self.apply_internship_mode(internship_search_mode(self.p))
        self.apply_apprentice_preference(bool(self.p.get("buscar_jovem_aprendiz",False)))
        self.prog.stop();self.refresh()
        rec=self.conn.execute("SELECT COUNT(*) FROM vagas WHERE status='Nova' AND COALESCE(selecionada_lote,0)=0 AND decisao='APROVADA'").fetchone()[0]
        rev=self.conn.execute("SELECT COUNT(*) FROM vagas WHERE status='Nova' AND COALESCE(selecionada_lote,0)=0 AND decisao='REVISAR'").fetchone()[0]
        self.info.set(f"Busca concluída: {rec} recomendadas • {rev} para conferir • {f} fora do perfil nesta busca")
        source_counts=self.__dict__.get("last_source_counts",{})
        if source_counts:
            LOGGER.info("Busca concluída por fonte: %s",", ".join(
                f"{name}={count}" for name,count in sorted(source_counts.items())))
        with SOURCE_RESULTS_LOCK:
            daily_metrics=dict(SOURCE_RESULT_METRICS)
        if daily_metrics:
            LOGGER.info("Métricas do cache de fontes: %s",json.dumps(daily_metrics,ensure_ascii=False,sort_keys=True))
        if total and src in ("all","linkedin"):
            self.after(350,lambda:self.update_pending_descriptions(limit=20,notify=False))
        if total==0:
            messagebox.showwarning("Fonte sem resultados",f"A fonte '{src}' não retornou resultados nesta tentativa.\n\nGoogle e LinkedIn podem limitar consultas automatizadas temporariamente.")

    def fail(self,e):
        self.search_running=False;self.search_button.configure(state="normal")
        self.stop_search_animation()
        self.prog.stop();self.info.set("Não conseguimos buscar vagas agora.")
        messagebox.showerror("Não foi possível buscar","Não conseguimos buscar vagas agora. Tente novamente.\n\nOs detalhes foram registrados no diagnóstico.")

    def start_search_animation(self):
        self.search_button.configure(text="Buscando…")
        if not self.search_activity.winfo_ismapped():self.search_activity.pack(side="left",padx=20)
        self.search_animation_frame=0
        def animate():
            if not self.search_running:return
            frames=("◐","◓","◑","◒")
            self.search_spinner.set(frames[self.search_animation_frame%len(frames)])
            self.search_animation_frame+=1
            self.search_animation_job=self.after(140,animate)
        animate()

    def stop_search_animation(self):
        if self.search_animation_job:
            try:self.after_cancel(self.search_animation_job)
            except tk.TclError:pass
        self.search_animation_job=None;self.search_button.configure(text="Buscar vagas")
        self.search_activity.pack_forget()


    def update_dashboard_counts(self):
        try:
            rec=self.conn.execute("SELECT COUNT(*) FROM vagas WHERE decisao='APROVADA' AND status='Nova' AND COALESCE(selecionada_lote,0)=0").fetchone()[0]
            rev=self.conn.execute("SELECT COUNT(*) FROM vagas WHERE decisao='REVISAR' AND status='Nova' AND COALESCE(selecionada_lote,0)=0").fetchone()[0]
            appc=self.conn.execute("SELECT COUNT(*) FROM vagas WHERE status IN ('Candidatado','Entrevista','Rejeitado')").fetchone()[0]
            out=self.conn.execute("""SELECT COUNT(*) FROM descartadas
                                     WHERE motivo_descarte!=? AND motivo_descarte NOT LIKE ?""",
                                  ("Descartada pelo usuário","localidade excluída pelo usuário — %")).fetchone()[0]
            excluded=self.conn.execute("SELECT COUNT(*) FROM descartadas WHERE motivo_descarte LIKE ?",
                                       ("localidade excluída pelo usuário — %",)).fetchone()[0]
            model_pending=self.conn.execute("SELECT COUNT(*) FROM descartadas WHERE motivo_descarte LIKE ?",
                                            ("%remoto não confirmado%",)).fetchone()[0]
            discarded=self.conn.execute("SELECT COUNT(*) FROM descartadas WHERE motivo_descarte=?",
                                        ("Descartada pelo usuário",)).fetchone()[0]
            self.stat_rec.set(f"Recomendadas: {rec}")
            self.stat_rev.set(f"Vale conferir: {rev}")
            self.stat_app.set(f"Candidaturas: {appc}")
            self.stat_out.set(f"Fora do perfil: {out}")
            self.stat_model_pending.set(f"Modelo a confirmar: {model_pending}")
            self.stat_location_excluded.set(f"Localidades excluídas: {excluded}")
            self.stat_discarded.set(f"Descartadas: {discarded}")
        except Exception:
            pass

    def refresh(self):
        if not hasattr(self,"tree"):return
        for x in self.tree.get_children():self.tree.delete(x)

        view=self.view_mode.get() if hasattr(self,"view_mode") else "todas"
        sql,pa=jobs_query(view,self.q.get())

        rows=list(self.conn.execute(sql,pa))
        for vid,score,titulo,empresa,local,modo,cat,status,queued,published_at in rows:
            lm=simple_location_mode(local,modo)
            tag="great" if (score or 0)>=80 else "possible"
            checked=vid in self.batch_selection if hasattr(self,"batch_selection") else bool(queued)
            published=format_date_br(published_at)
            if published=="Data não informada":published="Não informada"
            self.tree.insert("","end",iid=str(vid),values=("☑" if checked else "☐",f"{score}%",titulo,empresa,lm,published),tags=(tag,))
        if self.job_sort_state["column"]:
            self.sort_jobs(self.job_sort_state["column"],self.job_sort_state["reverse"])
        if hasattr(self,"empty_state") and hasattr(self,"main_pane"):
            show_empty=not self.cv.strip() and not rows
            if show_empty and not self.empty_state.winfo_manager():
                self.empty_state.pack(fill="x",pady=(0,10),before=self.main_pane)
            elif not show_empty and self.empty_state.winfo_manager():
                self.empty_state.pack_forget()
        self.update_dashboard_counts()

    def select(self,e=None):
        sel=self.tree.selection()
        if not sel:return
        self.current=int(sel[0])
        row=self.conn.execute("""SELECT titulo,empresa,local,modalidade,descricao,score,classificacao,status,salario,
                                COALESCE(location_evidence,''),COALESCE(data_publicacao,''),COALESCE(fonte,''),
                                COALESCE(confianca,'Baixa'),COALESCE(description_status,'')
                                FROM vagas WHERE id=?""",(self.current,)).fetchone()
        if not row:return
        t,emp,loc,mode,desc,score,cl,status,salario,location_evidence,published_at,source,confidence,description_status=row
        self.tv.set(f"{t}")
        model_detail=f" | Motivo: {location_evidence}" if mode=="Verificar modelo" and location_evidence else ""
        self.meta.set(f"{emp or 'Empresa não informada'} | {simple_location_mode(loc,mode)} | Publicação: {format_date_br(published_at)}{model_detail}")
        description_note=" • Descrição incompleta" if description_status!="disponivel" or len((desc or "").strip())<120 else ""
        self.data_quality.set(f"Fonte: {source or 'não informada'} • Qualidade dos dados: {confidence or 'Baixa'}{description_note}")
        self.desc_box.configure(state="normal");self.desc_box.delete("1.0","end");self.desc_box.insert("1.0",(desc or "Descrição ainda não disponível. Use Configurações > Atualizar descrições pendentes para tentar novamente.").strip());self.desc_box.configure(state="disabled")
        reqs=extract_requirements({"titulo":t or "","descricao":desc or ""})
        reqtxt="\n".join("• "+x[0] for x in reqs) if reqs else "Requisitos não identificados separadamente."
        self.req_box.configure(state="normal");self.req_box.delete("1.0","end");self.req_box.insert("1.0",reqtxt);self.req_box.configure(state="disabled")
        pay,bens=salary_and_benefits({"salario":salario or "","descricao":desc or ""})
        paytxt="Salário: "+pay
        if bens: paytxt+="\n"+"\n".join("• "+b for b in bens)
        else: paytxt+="\nBenefícios não identificados."
        self.pay_box.configure(state="normal");self.pay_box.delete("1.0","end");self.pay_box.insert("1.0",paytxt);self.pay_box.configure(state="disabled")

    def open_job(self):
        if not self.current:return
        row=self.conn.execute("SELECT url FROM vagas WHERE id=?",(self.current,)).fetchone()
        if row and row[0]:webbrowser.open(row[0])

    def discard_current(self):
        if not self.current:return
        row=self.conn.execute("""SELECT titulo,empresa,local,descricao,url,fonte,data_publicacao,salario
                                 FROM vagas WHERE id=?""",(self.current,)).fetchone()
        if not row:return
        if not messagebox.askyesno("Descartar vaga","Mover esta vaga para Descartadas?\n\nVocê poderá restaurá-la depois."):
            return
        t,e,l,d,u,f,dt,sal=row
        self.record_feedback(self.current,False)
        self.save_discarded({"titulo":t,"empresa":e,"local":l,"descricao":d,"url":u,
                             "fonte":f,"data_publicacao":dt,"salario":sal},"Descartada pelo usuário")
        with self.conn:
            self.conn.execute("UPDATE vagas SET status='Ignorada',selecionada_lote=0 WHERE id=?",(self.current,))
        if "batch_selection" in self.__dict__:self.batch_selection.discard(self.current);self.update_queue_action()
        self.current=None;self.info.set("Vaga movida para Descartadas.");self.refresh()

    def archive_application(self,vid):
        with self.conn:
            self.conn.execute("UPDATE vagas SET status='Arquivada',selecionada_lote=0 WHERE id=?",(vid,))
        if "batch_selection" in self.__dict__:self.batch_selection.discard(vid);self.update_queue_action()
        self.refresh()

    def show_applications(self):
        win,created=self.managed_window("candidaturas","Minhas candidaturas","1120x650")
        if not created:return
        body=ttk.Frame(win,padding=18);body.pack(fill="both",expand=True)
        head=ttk.Frame(body);head.pack(fill="x",pady=(0,14))
        title=ttk.Frame(head);title.pack(side="left")
        ttk.Label(title,text="Minhas candidaturas",font=("Segoe UI",18,"bold")).pack(anchor="w")
        summary=tk.StringVar(value="Nenhuma candidatura registrada")
        ttk.Label(title,textvariable=summary,foreground=getattr(self,"colors",{}).get("muted","#9aa7b5")).pack(anchor="w",pady=(2,0))
        ttk.Button(head,text="Voltar às vagas",command=win._managed_close).pack(side="right")

        cols=("data","titulo","empresa","remove")
        table_frame=ttk.Frame(body,style="Panel.TFrame",padding=1);table_frame.pack(fill="both",expand=True)
        tr=ttk.Treeview(table_frame,columns=cols,show="headings",selectmode="browse")
        self.applications_tree=tr
        for col,label,width in [("data","DATA",110),("titulo","VAGA",440),("empresa","EMPRESA",300),("remove","AÇÃO",90)]:
            tr.heading(col,text=label);tr.column(col,width=width,anchor="center" if col=="remove" else "w")
        scroll=ttk.Scrollbar(table_frame,orient="vertical",command=tr.yview);tr.configure(yscrollcommand=scroll.set)
        scroll.pack(side="right",fill="y");tr.pack(side="left",fill="both",expand=True)

        def load():
            for item in tr.get_children():tr.delete(item)
            rows=self.conn.execute("""SELECT id,titulo,empresa,COALESCE(candidatura_em,'')
                                    FROM vagas WHERE status IN ('Candidatado','Entrevista','Rejeitado')
                                    ORDER BY COALESCE(NULLIF(candidatura_em,''),criada_em) DESC,id DESC""").fetchall()
            for vid,job,company,applied_at in rows:
                date_text="Data não registrada"
                if applied_at:
                    try:date_text=datetime.fromisoformat(applied_at.replace("Z","+00:00")).strftime("%d/%m/%Y")
                    except Exception:date_text=applied_at[:10]
                tr.insert("","end",iid=str(vid),values=(date_text,job,company,"Remover"))
            summary.set(f"{len(rows)} candidatura(s) registrada(s)" if rows else "Nenhuma candidatura registrada")

        def open_selected(_event=None):
            selected=tr.selection()
            if not selected:return
            row=self.conn.execute("SELECT url FROM vagas WHERE id=?",(int(selected[0]),)).fetchone()
            if row and row[0]:webbrowser.open(row[0])

        def remove_selected():
            selected=tr.selection()
            if not selected:return
            if not messagebox.askyesno("Remover do resumo",
                "Remover esta vaga de Minhas candidaturas?\n\nO registro será arquivado, não apagado.",parent=win):return
            self.archive_application(int(selected[0]));load()

        def click_action(event):
            if tr.identify_region(event.x,event.y)!="cell" or tr.identify_column(event.x)!="#4":return
            iid=tr.identify_row(event.y)
            if not iid:return "break"
            tr.selection_set(iid);remove_selected();return "break"

        actions=ttk.Frame(body);actions.pack(fill="x",pady=(10,0))
        ttk.Button(actions,text="Abrir vaga selecionada",command=open_selected).pack(side="left")
        ttk.Button(actions,text="Atualizar resumo",command=load).pack(side="left",padx=6)
        tr.bind("<Button-1>",click_action,add="+");tr.bind("<Double-1>",open_selected);load()

    def show_batch(self):
        win,created=self.managed_window("fila","Fila de candidaturas","1050x560")
        if not created:return
        cols=("score","titulo","empresa","modalidade","resultado","remove")
        tr=ttk.Treeview(win,columns=cols,show="headings")
        for c,t,w in [("score","Score",60),("titulo","Título",300),("empresa","Empresa",180),
                      ("modalidade","Modalidade",150),("resultado","Último resultado",230),("remove","Ação",90)]:
            tr.heading(c,text=t);tr.column(c,width=w,anchor="center" if c=="remove" else "w")
        tr.pack(fill="both",expand=True,padx=8,pady=8)
        def load():
            for item in tr.get_children():tr.delete(item)
            rows=self.conn.execute("""SELECT id,score,titulo,empresa,modalidade,COALESCE(ultimo_resultado,'')
                                    FROM vagas WHERE selecionada_lote=1 ORDER BY score DESC""").fetchall()
            for vid,score,title,company,mode,result in rows:
                tr.insert("","end",iid=str(vid),values=(score,title,company,mode,result,"Remover"))
        def remove_selected():
            selected=tr.selection()
            if not selected:return
            self.remove_from_batch(int(selected[0]));load()
        def click_action(event):
            if tr.identify_region(event.x,event.y)!="cell" or tr.identify_column(event.x)!="#6":return
            iid=tr.identify_row(event.y)
            if iid:tr.selection_set(iid);remove_selected();return "break"
        tr.bind("<Button-1>",click_action,add="+");load()
        footer=ttk.Frame(win,padding=(8,0,8,8));footer.pack(fill="x")
        ttk.Label(footer,text="A fila não envia respostas desconhecidas, CAPTCHAs ou declarações por você.").pack(side="left")
        def clear_and_close():
            if self.clear_batch():win._managed_close()
        ttk.Button(footer,text="Limpar fila",style="Danger.TButton",command=clear_and_close).pack(side="right")
        ttk.Button(footer,text="Remover selecionada",command=remove_selected).pack(side="right",padx=6)

    def wait_for_user_review(self,titulo,vid):
        """
        Called by worker thread. Opens a Tk dialog on the UI thread and truly
        waits until the user chooses what to do before proceeding.
        """
        done=threading.Event()
        action={"value":"next"}

        def show():
            try:
                self.deiconify();self.lift();self.focus_force()
                win=tk.Toplevel(self);win.title("Confirmar candidatura");win.transient(self);win.withdraw()
                win.resizable(False,False)
                card=ttk.Frame(win,padding=(18,15,18,14));card.pack(fill="both",expand=True)
                ttk.Label(card,text="Você concluiu esta candidatura?",font=("Segoe UI",13,"bold")).pack(anchor="w")
                ttk.Label(card,text=titulo,font=("Segoe UI",10,"bold"),wraplength=540).pack(anchor="w",pady=(8,4))
                ttk.Label(card,text="Revise a vaga no navegador e informe o resultado. Assim a fila e Minhas candidaturas serão atualizadas.",
                          wraplength=540).pack(anchor="w",pady=(0,12))
                ttk.Separator(card).pack(fill="x",pady=(0,12))
                b=ttk.Frame(card);b.pack(fill="x")
                for column in range(3):b.columnconfigure(column,weight=1,uniform="review")

                def finish(value):
                    action["value"]=value
                    try:win.grab_release();win.destroy()
                    finally:done.set()

                ttk.Button(b,text="Sim, candidatei-me",style="Primary.TButton",command=lambda:finish("applied")).grid(row=0,column=0,sticky="ew",padx=(0,5))
                ttk.Button(b,text="Ainda não",command=lambda:finish("skip")).grid(row=0,column=1,sticky="ew",padx=5)
                ttk.Button(b,text="Parar lote",command=lambda:finish("stop")).grid(row=0,column=2,sticky="ew",padx=(5,0))
                win.protocol("WM_DELETE_WINDOW",lambda:finish("skip"))
                win.update_idletasks();self.center_window(win,580,win.winfo_reqheight());win.deiconify()
                win.attributes("-topmost",True);win.grab_set()
                win.after(100,lambda:(win.lift(),win.focus_force(),self.bell()))
            except Exception:
                LOGGER.exception("Falha ao exibir confirmação da candidatura")
                action["value"]="stop";done.set()

        if self.shutdown_requested():return "stop"
        self.ui_call(show)
        while not done.wait(.2):
            if self.shutdown_requested():return "stop"
        if action["value"]=="applied":
            self.mark_application_completed(vid)
            self.ui_call(self.sync_batch_selection)
        return action["value"]

    def start_batch(self):
        rows=self.conn.execute("""SELECT id FROM vagas WHERE selecionada_lote=1 AND status='Nova'
                                  ORDER BY score DESC""").fetchall()
        if not rows:
            messagebox.showinfo("Fila vazia","Adicione vagas à fila antes de iniciar o lote.")
            return
        p=load_profile();limit=int(p.get("max_candidaturas_sessao",20))
        if not messagebox.askyesno("Iniciar lote",
            f"Há {len(rows)} vagas na fila.\n\nNesta sessão serão abertas no máximo {limit}.\n"
            "O programa preencherá dados conhecidos e anexará o currículo. "
            "O envio automático fica DESATIVADO por padrão.\n\nContinuar?"):
            return
        self.shutdown_event.clear()
        self.start_worker(self.run_batch,[r[0] for r in rows[:limit]],name="candidaturas")

    def run_batch(self,ids):
        try:
            from playwright.sync_api import sync_playwright
        except Exception:
            self.ui_call(lambda:messagebox.showerror("Playwright",
                "Playwright não está instalado corretamente. Execute preparar_navegador.bat."))
            return

        p=load_profile()
        try:
            browser_cfg=get_browser_launch_settings(p)
            profile_dir=browser_cfg["profile_dir"]
            os.makedirs(profile_dir,exist_ok=True)
        except Exception as e:
            error_text=str(e)
            self.ui_call(lambda text=error_text:messagebox.showerror("Navegador",text))
            return

        try:
            with sync_playwright() as pw:
                launch_args={"user_data_dir":profile_dir,"headless":False}
                if browser_cfg.get("executable_path"):
                    launch_args["executable_path"]=browser_cfg["executable_path"]
                browser_type=pw.firefox if browser_cfg["browser"]=="firefox" else pw.chromium
                ctx=browser_type.launch_persistent_context(**launch_args)
                page=ctx.pages[0] if ctx.pages else ctx.new_page()
                close_previous=False

                for idx,vid in enumerate(ids,1):
                    if self.shutdown_requested():break
                    if idx>1:
                        previous_page=page
                        page=ctx.new_page()
                        page.bring_to_front()
                        if close_previous:
                            try:previous_page.close()
                            except Exception:pass
                        close_previous=False
                    row=self.conn.execute("""SELECT titulo,url,empresa,descricao,local,salario
                                             FROM vagas WHERE id=?""",(vid,)).fetchone()
                    if not row:continue
                    titulo,url,empresa,descricao,local,salario=row
                    self.ui_call(lambda i=idx,t=titulo:
                        self.info.set(f"Candidatura {i}/{len(ids)}: {t}"))
                    result="";page_opened=False;review_shown=False

                    try:
                        page_opened=True
                        page.goto(url,wait_until="domcontentloaded",timeout=45000)
                        page.wait_for_timeout(2200)
                        content=norm(page.locator("body").inner_text(timeout=8000))

                        # Live validity check: pages may have closed after collection.
                        invalid=live_page_invalid_reason(page)
                        if invalid:
                            result="Pulada: "+invalid
                            self.conn.execute("""UPDATE vagas SET selecionada_lote=0,
                                               ultimo_resultado=?,tentativas_envio=tentativas_envio+1
                                               WHERE id=?""",(result,vid))
                            self.conn.commit()
                            continue

                        if any(x in content for x in [
                            "captcha","verify you are human","verifique que voce e humano",
                            "security check","confirme que voce e humano"
                        ]):
                            result="Revisão necessária: CAPTCHA/verificação humana"
                        else:
                            filled=autofill_known_fields(page,p)
                            files=attach_resume(page,p)
                            unknown=find_unknown_required_fields(page)
                            if unknown and p.get("parar_em_pergunta_desconhecida",True):
                                result=f"Revisão necessária: {len(unknown)} campo(s) obrigatório(s) sem resposta"
                            else:
                                result=f"Pronto para revisão; preenchidos={filled}, anexos={files}"

                        self.conn.execute("""UPDATE vagas SET tentativas_envio=tentativas_envio+1,
                                           ultimo_resultado=? WHERE id=?""",(result,vid))
                        self.conn.commit()

                        # Crucial V7.4 fix: actually WAIT for the user's decision.
                        review_shown=True
                        action=self.wait_for_user_review(titulo,vid)
                        if action=="stop":
                            self.ui_call(lambda:self.info.set("Lote interrompido pelo usuário"))
                            break
                        close_previous=(action=="applied")
                        if idx<len(ids):
                            self.ui_call(lambda:self.info.set("Abrindo a próxima vaga..."))

                        # O próximo link abre logo após a confirmação, sem depender do site anterior.
                        time.sleep(.2)

                    except Exception as e:
                        result="Erro: "+str(e)[:180]
                        LOGGER.exception("Falha ao preparar candidatura para a vaga %s",vid)
                        self.conn.execute("""UPDATE vagas SET tentativas_envio=tentativas_envio+1,
                                           ultimo_resultado=? WHERE id=?""",(result,vid))
                        self.conn.commit()
                        # A página pode ter aberto mesmo quando um seletor/preenchimento falhou.
                        # Nesse caso a confirmação continua obrigatória para não perder a candidatura manual.
                        if page_opened and not review_shown:
                            action=self.wait_for_user_review(titulo,vid)
                            if action=="stop":
                                self.ui_call(lambda:self.info.set("Lote interrompido pelo usuário"));break
                            close_previous=(action=="applied")
                        continue

                try:ctx.close()
                except:pass

            self.ui_call(lambda:self.info.set("Lote concluído"))
            self.ui_call(self.refresh)

        except Exception as e:
            LOGGER.exception("Falha geral no lote de candidaturas")
            self.ui_call(lambda:messagebox.showerror("Não foi possível continuar","Não conseguimos abrir essa vaga. Tente novamente.\n\nOs detalhes foram registrados no diagnóstico."))


    def reset_for_new_user(self,parent=None):
        """Apaga somente os dados pessoais desta edição e mantém o aplicativo instalado."""
        worker_lock=self.__dict__.get("worker_lock")
        workers=self.__dict__.get("worker_threads",set())
        if worker_lock:
            with worker_lock:active=[worker for worker in workers if worker.is_alive()]
        else:active=[]
        if self.search_running or active:
            messagebox.showinfo("Tarefa em andamento","Aguarde a tarefa terminar antes de trocar de usuário.",parent=parent)
            return False
        if not messagebox.askyesno("Usar para outra pessoa",
            "Isso removerá deste aplicativo:\n\n"
            "• currículo e perfil\n• vagas, fila e candidaturas\n"
            "• descartes, preferências e cache\n\n"
            "O programa continuará instalado. Deseja continuar?",parent=parent):
            return False
        if not messagebox.askyesno("Confirmar limpeza completa",
            "Os dados desta pessoa serão apagados permanentemente. Confirmar?",parent=parent):
            return False
        try:
            self.conn.execute("PRAGMA secure_delete=ON")
            with self.conn:
                for table in ("vagas","descartadas","preferencias_feedback","app_meta"):
                    if self.conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",(table,)).fetchone():
                        self.conn.execute(f"DELETE FROM {table}")
                if self.conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='sqlite_sequence'").fetchone():
                    self.conn.execute("DELETE FROM sqlite_sequence WHERE name IN ('vagas','descartadas')")
            self.conn.execute("VACUUM")
            removable=[CV_PATH,CACHE_PATH,SOURCE_RESULTS_PATH,os.path.join(BASE_DIR,"curriculo.pdf"),
                       os.path.join(BASE_DIR,"curriculo.docx")]
            removable.extend(glob.glob(os.path.join(BASE_DIR,"curriculo_original.*")))
            for path in removable:
                if os.path.isfile(path):os.remove(path)
            recursive=[BACKUP_DIR,os.path.join(BASE_DIR,"browser_profiles")]
            base_abs=os.path.abspath(BASE_DIR)
            for directory in recursive:
                target=os.path.abspath(directory)
                if os.path.commonpath([base_abs,target])!=base_abs or target==base_abs:
                    raise RuntimeError("Diretório de limpeza fora da pasta de dados")
                if os.path.isdir(target):shutil.rmtree(target)
            # O log pode conter nomes de vagas e diagnósticos da sessão anterior.
            for handler in list(logging.getLogger().handlers):
                try:handler.close()
                finally:logging.getLogger().removeHandler(handler)
            for log_file in glob.glob(LOG_PATH+"*"):
                if os.path.isfile(log_file):os.remove(log_file)
            configure_logging()
            global DETAIL_CACHE,DETAIL_CACHE_DIRTY,SOURCE_API_CACHE,SOURCE_RESULTS_CACHE,SOURCE_RESULT_METRICS
            with CACHE_LOCK:
                DETAIL_CACHE={};DETAIL_CACHE_DIRTY=0
            with SOURCE_API_CACHE_LOCK:SOURCE_API_CACHE={}
            with SOURCE_RESULTS_LOCK:
                SOURCE_RESULTS_CACHE={"version":1,"entries":{}};SOURCE_RESULT_METRICS={}
            self.p=default_profile();save_json_file(PROFILE_PATH,self.p);self.cv="";self.current=None
            self.q.set("");self.view_mode.set("todas");self.batch_selection.clear()
            self.refresh();self.info.set("Aplicativo limpo. Carregue o currículo da nova pessoa.")
            if parent and parent.winfo_exists():parent._managed_close()
            messagebox.showinfo("Limpeza concluída","O aplicativo está zerado e pronto para um novo currículo.")
            self.after(200,self.first_run_privacy_flow)
            return True
        except Exception:
            LOGGER.exception("Falha ao redefinir os dados do usuário")
            messagebox.showerror("Não foi possível limpar","Feche outras janelas do aplicativo e tente novamente.",parent=parent)
            return False

    def edit_profile(self):
        w,created=self.managed_window("configuracoes","Configurações","620x620",modal=True)
        if not created:return
        actions=ttk.Frame(w,padding=(16,7,16,12));actions.pack(side="bottom",fill="x")
        content=ttk.Frame(w);content.pack(side="top",fill="both",expand=True)
        canvas=tk.Canvas(content,bg=self.colors["bg"],highlightthickness=0,borderwidth=0)
        settings_scroll=ttk.Scrollbar(content,orient="vertical",command=canvas.yview)
        canvas.configure(yscrollcommand=settings_scroll.set)
        settings_scroll.pack(side="right",fill="y");canvas.pack(side="left",fill="both",expand=True)
        body=ttk.Frame(canvas,padding=16)
        body_window=canvas.create_window((0,0),window=body,anchor="nw")
        body.bind("<Configure>",lambda _event:canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.bind("<Configure>",lambda event:canvas.itemconfigure(body_window,width=event.width))
        w.bind("<MouseWheel>",lambda event:canvas.yview_scroll(-1*(event.delta//120),"units"))
        ttk.Label(body,text="Configurações",font=("Segoe UI",15,"bold")).pack(anchor="w")
        ttk.Label(body,text="Ajuste apenas o necessário. O restante é definido pelo currículo.").pack(anchor="w",pady=(2,14))

        cvf=ttk.LabelFrame(body,text="Currículo",padding=10);cvf.pack(fill="x",pady=(0,10))
        cvstatus=tk.StringVar(value="Currículo carregado" if self.cv.strip() else "Nenhum currículo carregado")
        ttk.Label(cvf,textvariable=cvstatus).pack(side="left")
        def load_cv_click():
            path=filedialog.askopenfilename(title="Escolha seu currículo",filetypes=[("Currículo","*.pdf *.docx *.txt"),("PDF","*.pdf"),("Word","*.docx"),("Texto","*.txt")])
            if not path:return
            try:
                text=parse_cv_file(path)
                if len(text.strip())<80:raise RuntimeError("O currículo parece estar vazio ou não pôde ser lido.")
                open(CV_PATH,"w",encoding="utf-8").write(text)
                ext=os.path.splitext(path)[1].lower(); dest=os.path.join(BASE_DIR,"curriculo_original"+ext)
                import shutil; shutil.copy2(path,dest)
                self.p["arquivo_curriculo_original"]=os.path.basename(dest)
                self.cv=text
                self.p["areas_estagio_manual"]=False
                prof=adapt_profile_to_cv(self.p,text);save_json_file(PROFILE_PATH,self.p)
                recalculated=self.recalculate_all_jobs();self.refresh()
                cvstatus.set("Currículo carregado — perfil atualizado")
                areas.set(", ".join(prof["areas"]))
                language.set(self.p.get("idioma_curriculo_detectado","Misto/indefinido"))
                english_choice.set(english_level(self.p))
                rebuild_internship_areas()
                international.set(international_search_enabled(self.p))
                entry_hint.set("Currículo de entrada identificado. Ative a busca ampliada abaixo se desejar."
                               if self.p.get("perfil_inicio_carreira_detectado",False) else "")
                messagebox.showinfo("Currículo pronto",f"Currículo lido com sucesso. A busca foi adaptada ao seu perfil e {recalculated} vaga(s) foram reavaliadas.",parent=w)
            except Exception as e:messagebox.showerror("Não foi possível carregar",str(e),parent=w)
        ttk.Button(cvf,text="Carregar meu currículo",command=load_cv_click).pack(side="right")
        areas=tk.StringVar(value=", ".join(self.p.get("areas_curriculo_detectadas",[])) or "Será identificado pelo currículo")
        ttk.Label(body,text="Áreas identificadas:").pack(anchor="w")
        ttk.Label(body,textvariable=areas,wraplength=560).pack(anchor="w")
        language=tk.StringVar(value=self.p.get("idioma_curriculo_detectado","Ainda não identificado"))
        ttk.Label(body,text="Idioma identificado:").pack(anchor="w",pady=(7,0))
        ttk.Label(body,textvariable=language).pack(anchor="w",pady=(0,12))
        ttk.Label(body,text="Meu nível de inglês:").pack(anchor="w")
        english_choice=tk.StringVar(value=english_level(self.p))
        english_combo=ttk.Combobox(body,textvariable=english_choice,state="readonly",width=22,
                                   values=ENGLISH_LEVELS)
        english_combo.pack(anchor="w",pady=(3,8))
        ttk.Label(body,text="Como buscar estágios:").pack(anchor="w")
        internship_choice=tk.StringVar(value=INTERNSHIP_MODE_LABELS[internship_search_mode(self.p)])
        internship_combo=ttk.Combobox(body,textvariable=internship_choice,state="readonly",width=34,
                                      values=list(INTERNSHIP_MODE_LABELS.values()))
        internship_combo.pack(anchor="w",pady=(3,7))
        internship_areas_frame=ttk.LabelFrame(body,text="Áreas de estágio",padding=8)
        internship_areas_frame.pack(fill="x",pady=(0,12))
        internship_area_vars={}
        extra_internship_areas=tk.StringVar()
        def rebuild_internship_areas():
            for child in internship_areas_frame.winfo_children():child.destroy()
            internship_area_vars.clear()
            selected=internship_selected_areas(self.p)
            detected=list(self.p.get("cursos_curriculo_detectados",[]))
            candidates=list(dict.fromkeys(detected+selected))
            if candidates:
                ttk.Label(internship_areas_frame,text="Selecione uma ou mais áreas:").pack(anchor="w",pady=(0,3))
                for course in candidates:
                    variable=tk.BooleanVar(value=course in selected);internship_area_vars[course]=variable
                    ttk.Checkbutton(internship_areas_frame,text=course.title(),variable=variable).pack(anchor="w")
            else:
                ttk.Label(internship_areas_frame,text="Nenhum curso identificado; informe a área desejada.").pack(anchor="w")
            ttk.Label(internship_areas_frame,text="Outras áreas (separadas por vírgula)").pack(anchor="w",pady=(5,0))
            ttk.Entry(internship_areas_frame,textvariable=extra_internship_areas).pack(fill="x",pady=(3,0))
        rebuild_internship_areas()
        def sync_internship_area_state(_event=None):
            if internship_search_mode({"modo_estagios":next(
                    (key for key,label in INTERNSHIP_MODE_LABELS.items() if label==internship_choice.get()),"nao_buscar")})=="nao_buscar":
                internship_areas_frame.pack_forget()
            elif not internship_areas_frame.winfo_manager():
                internship_areas_frame.pack(fill="x",pady=(0,12),after=internship_combo)
        internship_combo.bind("<<ComboboxSelected>>",sync_internship_area_state)
        sync_internship_area_state()
        entry_hint=tk.StringVar(value=("Currículo de entrada identificado. Ative a busca ampliada abaixo se desejar."
            if self.p.get("perfil_inicio_carreira_detectado",False) else ""))
        ttk.Label(body,textvariable=entry_hint,foreground=self.colors["blue_dark"],wraplength=580).pack(anchor="w",pady=(0,7))

        locf=ttk.LabelFrame(body,text="Onde quero trabalhar",padding=10);locf.pack(fill="x",pady=(0,10))
        ttk.Label(locf,text="Cidades para presencial/híbrido (separadas por vírgula)").pack(anchor="w")
        cities=tk.StringVar(value=", ".join(self.p.get("cidades_presencial",self.p.get("cidades_presencial_hibrido",[]))))
        location_row=ttk.Frame(locf);location_row.pack(fill="x",pady=(3,7))
        ttk.Entry(location_row,textvariable=cities).pack(side="left",fill="x",expand=True)
        ttk.Label(location_row,text="UF").pack(side="left",padx=(10,4))
        state=tk.StringVar(value=str(self.p.get("estado_local","")).upper())
        ttk.Entry(location_row,textvariable=state,width=4).pack(side="left")
        ttk.Label(locf,text="Localidades que não quero receber (separadas por vírgula)").pack(anchor="w",pady=(2,0))
        excluded_locations=tk.StringVar(value=", ".join(self.p.get("localidades_excluidas",[])))
        ttk.Entry(locf,textvariable=excluded_locations).pack(fill="x",pady=(3,7))
        remote=tk.BooleanVar(value=self.p.get("aceitar_remoto",True));ttk.Checkbutton(locf,text="Aceito vagas remotas",variable=remote).pack(anchor="w")
        pcd_vacancies=tk.BooleanVar(value=pcd_vacancies_enabled(self.p))
        ttk.Checkbutton(locf,text="Buscar vagas para PCD",variable=pcd_vacancies).pack(anchor="w",pady=(4,0))
        international=tk.BooleanVar(value=international_search_enabled(self.p))
        international_check=ttk.Checkbutton(
            locf,text="Buscar vagas remotas internacionais (requer inglês fluente)",variable=international)
        international_check.pack(anchor="w",pady=(4,0))
        def sync_international_option(_event=None):
            fluent=normalize_english_level(english_choice.get())=="Fluente"
            if not fluent:international.set(False)
            international_check.configure(state="normal" if fluent else "disabled")
        english_combo.bind("<<ComboboxSelected>>",sync_international_option)
        sync_international_option()
        entry_market=tk.BooleanVar(value=bool(self.p.get("buscar_vagas_inicio_carreira",False)))
        ttk.Checkbutton(locf,text="Buscar vagas sem experiência / primeiro emprego",
                        variable=entry_market).pack(anchor="w",pady=(4,0))
        apprentice=tk.BooleanVar(value=bool(self.p.get("buscar_jovem_aprendiz",False)))
        ttk.Checkbutton(locf,text="Buscar Jovem Aprendiz (normalmente até 24 anos)",
                        variable=apprentice).pack(anchor="w",pady=(4,0))

        advanced=ttk.Frame(body)
        navf=ttk.LabelFrame(advanced,text="Navegador para candidaturas",padding=10);navf.pack(fill="x",pady=(8,10))
        ttk.Label(navf,text="Automático é recomendado. O app tenta usar um navegador instalado e, se necessário, o Chromium.").pack(anchor="w")
        browser_choice=tk.StringVar(value=self.p.get("navegador_automacao","Automático"))
        browser_combo=ttk.Combobox(navf,textvariable=browser_choice,state="readonly",width=24,
                                  values=["Automático","Google Chrome","Microsoft Edge","Brave","Chromium interno","Firefox"])
        browser_combo.pack(anchor="w",pady=(5,0))


        agef=ttk.Frame(advanced);agef.pack(fill="x",pady=(2,10))
        ttk.Label(agef,text="Mostrar vagas publicadas nos últimos").pack(side="left")
        days=tk.StringVar(value=str(self.p.get("idade_maxima_dias",self.p.get("max_age_days",60))))
        ttk.Entry(agef,textvariable=days,width=5).pack(side="left",padx=5);ttk.Label(agef,text="dias").pack(side="left")

        def save():
            pcd_enabled=bool(pcd_vacancies.get())
            if pcd_enabled and not self.pcd_consent_valid():
                pcd_enabled=self.request_pcd_consent(w)
                pcd_vacancies.set(pcd_enabled)
            if not pcd_enabled:
                self.p["consentimento_pcd_versao"]=""
                self.p["consentimento_pcd_em"]=""
            vals=[x.strip() for x in cities.get().split(",") if x.strip()]
            self.p["cidades_presencial"]=vals;self.p["cidades_presencial_hibrido"]=vals
            self.p["estado_local"]=state.get().strip().upper()[:2]
            self.p["localidades_excluidas"]=list(dict.fromkeys(
                value.strip() for value in excluded_locations.get().split(",") if value.strip()))
            self.p["aceitar_remoto"]=bool(remote.get())
            selected_internship_mode=next(
                (key for key,label in INTERNSHIP_MODE_LABELS.items() if label==internship_choice.get()),"nao_buscar")
            selected_areas=[course for course,variable in internship_area_vars.items() if variable.get()]
            selected_areas += [value.strip() for value in extra_internship_areas.get().split(",") if value.strip()]
            self.p["modo_estagios"]=selected_internship_mode
            self.p["buscar_estagios"]=selected_internship_mode!="nao_buscar"
            self.p["areas_estagio"]=list(dict.fromkeys(selected_areas))
            self.p["areas_estagio_manual"]=True
            self.p["buscar_vagas_pcd"]=pcd_enabled
            self.p["nivel_ingles"]=normalize_english_level(english_choice.get())
            self.p["nivel_ingles_manual"]=True
            self.p["buscar_vagas_internacionais"]=(bool(international.get()) and
                                                    self.p["nivel_ingles"]=="Fluente")
            self.p["preferencia_internacional_manual"]=True
            self.p["buscar_vagas_inicio_carreira"]=bool(entry_market.get())
            self.p["buscar_jovem_aprendiz"]=bool(apprentice.get())
            self.p["descartar_vagas_exclusivas_pcd"]=not pcd_enabled
            self.p["mostrar_compativeis_fora_regiao"]=False
            browser_map={"Automático":"automatico","Google Chrome":"chrome","Microsoft Edge":"edge",
                         "Brave":"brave","Chromium interno":"chromium","Firefox":"firefox"}
            self.p["navegador_automacao"]=browser_map.get(browser_choice.get(),"automatico")
            try:
                selected_days=max(1,int(days.get()))
                self.p["idade_maxima_dias"]=selected_days
                self.p["idade_maxima_vaga_dias"]=selected_days
            except:pass
            if self.cv.strip():adapt_profile_to_cv(self.p,self.cv)
            save_json_file(PROFILE_PATH,self.p)
            self.apply_publication_age_preference()
            self.apply_pcd_preference(pcd_enabled)
            self.apply_internship_mode(selected_internship_mode)
            self.apply_apprentice_preference(bool(apprentice.get()))
            self.apply_international_preference(bool(international.get()))
            self.apply_excluded_locations()
            self.recalculate_all_jobs();self.refresh()
            messagebox.showinfo("Configurações","Alterações salvas.",parent=w);w._managed_close()
        maintenance=ttk.LabelFrame(advanced,text="Manutenção",padding=8);maintenance.pack(fill="x",pady=(0,8))
        maintenance_row=ttk.Frame(maintenance);maintenance_row.pack(fill="x")
        ttk.Button(maintenance_row,text="Atualizar descrições",command=lambda:self.update_pending_descriptions()).pack(side="left")
        ttk.Button(maintenance_row,text="Reavaliar vagas",command=lambda:self.recalculate_with_notice(w)).pack(side="left",padx=6)
        ttk.Button(maintenance_row,text="Limpar tudo / trocar currículo",style="Danger.TButton",
                   command=lambda:self.reset_for_new_user(w)).pack(side="right")

        advanced_text=tk.StringVar(value="Mostrar opções avançadas")
        def toggle_advanced():
            if advanced.winfo_manager():
                advanced.pack_forget();advanced_text.set("Mostrar opções avançadas");self.center_window(w,620,620)
            else:
                advanced.pack(fill="x");advanced_text.set("Ocultar opções avançadas");self.center_window(w,680,700)
                canvas.update_idletasks();canvas.yview_moveto(1.0)

        ttk.Button(actions,textvariable=advanced_text,style="Soft.TButton",command=toggle_advanced).pack(side="left")
        ttk.Button(actions,text="Privacidade e termos",style="Soft.TButton",
                   command=lambda:self.show_privacy_notice(False)).pack(side="left",padx=6)
        ttk.Button(actions,text="Salvar",style="Primary.TButton",command=save).pack(side="right")

if __name__=="__main__":
    App().mainloop()
